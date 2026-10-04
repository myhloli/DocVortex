"""Convert Excel 97–2003 BIFF workbook to DocVortex paginated model-list."""

from __future__ import annotations

import collections
import re
from typing import Any, BinaryIO

from loguru import logger
from openpyxl import Workbook  # type: ignore[reportMissingModuleSource]
from openpyxl.cell.rich_text import CellRichText, TextBlock  # type: ignore[reportMissingModuleSource]
from openpyxl.cell.text import InlineFont  # type: ignore[reportMissingModuleSource]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[reportMissingModuleSource]

from ..errors import LegacyOfficeEncryptedError, LegacyOfficeResourceLimitError
from ..limits import MAX_GRID_SLOTS
from ..equation.mtef import decode_equation_native
from ..legacy.ole import BoundedOleReader
from ..spreadsheet.html import render_spreadsheet_table
from ..spreadsheet.models import AnchoredBlock, ExcelTable, FormulaMap, SheetImage
from ..spreadsheet.projector import SpreadsheetProjector
from ..streams import read_stream_bytes_from_start
from .....schema import BlockType

from .models import XlsChart, XlsChartSheet, XlsRichText, XlsSheet, XlsWorkbook
from .parser import parse_xls_workbook

_XLS_EQUATION_STREAM_RE = re.compile(
    r"^(MBD[0-9A-F]{8})/Equation Native$",
    re.IGNORECASE,
)


def _read_embedded_equations(ole: BoundedOleReader) -> dict[str, str]:
    """Read XLS embedding storages Equation Native that can be safely decoded."""

    equations: dict[str, str] = {}
    for stream_name in ole.stream_names(prefix="MBD"):
        match = _XLS_EQUATION_STREAM_RE.match(stream_name)
        if match is None:
            continue
        storage = match.group(1)
        latex = decode_equation_native(ole.read_stream(stream_name))
        if latex is None:
            logger.warning(
                "XLS_MTEF_FALLBACK: storage={!r} has an invalid or unsupported Equation Native stream",
                storage,
            )
            continue
        equations.setdefault(storage, latex)
    return equations


def _inline_font(rich_text: XlsRichText, start: int) -> InlineFont | None:
    """Returns the openpyxl inline font covering the specified character position."""

    for run in rich_text.runs:
        if run.start <= start < run.end:
            style = run.style
            return InlineFont(
                b=style.bold,
                i=style.italic,
                u="single" if style.underline else None,
                strike=style.strike,
                vertAlign=("superscript" if style.superscript else "subscript" if style.subscript else None),
            )
    return None


def _rich_text_boundaries(value: XlsRichText) -> list[int]:
    """Collect the starting and ending boundaries of rich text and clip it to the valid character range."""

    boundaries = {0, len(value.text)}
    for run in value.runs:
        boundaries.add(max(0, min(run.start, len(value.text))))
        boundaries.add(max(0, min(run.end, len(value.text))))
    return sorted(boundaries)


def _to_openpyxl_rich_text(value: XlsRichText) -> str | CellRichText:
    """Convert internal rich text to XlsxConverter already supported by CellRichText."""

    if not value.runs:
        return value.text
    parts: list[str | TextBlock] = []
    boundaries = _rich_text_boundaries(value)
    for start, end in zip(boundaries, boundaries[1:]):
        if start >= end:
            continue
        text = value.text[start:end]
        font = _inline_font(value, start)
        parts.append(TextBlock(font, text) if font is not None else text)
    return CellRichText(parts)


class _XlsPageBuilder(SpreadsheetProjector):
    """Reuse existing meshes with HTML projections using the lightweight openpyxl worksheet adapter."""

    def __init__(self, workbook_model: XlsWorkbook) -> None:
        """Initialize parsing results mapping and full workbook grid budget."""

        super().__init__(include_hidden_sheets=False)
        self.workbook_model = workbook_model
        self._sheet_by_title: dict[str, XlsSheet] = {}
        self._active_xls_sheet: XlsSheet | None = None
        self._used_cells: set[tuple[int, int]] = set()
        self._grid_slots = 0

    def _build_openpyxl_workbook(self) -> Workbook:
        """Project the BIFF semantic model into a non-displaced worksheet object."""

        workbook = Workbook()
        merge_grid_slots = 0
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)
        for sheet_model in self.workbook_model.sheets:
            worksheet = workbook.create_sheet(sheet_model.name)
            self._sheet_by_title[worksheet.title] = sheet_model
            if not sheet_model.visible:
                worksheet.sheet_state = Worksheet.SHEETSTATE_HIDDEN
            for cell_model in sheet_model.cells.values():
                cell = worksheet.cell(row=cell_model.row + 1, column=cell_model.col + 1)
                cell.value = _to_openpyxl_rich_text(cell_model.value)
                if cell_model.hyperlink:
                    cell.hyperlink = cell_model.hyperlink
            for row_first, col_first, row_last, col_last in sheet_model.merges:
                merge_slots = (row_last - row_first + 1) * (col_last - col_first + 1)
                if merge_slots > MAX_GRID_SLOTS - merge_grid_slots:
                    raise LegacyOfficeResourceLimitError(f"workbook extent exceeds max_grid_slots={MAX_GRID_SLOTS}")
                merge_grid_slots += merge_slots
                worksheet.merge_cells(
                    start_row=row_first + 1,
                    start_column=col_first + 1,
                    end_row=row_last + 1,
                    end_column=col_last + 1,
                )
        return workbook

    def _ensure_workbook(self) -> Workbook:
        """Lazily construct the openpyxl adaptation workbook once."""

        workbook = getattr(self, "workbook", None)
        if workbook is None:
            workbook = self._build_openpyxl_workbook()
            self.workbook = workbook
        return workbook

    def render_chart_selection(
        self,
        sheet_name: str,
        rows: tuple[int, ...] | list[int],
        cols: tuple[int, ...] | list[int],
    ) -> str | None:
        """Render the specified worksheet row and column selection as shared chart HTML."""

        if not rows or not cols:
            return None
        grid_slots = len(rows) * len(cols)
        self._grid_slots += grid_slots
        if self._grid_slots > MAX_GRID_SLOTS:
            raise LegacyOfficeResourceLimitError(f"workbook extent exceeds max_grid_slots={MAX_GRID_SLOTS}")
        workbook = self._ensure_workbook()
        if sheet_name not in workbook.sheetnames:
            return None
        sheet = workbook[sheet_name]
        table = self._build_synthetic_table_from_sheet_selection(
            sheet,
            list(rows),
            list(cols),
        )
        return render_spreadsheet_table(table)

    def _chart_sheet_page(self, chart_sheet: XlsChartSheet) -> list[dict[str, Any]]:
        """Project standalone chart sheet to a logical page containing only chart block."""

        if chart_sheet.source_sheet_name is None:
            return []
        content = self.render_chart_selection(
            chart_sheet.source_sheet_name,
            chart_sheet.source_rows,
            chart_sheet.source_cols,
        )
        return [{"type": BlockType.CHART, "content": content}] if content else []

    def build_pages(self) -> list[list[dict[str, Any]]]:
        """Visible worksheet/chart sheet pages are generated in the order of the original directory."""

        self.workbook = self._build_openpyxl_workbook()
        pages_by_name: dict[str, list[dict[str, Any]]] = {}
        for worksheet in self._iter_sheets_to_convert():
            self._active_xls_sheet = self._sheet_by_title.get(worksheet.title)
            self.cur_page = []
            self._convert_sheet(worksheet)
            pages_by_name[worksheet.title] = self.cur_page

        ordered_pages: list[tuple[int, int, str, list[dict[str, Any]]]] = []
        for index, sheet in enumerate(self.workbook_model.sheets):
            if not sheet.visible:
                continue
            order = sheet.order if sheet.order >= 0 else index
            ordered_pages.append((order, 0, sheet.name, pages_by_name.get(sheet.name, [])))
        for index, chart_sheet in enumerate(self.workbook_model.chart_sheets):
            if not chart_sheet.visible:
                continue
            order = chart_sheet.order if chart_sheet.order >= 0 else len(ordered_pages) + index
            ordered_pages.append((order, 1, chart_sheet.name, self._chart_sheet_page(chart_sheet)))
        ordered_pages.sort(key=lambda item: (item[0], item[1]))
        sheet_pages = [(name, page) for _order, _kind, name, page in ordered_pages]
        if self._should_emit_sheet_titles([page for _, page in sheet_pages]):
            self._prepend_sheet_titles(sheet_pages)
        return [page for _, page in sheet_pages]

    def _collect_sheet_images(self, sheet: Worksheet) -> list[SheetImage]:
        """Returns the image that the parser has bound to the current sheet."""

        sheet_model = self._sheet_by_title.get(sheet.title)
        if sheet_model is None:
            return []
        return [
            SheetImage(
                anchor=(image.row, image.col),
                image_base64=image.image_base64,
            )
            for image in sheet_model.images
        ]

    def _map_math_formulas_to_cells(self, sheet: Worksheet) -> FormulaMap:
        """Map the legacy Equation Editor formula to the table cell anchor."""

        math_map: dict[tuple[int, int], list[str]] = collections.defaultdict(list)
        sheet_model = self._sheet_by_title.get(sheet.title)
        if sheet_model is None:
            return math_map
        for equation in sheet_model.equations:
            math_map[(equation.row, equation.col)].append(equation.latex)
        return math_map

    def _find_tables_in_sheet(
        self,
        sheet: Worksheet,
    ) -> tuple[set[tuple[int, int]], list[tuple[tuple[int, int], int, dict]]]:
        """The record table has absorbed the coordinates for independent formula deduplication."""

        used_cells, artifacts = super()._find_tables_in_sheet(sheet)
        self._used_cells = used_cells
        return used_cells, artifacts

    def _find_data_tables(self, sheet: Worksheet) -> list[ExcelTable]:
        """Reuse XLSX area discovery and cumulative constraints on the actual materialized mesh."""

        tables = super()._find_data_tables(sheet)
        self._grid_slots += sum(table.num_rows * table.num_cols for table in tables)
        if self._grid_slots > MAX_GRID_SLOTS:
            raise LegacyOfficeResourceLimitError(f"workbook extent exceeds max_grid_slots={MAX_GRID_SLOTS}")
        return tables

    def _chart_block(self, sheet: Worksheet, chart: XlsChart) -> dict[str, Any] | None:
        """Convert a simple chart reference into a data table and keep the preview image available."""

        if chart.source_rows and chart.source_cols:
            content = self.render_chart_selection(
                sheet.title,
                chart.source_rows,
                chart.source_cols,
            )
            if content is None:
                return None
            block = {
                "type": BlockType.CHART,
                "content": content,
            }
            if chart.image_base64:
                block["image_base64"] = chart.image_base64
            return block
        if chart.image_base64:
            return {
                "type": BlockType.CHART,
                "content": "",
                "image_base64": chart.image_base64,
            }
        return None

    def _find_charts_in_sheet(
        self,
        sheet: Worksheet,
    ) -> list[AnchoredBlock]:
        """Press cell anchor to output legacy chart blocks of the current worksheet."""

        sheet_model = self._sheet_by_title.get(sheet.title)
        if sheet_model is None:
            return []
        artifacts = []
        for order, chart in enumerate(sheet_model.charts):
            block = self._chart_block(sheet, chart)
            if block is None:
                continue
            artifacts.append(((chart.row, chart.col), 10_000 + order, block))
        for order, equation in enumerate(sheet_model.equations):
            anchor = (equation.row, equation.col)
            if anchor in self._used_cells:
                continue
            artifacts.append(
                (
                    anchor,
                    20_000 + order,
                    {"type": BlockType.EQUATION, "content": equation.latex},
                )
            )
        return artifacts


def render_xls_chart_html(
    workbook: XlsWorkbook,
    sheet_name: str,
    rows: tuple[int, ...] | list[int],
    cols: tuple[int, ...] | list[int],
) -> str | None:
    """Stable HTML projection of XLS multiplexed for DOC/PPT embedded chart."""

    return _XlsPageBuilder(workbook).render_chart_selection(
        sheet_name,
        rows,
        cols,
    )


class XlsConverter:
    """Convert Excel 97–2003 OLE/BIFF binary stream to model-list."""

    def __init__(self) -> None:
        """Initialize empty paged output."""

        self.pages: list[list[dict[str, Any]]] = []

    def convert(self, file_binary: BinaryIO) -> None:
        """Read Workbook/Book stream, parse BIFF and generate visible sheet page."""

        file_bytes = read_stream_bytes_from_start(file_binary)
        with BoundedOleReader(file_bytes) as ole:
            if ole.has_stream("EncryptionInfo") or ole.has_stream("EncryptedPackage"):
                raise LegacyOfficeEncryptedError("password-protected XLS is not supported")
            if ole.has_stream("Workbook"):
                workbook_stream = ole.read_stream("Workbook")
            else:
                workbook_stream = ole.read_stream("Book")
            native_equations = _read_embedded_equations(ole)
            workbook = parse_xls_workbook(
                workbook_stream,
                native_equations=native_equations,
            )
        builder = _XlsPageBuilder(workbook)
        self.pages = builder.build_pages()
        logger.debug("XLS parsing produced {} visible sheet pages", len(self.pages))
