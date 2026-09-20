"""XLS 与 XLSX 复用的中立工作表投影器。"""

from __future__ import annotations

import collections
import html
from collections.abc import Callable, Iterator
from typing import Any

from loguru import logger
from openpyxl.cell.rich_text import CellRichText
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .....schema import BlockType
from ..._shared.hyperlink import OFFICE_EXTERNAL_HYPERLINK_SCHEMES, sanitize_hyperlink_target
from .....content.spans import text_spans
from .html import EQUATION_BOOKENDS, render_spreadsheet_table
from .models import AnchoredBlock, DataRegion, ExcelCell, ExcelTable, FormulaMap, SheetImage
from .region_discovery import (
    ConnectedRegion,
    discover_connected_regions,
    keep_maximal_by_semantic_sets,
    select_best_gap_candidate,
)


class _MergedCellLookup:
    """按行缓存合并单元格范围，避免解析时反复扫描 openpyxl 合并区域。"""

    def __init__(self, sheet: Worksheet):
        """从工作表合并区域构建 0-based 坐标索引。"""
        self._merged_row_intervals: dict[int, list[tuple[int, int]]] = collections.defaultdict(list)
        self._hidden_row_intervals: dict[int, list[tuple[int, int]]] = collections.defaultdict(list)
        self._anchor_spans: dict[tuple[int, int], tuple[int, int]] = {}

        for merged in sheet.merged_cells.ranges:
            min_row = merged.min_row - 1
            max_row = merged.max_row - 1
            min_col = merged.min_col - 1
            max_col = merged.max_col - 1

            self._anchor_spans[(min_row, min_col)] = (
                max_row - min_row + 1,
                max_col - min_col + 1,
            )

            for row in range(min_row, max_row + 1):
                self._merged_row_intervals[row].append((min_col, max_col))
                hidden_start_col = min_col + 1 if row == min_row else min_col
                if hidden_start_col <= max_col:
                    self._hidden_row_intervals[row].append((hidden_start_col, max_col))

        for intervals in self._merged_row_intervals.values():
            intervals.sort()
        for intervals in self._hidden_row_intervals.values():
            intervals.sort()

    @staticmethod
    def _contains_interval(
        row_intervals: dict[int, list[tuple[int, int]]],
        row: int,
        col: int,
    ) -> bool:
        """判断 0-based 坐标是否落入指定行的任一列区间。"""
        for start_col, end_col in row_intervals.get(row, []):
            if start_col <= col <= end_col:
                return True
            if start_col > col:
                break
        return False

    def contains_merged_cell(self, row: int, col: int) -> bool:
        """判断 0-based 坐标是否属于任一合并区域。"""
        return self._contains_interval(self._merged_row_intervals, row, col)

    def is_hidden_merged_cell(self, row: int, col: int) -> bool:
        """判断 0-based 坐标是否为合并区域内非左上角的隐藏格。"""
        return self._contains_interval(self._hidden_row_intervals, row, col)

    def get_anchor_span(self, row: int, col: int) -> tuple[int, int]:
        """返回合并区域左上角坐标对应的 rowspan/colspan，非合并锚点返回 1x1。"""
        return self._anchor_spans.get((row, col), (1, 1))


class SpreadsheetProjector:
    """把 openpyxl 工作表投影为稳定的分页 model-list。"""

    def __init__(
        self,
        *,
        treat_singleton_as_text: bool = True,
        gap_tolerance: int | None = None,
        include_hidden_sheets: bool = False,
    ) -> None:
        """保存工作表投影配置并初始化无格式专属依赖的运行状态。"""
        self.treat_singleton_as_text = treat_singleton_as_text
        self.gap_tolerance = gap_tolerance
        self.include_hidden_sheets = include_hidden_sheets
        self._reset_projection_state()

    def _reset_projection_state(self) -> None:
        """重置工作簿、分页和逐 sheet 的共享投影状态。"""
        self.workbook: Workbook | None = None
        self.pages: list[list[dict[str, Any]]] = []
        self.cur_page: list[dict[str, Any]] = []
        self.math_map: FormulaMap = {}
        self.sheet_images: list[SheetImage] = []
        self.table_image_map: dict[tuple[int, int], list[str]] = collections.defaultdict(list)
        self._merged_cell_lookup_cache: dict[int, _MergedCellLookup] = {}

    def _prepare_sheet_assets(self, sheet: Worksheet) -> None:
        """准备普通公式和图片，并构建表格单元格使用的媒体映射。"""
        self.math_map = self._map_math_formulas_to_cells(sheet)
        self.sheet_images = self._collect_sheet_images(sheet)
        self.table_image_map = collections.defaultdict(list)
        for image in self.sheet_images:
            row, col = image.anchor
            if row is None or col is None:
                continue
            if image.latex:
                self.table_image_map[(row, col)].append(EQUATION_BOOKENDS.format(EQ=image.latex))
            elif image.image_base64:
                self.table_image_map[(row, col)].append(f'<img src="{image.image_base64}" />')

    def _convert_sheet(self, sheet: Worksheet) -> None:
        """按表格、图表、附加素材和独立图片的稳定顺序投影一个工作表。"""
        self._prepare_sheet_assets(sheet)
        used_cells, visual_artifacts = self._find_tables_in_sheet(sheet)
        visual_artifacts.extend(self._find_charts_in_sheet(sheet))
        visual_artifacts.extend(self._find_additional_visual_artifacts(used_cells))
        for _, _, block in sorted(
            visual_artifacts,
            key=lambda item: (item[0][0], item[0][1], item[1]),
        ):
            self.cur_page.append(block)
        self._find_images_in_sheet(used_cells)

    def _map_math_formulas_to_cells(self, sheet: Worksheet) -> FormulaMap:
        """返回当前工作表按 0-based cell anchor 分组的公式。"""
        return {}

    def _collect_sheet_images(self, sheet: Worksheet) -> list[SheetImage]:
        """返回当前工作表按 anchor 排序的图片或图片公式。"""
        return []

    def _find_charts_in_sheet(self, sheet: Worksheet) -> list[AnchoredBlock]:
        """返回当前工作表的格式专属图表 blocks。"""
        return []

    def _find_additional_visual_artifacts(
        self,
        used_cells: set[tuple[int, int]],
    ) -> list[AnchoredBlock]:
        """返回未被表格吸收的格式专属公式或图片 blocks。"""
        return []

    def _resolve_cell_image(self, raw_cell_text: str) -> str:
        """解析格式专属的单元格图片函数，默认不产生媒体。"""
        return ""

    def _iter_sheets_to_convert(self) -> Iterator[Worksheet]:
        """按工作簿顺序遍历允许输出的可见工作表。"""
        if self.workbook is None:
            return

        for sheet in self.workbook.worksheets:
            if not self.include_hidden_sheets and sheet.sheet_state != Worksheet.SHEETSTATE_VISIBLE:
                logger.debug(f"跳过隐藏工作表：{sheet.title}")
                continue
            yield sheet

    @staticmethod
    def _build_sheet_title_block(sheet_title: str) -> dict:
        """构造工作表标题块，复用 Office 标题渲染链路输出 Markdown 标题。"""
        return {
            "type": BlockType.PARAGRAPH_TITLE,
            "level": 2,
            "content": text_spans(sheet_title),
        }

    @staticmethod
    def _should_emit_sheet_titles(pages: list[list[dict]]) -> bool:
        """仅当存在多个非空输出 sheet 时才添加标题，避免单表或空表噪声。"""
        return sum(1 for page in pages if page) > 1

    def _prepend_sheet_titles(self, sheet_pages: list[tuple[str, list[dict]]]) -> None:
        """将 sheet 标题插入每个非空 page 开头，不参与表格/图表视觉排序。"""
        for sheet_title, page in sheet_pages:
            if not page:
                continue
            page.insert(0, self._build_sheet_title_block(sheet_title))

    def _get_block_sort_anchor(self, row: int | None, col: int | None) -> tuple[int, int]:
        """把缺失 anchor 稳定放到全部有效工作表坐标之后。"""
        if row is None or col is None:
            return (10**9, 10**9)
        return row, col

    def _build_block_from_excel_table(self, excel_table: ExcelTable) -> dict:
        """按 singleton 规则把表格 IR 投影为文本或表格 block。"""
        if self.treat_singleton_as_text and len(excel_table.data) == 1 and self._can_render_singleton_as_text(excel_table):
            return {
                "type": BlockType.TEXT,
                "content": text_spans(excel_table.data[0].text),
            }

        return {
            "type": BlockType.TABLE,
            "content": render_spreadsheet_table(excel_table),
        }

    def _find_tables_in_sheet(self, sheet: Worksheet) -> tuple[set[tuple[int, int]], list[tuple[tuple[int, int], int, dict]]]:
        """发现当前 sheet 表格并返回已吸收 cell 与锚定 blocks。"""
        used_cells = set()
        visual_artifacts = []
        if self.workbook is not None:
            tables = self._find_data_tables(sheet)  # 检测工作表中的所有数据表格

            for order, excel_table in enumerate(tables):
                # Record used cells
                anchor_c, anchor_r = excel_table.anchor
                for cell in excel_table.data:
                    source_row, source_col = self._resolve_excel_cell_source_position(
                        excel_table.anchor,
                        cell,
                    )
                    used_cells.add((source_row, source_col))

                visual_artifacts.append(
                    (
                        self._get_block_sort_anchor(anchor_r, anchor_c),
                        order,
                        self._build_block_from_excel_table(excel_table),
                    )
                )

        return used_cells, visual_artifacts

    def _build_excel_cell(
        self,
        sheet: Worksheet,
        display_row: int,
        display_col: int,
        source_row: int,
        source_col: int,
        row_span: int = 1,
        col_span: int = 1,
    ) -> ExcelCell:
        """把源工作表单元格完整物化为中立 ExcelCell。"""
        cell = sheet.cell(row=source_row + 1, column=source_col + 1)
        raw_cell_text = str(cell.value) if cell.value is not None else ""
        cell_text = ""
        text_is_html = False
        media_content = []
        if "DISPIMG" in raw_cell_text:
            cell_image = self._resolve_cell_image(raw_cell_text)
            if cell_image:
                media_content.append(cell_image)
        else:
            cell_text, text_is_html = self._cell_value_to_html(cell)
        media_content.extend(self.table_image_map.get((source_row, source_col), []))

        return ExcelCell(
            row=display_row,
            col=display_col,
            text=cell_text,
            row_span=row_span,
            col_span=col_span,
            styles=self._extract_cell_style(cell),
            media=media_content,
            equations=list(self.math_map.get((source_row, source_col), [])),
            text_is_html=text_is_html,
            source_row=source_row,
            source_col=source_col,
        )

    def _build_synthetic_table_from_sheet_selection(self, sheet: Worksheet, rows: list[int], cols: list[int]) -> ExcelTable:
        """把指定源行列选择物化为紧凑的表格 IR。"""
        selected_coords = {(row, col) for row in rows for col in cols}
        hidden_merge_cells = set()
        merge_spans = {}

        for mr in sheet.merged_cells.ranges:
            top_left = (mr.min_row - 1, mr.min_col - 1)
            if top_left not in selected_coords:
                continue

            selected_rows = [row for row in rows if mr.min_row - 1 <= row <= mr.max_row - 1]
            selected_cols = [col for col in cols if mr.min_col - 1 <= col <= mr.max_col - 1]
            if not selected_rows or not selected_cols:
                continue

            merge_spans[top_left] = (len(selected_rows), len(selected_cols))
            for row in selected_rows:
                for col in selected_cols:
                    if (row, col) != top_left:
                        hidden_merge_cells.add((row, col))

        data = []
        for display_row, source_row in enumerate(rows):
            for display_col, source_col in enumerate(cols):
                if (source_row, source_col) in hidden_merge_cells:
                    continue

                row_span, col_span = merge_spans.get((source_row, source_col), (1, 1))
                data.append(
                    self._build_excel_cell(
                        sheet,
                        display_row,
                        display_col,
                        source_row,
                        source_col,
                        row_span=row_span,
                        col_span=col_span,
                    )
                )

        return ExcelTable(
            anchor=(cols[0], rows[0]),
            num_rows=len(rows),
            num_cols=len(cols),
            data=data,
        )

    def _resolve_excel_cell_source_position(
        self,
        table_anchor: tuple[int, int],
        excel_cell: ExcelCell | None,
        row: int | None = None,
        col: int | None = None,
    ) -> tuple[int, int]:
        """优先使用显式源坐标，否则通过表格 anchor 还原源坐标。"""
        if excel_cell is not None:
            if excel_cell.source_row is not None and excel_cell.source_col is not None:
                return excel_cell.source_row, excel_cell.source_col
            row = excel_cell.row
            col = excel_cell.col

        if row is None or col is None:
            raise ValueError("row and col must be provided when excel_cell is None")

        return table_anchor[1] + row, table_anchor[0] + col

    def _can_render_singleton_as_text(self, excel_table: ExcelTable) -> bool:
        """判断单格表是否可安全降级为普通文本 block。"""
        cell = excel_table.data[0]
        return cell.row_span == 1 and cell.col_span == 1 and not cell.media and not cell.text_is_html and not cell.equations

    def _cell_has_semantic_content(self, excel_table: ExcelTable, cell: ExcelCell) -> bool:
        """判断单元格是否包含文本、媒体或公式语义。"""
        return bool(cell.text.strip() or any(media.strip() for media in cell.media) or cell.equations)

    def _get_table_semantic_positions(self, excel_table: ExcelTable) -> set[tuple[int, int]]:
        """返回表格内具有语义内容的源工作表坐标。"""
        semantic_positions = set()
        for cell in excel_table.data:
            if not self._cell_has_semantic_content(excel_table, cell):
                continue
            semantic_positions.add(
                self._resolve_excel_cell_source_position(
                    excel_table.anchor,
                    excel_cell=cell,
                )
            )
        return semantic_positions

    def _filter_semantic_subset_tables(self, tables: list[ExcelTable]) -> list[ExcelTable]:
        """删除语义坐标严格包含于其它候选的重复表格。"""
        semantic_sets = [self._get_table_semantic_positions(table) for table in tables]
        return [tables[index] for index in keep_maximal_by_semantic_sets(semantic_sets)]

    def _sheet_semantic_predicates(
        self, sheet: Worksheet
    ) -> tuple[Callable[[int, int], bool], Callable[[int, int], tuple[int, int]]]:
        """构造 gap 评分使用的语义内容与合并跨度查询。

        语义判断与 _build_excel_cell 物化结果保持一致：普通值取非空白文本，
        DISPIMG 公式仅在解析出媒体时算内容，另叠加锚定图片与公式映射。
        """
        merged_lookup = self._get_merged_cell_lookup(sheet)

        def has_semantic_content(row: int, col: int) -> bool:
            """判断源坐标是否含文本、媒体或公式语义。"""
            cell = sheet._cells.get((row + 1, col + 1))
            if cell is not None and cell.value is not None:
                raw_text = str(cell.value)
                if "DISPIMG" in raw_text:
                    resolved_image = self._resolve_cell_image(raw_text)
                    if resolved_image and resolved_image.strip():
                        return True
                elif raw_text.strip():
                    return True
            if any(media.strip() for media in self.table_image_map.get((row, col), [])):
                return True
            return bool(self.math_map.get((row, col)))

        return has_semantic_content, merged_lookup.get_anchor_span

    def _select_best_gap_candidate(self, sheet: Worksheet) -> tuple[int, float, list[ExcelTable]]:
        """按固定候选与偏好顺序选择最稳定的 gap tolerance。"""
        bounds: DataRegion = self._find_true_data_bounds(sheet)
        has_semantic_content, span_at = self._sheet_semantic_predicates(sheet)
        gap_tolerance, penalty, best_regions = select_best_gap_candidate(
            lambda tolerance: self._discover_sheet_regions(sheet, bounds, tolerance),
            has_semantic_content,
            span_at,
        )
        merged_lookup = self._get_merged_cell_lookup(sheet)
        tables = self._filter_semantic_subset_tables(
            [self._materialize_region(sheet, region, merged_lookup) for region in best_regions]
        )
        return gap_tolerance, penalty, tables

    def _select_best_tables(self, sheet: Worksheet) -> list[ExcelTable]:
        """选择并记录当前工作表的最佳表格候选集合。"""
        gap_tolerance, penalty, tables = self._select_best_gap_candidate(sheet)
        logger.debug(
            "Selected gap_tolerance={} for sheet '{}' with penalty={:.4f}",
            gap_tolerance,
            sheet.title,
            penalty,
        )
        return tables

    def _find_images_in_sheet(self, used_cells: set[tuple[int, int]] | None = None) -> None:
        """输出没有被表格吸收且不是公式载体的独立图片。"""
        if self.workbook is not None:
            for image in self.sheet_images:
                r, c = image.anchor
                if used_cells and r is not None and c is not None and (r, c) in used_cells:
                    continue

                if image.latex:
                    continue
                if image.image_base64:
                    self.cur_page.append(
                        {
                            "type": BlockType.IMAGE,
                            "image_base64": image.image_base64,
                        }
                    )

    def _find_data_tables(self, sheet: Worksheet) -> list[ExcelTable]:
        """在 Excel 工作表中查找所有紧凑的矩形数据表格。

        参数：
            sheet: 待解析的 Excel 工作表。

        返回：
            表示所有数据表格的 ExcelTable 对象列表。
        """
        if self.gap_tolerance is None:
            return self._select_best_tables(sheet)
        return self._find_data_tables_with_gap(sheet, self.gap_tolerance)

    def _find_data_tables_with_gap(self, sheet: Worksheet, gap_tolerance: int) -> list[ExcelTable]:
        """按固定 gap 发现表格并移除语义子集候选。"""
        return self._filter_semantic_subset_tables(self._find_data_tables_with_gap_raw(sheet, gap_tolerance))

    def _find_data_tables_with_gap_raw(self, sheet: Worksheet, gap_tolerance: int) -> list[ExcelTable]:
        """在固定 gap_tolerance 下查找工作表中的所有数据表格。"""
        bounds: DataRegion = self._find_true_data_bounds(sheet)  # 获取真实数据边界
        merged_lookup = self._get_merged_cell_lookup(sheet)
        return [
            self._materialize_region(sheet, region, merged_lookup)
            for region in self._discover_sheet_regions(sheet, bounds, gap_tolerance)
        ]

    def _discover_sheet_regions(
        self,
        sheet: Worksheet,
        bounds: DataRegion,
        gap_tolerance: int,
    ) -> list[ConnectedRegion]:
        """对当前工作表按 gap_tolerance 洪水填充发现连通数据区域。"""
        merged_lookup = self._get_merged_cell_lookup(sheet)
        max_row, max_col = bounds.max_row - 1, bounds.max_col - 1

        def has_content(row: int, col: int) -> bool:
            """检查指定单元格（0-based索引）是否有内容（有值或属于合并区域）。"""
            cell = sheet._cells.get((row + 1, col + 1))
            if cell is not None and cell.value is not None:
                return True
            return merged_lookup.contains_merged_cell(row, col)

        # 仅遍历已存在且有值的单元格，避免 iter_rows 在稀疏大表上创建大量空单元格。
        return discover_connected_regions(
            has_content,
            self._get_non_empty_cell_positions(sheet, bounds),
            max_row,
            max_col,
            gap_tolerance,
        )

    def _materialize_region(
        self,
        sheet: Worksheet,
        region: ConnectedRegion,
        merged_lookup: _MergedCellLookup,
    ) -> ExcelTable:
        """把连通区域包围盒物化为紧凑的表格 IR，正确处理合并单元格。

        遍历发现区域的边界框（bbox 内部的空格作为空单元格保留，维持矩形布局），
        跳过被合并单元格遮蔽的坐标并把跨度挂到合并锚点上。
        """
        data: list[ExcelCell] = []
        for source_row in range(region.row_start, region.row_end + 1):
            for source_col in range(region.col_start, region.col_end + 1):
                # 跳过被合并单元格遮蔽的单元格（非左上角）。
                if merged_lookup.is_hidden_merged_cell(source_row, source_col):
                    continue

                # 计算合并跨度（默认为 1x1）。
                row_span, col_span = merged_lookup.get_anchor_span(source_row, source_col)

                data.append(
                    self._build_excel_cell(
                        sheet,
                        source_row - region.row_start,
                        source_col - region.col_start,
                        source_row,
                        source_col,
                        row_span=row_span,
                        col_span=col_span,
                    )
                )

        return ExcelTable(
            anchor=(region.col_start, region.row_start),
            num_rows=region.num_rows,
            num_cols=region.num_cols,
            data=data,
        )

    def _get_non_empty_cell_positions(
        self,
        sheet: Worksheet,
        bounds: DataRegion,
    ) -> list[tuple[int, int]]:
        """按行列顺序返回真实边界内已有值单元格的 0-based 坐标。"""
        positions = []
        for cell in sheet._cells.values():
            if cell.value is None:
                continue
            if not (bounds.min_row <= cell.row <= bounds.max_row and bounds.min_col <= cell.column <= bounds.max_col):
                continue
            positions.append((cell.row - 1, cell.column - 1))
        return sorted(positions)

    def _find_true_data_bounds(self, sheet: Worksheet) -> DataRegion:
        """查找工作表中真实的数据边界（最小/最大行列）。

        该函数扫描所有单元格，找到包含所有非空单元格或合并单元格区域的
        最小矩形范围，返回边界的行列索引。

        参数：
            sheet: 待分析的工作表。

        返回：
            覆盖所有数据和合并单元格的最小矩形区域 DataRegion。
            若工作表为空，则默认返回 (1, 1, 1, 1)。
        """
        min_row, min_col = None, None
        max_row, max_col = 0, 0

        # 遍历所有有值的单元格，动态更新边界
        for cell in sheet._cells.values():
            if cell.value is not None:
                r, c = cell.row, cell.column
                min_row = r if min_row is None else min(min_row, r)
                min_col = c if min_col is None else min(min_col, c)
                max_row = max(max_row, r)
                max_col = max(max_col, c)

        # 将合并单元格的范围也纳入边界计算
        for merged in sheet.merged_cells.ranges:
            min_row = merged.min_row if min_row is None else min(min_row, merged.min_row)
            min_col = merged.min_col if min_col is None else min(min_col, merged.min_col)
            max_row = max(max_row, merged.max_row)
            max_col = max(max_col, merged.max_col)

        # 若工作表中没有任何数据，默认返回 (1, 1, 1, 1)
        if min_row is None or min_col is None:
            min_row = min_col = max_row = max_col = 1

        return DataRegion(min_row, max_row, min_col, max_col)

    def _get_merged_cell_lookup(self, sheet: Worksheet) -> _MergedCellLookup:
        """获取工作表合并单元格缓存，同一轮转换内每个 sheet 只构建一次。"""
        cache_key = id(sheet)
        lookup = self._merged_cell_lookup_cache.get(cache_key)
        if lookup is None:
            lookup = _MergedCellLookup(sheet)
            self._merged_cell_lookup_cache[cache_key] = lookup
        return lookup

    @staticmethod
    def _escape_text_with_line_breaks(text: str) -> str:
        """转义文本并把平台换行统一投影为 HTML 换行。"""
        return html.escape(text).replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")

    @staticmethod
    def _get_cell_hyperlink_target(cell: Any) -> str:
        """读取单元格外链或工作簿内 location。"""
        hyperlink = getattr(cell, "hyperlink", None)
        if not hyperlink:
            return ""

        target = getattr(hyperlink, "target", None)
        if target:
            return str(target)

        location = getattr(hyperlink, "location", None)
        if location:
            return f"#{location}"

        return ""

    @staticmethod
    def _apply_inline_font_tags(text_html: str, inline_font: Any) -> str:
        """按 openpyxl 行内字体顺序包装可见 HTML 标签。"""
        if not text_html or inline_font is None:
            return text_html

        wrapped = text_html
        if getattr(inline_font, "strike", False) or getattr(inline_font, "u", None):
            wrapped = wrapped.replace(" ", "&nbsp;")
        vert_align = getattr(inline_font, "vertAlign", None)
        if vert_align == "superscript":
            wrapped = f"<sup>{wrapped}</sup>"
        elif vert_align == "subscript":
            wrapped = f"<sub>{wrapped}</sub>"

        if getattr(inline_font, "strike", False):
            wrapped = f"<s>{wrapped}</s>"
        if getattr(inline_font, "u", None):
            wrapped = f"<u>{wrapped}</u>"
        if getattr(inline_font, "i", False):
            wrapped = f"<em>{wrapped}</em>"
        if getattr(inline_font, "b", False):
            wrapped = f"<strong>{wrapped}</strong>"

        return wrapped

    def _cell_value_to_html(self, cell: Any) -> tuple[str, bool]:
        """把普通或富文本单元格转换为安全 HTML 与内容类型标记。"""
        if cell.value is None:
            return "", False

        safe_target = sanitize_hyperlink_target(
            self._get_cell_hyperlink_target(cell),
            allowed_schemes=OFFICE_EXTERNAL_HYPERLINK_SCHEMES,
            allow_relative=True,
            allow_fragment=True,
        )
        link_target = html.escape(safe_target, quote=True) if safe_target else ""

        if isinstance(cell.value, CellRichText):
            html_parts = []
            for part in cell.value:
                if hasattr(part, "text"):
                    part_text = self._escape_text_with_line_breaks(str(getattr(part, "text", "")))
                    html_parts.append(
                        self._apply_inline_font_tags(
                            part_text,
                            getattr(part, "font", None),
                        )
                    )
                else:
                    html_parts.append(self._escape_text_with_line_breaks(str(part)))

            rich_text_html = "".join(html_parts)
            if link_target and rich_text_html:
                rich_text_html = f'<a href="{link_target}">{rich_text_html}</a>'
            return rich_text_html, True

        plain_text = str(cell.value)
        if link_target and plain_text:
            escaped_text = self._escape_text_with_line_breaks(plain_text)
            return f'<a href="{link_target}">{escaped_text}</a>', True

        return plain_text, False

    def _extract_cell_style(self, cell: Any) -> dict[str, Any]:
        """从 openpyxl 单元格提取当前 IR 保留的可见样式。"""
        style: dict[str, Any] = {}
        if cell.font:
            if cell.font.b:
                style["font-weight"] = "bold"
            if cell.font.i:
                style["font-style"] = "italic"
            if cell.font.u:
                style["text-decoration"] = "underline"
            if cell.font.strike:
                style["text-decoration"] = "line-through"
            if cell.font.color and hasattr(cell.font.color, "rgb") and cell.font.color.rgb:
                # Color might be ARGB "FF000000"
                color = cell.font.color.rgb
                if isinstance(color, str) and len(color) == 8:
                    style["color"] = "#" + color[2:]
                elif isinstance(color, str):
                    style["color"] = "#" + color

        if cell.alignment:
            if cell.alignment.horizontal:
                style["text-align"] = cell.alignment.horizontal
            if cell.alignment.vertical:
                style["vertical-align"] = cell.alignment.vertical

        # 渐变填充没有 patternType；只提取既有纯色背景，其他填充仍保留内容及其余样式。
        fill = cell.fill
        if getattr(fill, "patternType", None) == "solid" and fill.fgColor:
            color = fill.fgColor.rgb
            if hasattr(fill.fgColor, "type") and fill.fgColor.type == "rgb" and color:
                if isinstance(color, str) and len(color) == 8:
                    style["background-color"] = "#" + color[2:]
        return style


__all__ = ["SpreadsheetProjector"]
