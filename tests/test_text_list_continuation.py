"""正文列表直属末项续接的真实论文、保护条件及共享渲染回归。"""

from copy import deepcopy
import json
from pathlib import Path
from io import BytesIO
import zipfile

import pytest

from docvortex.content.inline import inline_plain_text
from docvortex.postprocess.paragraphs import merge_para_text_blocks
from docvortex.render import render, RenderFormat, RenderMode, PdfRenderOptions, PdfLayout
from docvortex.render._internal.common.planner import build_render_plan
from docvortex.render._internal.common.context import owned_render_document
from docvortex.schema import (
    MiddleJson,
    PageInfo,
    ListBlock,
    TextBlock,
    TextSpan,
    HyperlinkSpan,
    EquationInlineSpan,
    PageFootnoteBlock,
)


def _fixture() -> list[dict]:
    """读取修改前保存的真实论文行框，避免依赖模型再次生成版面。"""
    return json.loads((Path(__file__).parent / "fixtures/baige_text_list_continuation.json").read_text())


def test_paper_list_tail_and_column_chain() -> None:
    """第 1→2、8→9 页及第 9 页左右栏均建立续接，原页索引和直属成员不移动。"""
    pages = _fixture()
    original = deepcopy(pages)
    merge_para_text_blocks(pages)
    for page, indices in ((pages[1], [3]), (pages[3], [3, 4, 28])):
        for index in indices:
            assert next(block for block in page["blocks"] if block["index"] == index)["continues_prev"] is True
    for page, before in zip(pages, original):
        assert [(b["index"], b["type"], b.get("bbox")) for b in page["blocks"]] == [
            (b["index"], b["type"], b.get("bbox")) for b in before["blocks"]
        ]
        for block in page["blocks"]:
            if block["type"] == "list":
                assert all("continues_prev" not in child for child in block["content"])
    once = deepcopy(pages)
    merge_para_text_blocks(pages)
    assert pages == once


@pytest.mark.parametrize("case", ["reference", "nested", "geometry", "number", "bullet", "heading", "terminated", "page_gap"])
def test_reject_invalid_list_continuations(case: str) -> None:
    """只修改直属正文末项候选，所有已约定的屏障和新条目均拒绝合并。"""
    pages = _fixture()[:2]
    previous = next(b for b in pages[0]["blocks"] if b["type"] == "list")
    current = next(b for b in pages[1]["blocks"] if b["index"] == 3)
    tail = previous["content"][-1]
    if case == "reference":
        previous["sub_type"] = "ref_text"
    elif case == "nested":
        previous["content"][-1] = {"type": "list", "sub_type": "text", "content": [tail]}
    elif case == "geometry":
        tail.pop("bbox")
    elif case in ("number", "bullet"):
        current["content"][0]["content"] = ("3）" if case == "number" else "• ") + current["content"][0]["content"]
    elif case == "heading":
        pages[0]["blocks"].append({"type": "paragraph_title", "index": 100, "content": []})
    elif case == "terminated":
        tail["content"][-1]["content"] += "。"
    elif case == "page_gap":
        pages[1]["page_idx"] = 2
    merge_para_text_blocks(pages)
    assert current.get("continues_prev") is not True


def _document() -> MiddleJson:
    """建立列表、脚注、跨页片段和同页跨栏链，并保留公式、链接和样式。"""

    def text(index: int, content: str, *, continued: bool = False) -> TextBlock:
        """构造顶层普通正文或续文。"""
        return TextBlock(type="text", index=index, content=[TextSpan(type="text", content=content)], continues_prev=continued)

    item = text(1, "2）卫星技术，")
    item.continues_prev = None
    item.model_fields_set.discard("continues_prev")
    return MiddleJson(
        pages=[
            PageInfo(
                page_idx=0,
                blocks=[
                    text(0, "此前普通正文"),
                    ListBlock(type="list", index=1, sub_type="text", content=[item]),
                    PageFootnoteBlock(type="page_footnote", index=2, content=[TextSpan(type="text", content="收稿日期脚注")]),
                ],
            ),
            PageInfo(
                page_idx=1,
                blocks=[
                    TextBlock(
                        type="text",
                        index=0,
                        continues_prev=True,
                        content=[
                            TextSpan(type="text", content="覆盖", styles=["bold"]),
                            EquationInlineSpan(type="equation_inline", content="x^2"),
                            HyperlinkSpan(
                                type="hyperlink", url="https://example.com", content=[TextSpan(type="text", content="链接")]
                            ),
                        ],
                    ),
                    text(1, "范围广。", continued=True),
                ],
            ),
        ],
        metadata={"file_suffix": "html", "producer": {"name": "docvortex", "version": "0.5.13"}},
        is_full_document=True,
    )


@pytest.mark.parametrize("owned", [False, True])
@pytest.mark.parametrize("mode", [RenderMode.DEFAULT, RenderMode.FULL])
def test_planner_targets_list_tail_without_mutating_input(owned: bool, mode: RenderMode) -> None:
    """默认合并到列表末项，FULL 仅合并同页片段，所有权优化也不修改输入树。"""
    middle = _document()
    before = middle.model_dump()
    if owned:
        with owned_render_document(middle):
            plan = build_render_plan(middle, mode)
    else:
        plan = build_render_plan(middle, mode)
    assert plan[0][0].text_contents == [middle.pages[0].blocks[0].content]
    if mode is RenderMode.DEFAULT:
        tail = plan[0][1].block.content[-1]
        assert inline_plain_text(tail.content) == "2）卫星技术，覆盖x^2链接范围广。"
        assert tail.content[1].styles == ["bold"]
        assert any(isinstance(s, HyperlinkSpan) for s in tail.content)
        assert any(isinstance(s, EquationInlineSpan) for s in tail.content)
        assert plan[1][0].removed and plan[1][1].removed
    else:
        assert not plan[1][0].removed and plan[1][1].removed
        assert inline_plain_text(plan[0][1].block.content[-1].content) == "2）卫星技术，"
    assert middle.model_dump() == before


@pytest.mark.parametrize(
    "target", [RenderFormat.MARKDOWN, RenderFormat.HTML, RenderFormat.DOCX, RenderFormat.EPUB, RenderFormat.PDF]
)
def test_shared_renderers_consume_list_tail_plan(target: RenderFormat) -> None:
    """各共享格式消费同一合并目标，保留脚注及链接，语义 PDF 显式采用重排。"""
    from docx import Document
    from pypdf import PdfReader

    middle = _document()
    before = middle.model_dump()
    rendered = render(middle, target, options=PdfRenderOptions(layout=PdfLayout.REFLOW) if target is RenderFormat.PDF else None)
    if target is RenderFormat.DOCX:
        content = "\n".join(p.text for p in Document(BytesIO(rendered)).paragraphs)
    elif target is RenderFormat.EPUB:
        with zipfile.ZipFile(BytesIO(rendered)) as archive:
            content = "\n".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xhtml"))
    elif target is RenderFormat.PDF:
        content = "\n".join(page.extract_text() for page in PdfReader(BytesIO(rendered)).pages)
    else:
        content = rendered
    assert content.index("覆盖") < content.index("收稿日期脚注")
    assert content.count("范围广") == 1
    assert middle.model_dump() == before


@pytest.mark.parametrize("prefix", ["[3]", "［3］", "a)", "(3)", "一、", "③", "● ", "▫ ", "- "])
def test_new_item_markers_do_not_continue_tail(prefix: str) -> None:
    """数字、字母、中文编号和常见项目符号均保留为新条目。"""
    pages = _fixture()[:2]
    current = next(b for b in pages[1]["blocks"] if b["index"] == 3)
    current["content"][0]["content"] = prefix + current["content"][0]["content"]
    merge_para_text_blocks(pages)
    assert current.get("continues_prev") is not True


@pytest.mark.parametrize("barrier", ["terminated", "geometry"])
def test_list_tail_does_not_bypass_guards_with_reference_hint(barrier: str) -> None:
    """正文列表不能借助参考条目侧车绕过句末终止或缺失几何的保护。"""
    pages = _fixture()[:2]
    previous = next(b for b in pages[0]["blocks"] if b["type"] == "list")["content"][-1]
    current = next(b for b in pages[1]["blocks"] if b["index"] == 3)
    previous["_reference_start"] = current["_reference_start"] = False
    if barrier == "terminated":
        previous["content"][-1]["content"] += "。"
    else:
        previous.pop("bbox")
    merge_para_text_blocks(pages)
    assert current.get("continues_prev") is not True
