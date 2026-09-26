"""验证职责拆分后转换器状态、类型入口及失败重试的隔离性。"""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO, UnsupportedOperation
from unittest.mock import Mock
from zipfile import ZipFile

import pytest

from docvortex.analyzers.native.office.errors import LegacyOfficeResourceLimitError


class _NonSeekableStream(BytesIO):
    """模拟仅支持顺序读取的输入，覆盖不可复位流分支。"""

    def seek(self, *args, **kwargs):
        """拒绝复位，让转换器使用一次性读取路径。"""
        raise UnsupportedOperation("not seekable")


@pytest.mark.parametrize("seekable", [False, True])
@pytest.mark.parametrize("error_type", [LegacyOfficeResourceLimitError, MemoryError])
def test_pptx_resource_errors_do_not_normalize(
    monkeypatch: pytest.MonkeyPatch, seekable: bool, error_type: type[Exception]
) -> None:
    """验证两种输入流的资源异常原样传播，且不会复制整包或触发规范化。"""
    from docvortex.analyzers.native.office.pptx import pptx_converter as module

    converter = module.PptxConverter()
    error = error_type("resource exhausted")
    convert_stream = Mock(side_effect=error)
    normalize = Mock(side_effect=AssertionError("资源异常不应触发规范化"))
    read_again = Mock(side_effect=AssertionError("资源异常不应重新读取整包"))
    monkeypatch.setattr(converter, "_convert_package_stream", convert_stream)
    monkeypatch.setattr(module, "normalize_pptx_package", normalize)
    monkeypatch.setattr(module, "read_stream_bytes_from_start", read_again)
    stream = (BytesIO if seekable else _NonSeekableStream)(b"package")

    with pytest.raises(error_type) as caught:
        converter.convert(stream)
    assert caught.value is error
    convert_stream.assert_called_once()
    normalize.assert_not_called()
    read_again.assert_not_called()
    assert not stream.closed


@pytest.mark.parametrize("seekable", [False, True])
def test_pptx_strict_package_still_retries_after_normalization(
    monkeypatch: pytest.MonkeyPatch, seekable: bool
) -> None:
    """验证真正的 Strict OOXML 包兼容问题仍能规范化重试并完整输出。"""
    from pptx import Presentation
    from docvortex.analyzers.native.office.pptx import pptx_converter as module

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0]).shapes.title.text = "Strict presentation"
    original = BytesIO()
    presentation.save(original)
    strict = BytesIO()
    with ZipFile(original) as source, ZipFile(strict, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.endswith(".xml"):
                data = data.replace(
                    b"http://schemas.openxmlformats.org/presentationml/2006/main",
                    b"http://purl.oclc.org/ooxml/presentationml/main",
                )
            target.writestr(info, data)
    expected = module.PptxConverter()
    expected.convert(original)
    normalize = Mock(wraps=module.normalize_pptx_package)
    monkeypatch.setattr(module, "normalize_pptx_package", normalize)
    converter = module.PptxConverter()
    stream = (BytesIO if seekable else _NonSeekableStream)(strict.getvalue())

    converter.convert(stream)

    normalize.assert_called_once()
    assert converter.pages == expected.pages
    assert len(converter.pages) == 1
    assert not stream.closed


@pytest.mark.parametrize("suffix", ["docx", "pptx"])
def test_converter_reuse_after_failure_matches_fresh_instance(suffix: str) -> None:
    """不同文档及失败重试使用独立状态，不修改之前交付的页面列表。"""
    if suffix == "docx":
        from docx import Document

        from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter as Converter

        def payload(text: str) -> bytes:
            """生成包含单段正文的独立 DOCX。"""
            document = Document()
            document.add_paragraph(text)
            stream = BytesIO()
            document.save(stream)
            return stream.getvalue()
    else:
        from pptx import Presentation

        from docvortex.analyzers.native.office.pptx.pptx_converter import PptxConverter as Converter

        def payload(text: str) -> bytes:
            """生成包含一个标题页的独立 PPTX。"""
            document = Presentation()
            slide = document.slides.add_slide(document.slide_layouts[0])
            slide.shapes.title.text = text
            stream = BytesIO()
            document.save(stream)
            return stream.getvalue()

    converter = Converter()
    converter.convert(BytesIO(payload("first document")))
    first_pages = converter.pages
    first_snapshot = deepcopy(first_pages)
    with pytest.raises(Exception):
        converter.convert(BytesIO(b"invalid package"))
    second = payload("second document")
    converter.convert(BytesIO(second))
    fresh = Converter()
    fresh.convert(BytesIO(second))
    assert converter.pages == fresh.pages
    assert converter.pages is not first_pages
    assert first_pages == first_snapshot


def test_pdf_contract_identity_survives_module_split() -> None:
    """原 PDF 模块的类型仍是唯一对象，已存在的 pickle 数据可以继续读取。"""
    import pickle

    from docvortex.document.pdf import _document as document
    from docvortex.document.pdf import native_contracts

    assert document.PDFDrawingLine is native_contracts.PDFDrawingLine
    value = document.PDFDrawingLine((0, 0), (10, 0), (0, 0, 10, 0), 1, "horizontal")
    assert pickle.loads(pickle.dumps(value)) == value
