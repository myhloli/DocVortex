"""以冻结 Python 算法验证电子表格 Rust 区域、评分、筛选和真实入口。"""

from __future__ import annotations

from dataclasses import asdict, dataclass
import hashlib
from io import BytesIO
from pathlib import Path
import random
from zipfile import ZipFile

import pytest

import _spreadsheet_regions_oracle as oracle
from docvortex._compute_backend import get_native


@dataclass(frozen=True)
class _Case:
    """独立描述起点、连接位置、语义原点和合并跨度，不预先合并四类信息。"""

    name: str
    starts: tuple[tuple[int, int], ...] = ()
    points: tuple[tuple[int, int], ...] = ()
    rectangles: tuple[tuple[int, int, int, int], ...] = ()
    semantic: tuple[tuple[int, int], ...] = ()
    spans: tuple[tuple[int, int, int, int], ...] = ()
    max_row: int = 5
    max_col: int = 5


def _callbacks(case: _Case):
    """只把输入事实适配成旧算法的三个回调，不自行实现连接或评分逻辑。"""
    points = set(case.points)
    semantic = set(case.semantic)
    spans = {(row, col): (row_span, col_span) for row, col, row_span, col_span in case.spans}

    def has_content(row: int, col: int) -> bool:
        """查询离散连接格或闭区间矩形，不把语义位置自动加入连接集合。"""
        return (row, col) in points or any(
            row_start <= row <= row_end and col_start <= col <= col_end
            for row_start, row_end, col_start, col_end in case.rectangles
        )

    def has_semantic_content(row: int, col: int) -> bool:
        """独立查询真实语义原点，合并覆盖格不会自行成为语义原点。"""
        return (row, col) in semantic

    def span_at(row: int, col: int) -> tuple[int, int]:
        """返回包括无语义合并原点在内的跨度，其余坐标保持默认单格。"""
        return spans.get((row, col), (1, 1))

    return has_content, has_semantic_content, span_at


def _oracle_regions(regions) -> list[tuple[list[tuple[int, int]], tuple[int, int, int, int]]]:
    """把冻结算法结果变成可比较证据，仅规范无序 cells，不重排区域。"""
    return [(sorted(region.cells), (region.row_start, region.row_end, region.col_start, region.col_end)) for region in regions]


def _native_regions(regions) -> list[tuple[list[tuple[int, int]], tuple[int, int, int, int]]]:
    """保留 Rust 区域顺序并拒绝重复 cells，仅消除集合内部无意义的次序差异。"""
    result = []
    for cells, bounds in regions:
        normalized = [tuple(point) for point in cells]
        assert len(normalized) == len(set(normalized))
        assert len(bounds) == 4
        result.append((sorted(normalized), tuple(bounds)))
    return result


def _oracle_trace(case: _Case) -> tuple:
    """独立运行冻结的 BFS、逐候选汇总、选择和严格子集筛选，返回各阶段证据。"""
    has_content, has_semantic_content, span_at = _callbacks(case)
    by_gap = {}
    candidates = []
    for gap in oracle.AUTO_GAP_TOLERANCE_CANDIDATES:
        regions = oracle.discover_connected_regions(has_content, case.starts, case.max_row, case.max_col, gap)
        by_gap[gap] = regions
        summary = oracle.summarize_connected_regions(regions, has_semantic_content, span_at)
        candidates.append((gap, oracle.gap_candidate_penalty(summary), _oracle_regions(regions), summary))

    def discover(gap: int):
        """复用本次旧算法已生成的候选，选择过程仍执行原始评分和阈值规则。"""
        return by_gap[gap]

    gap, penalty, regions = oracle.select_best_gap_candidate(discover, has_semantic_content, span_at)
    semantic_sets = [oracle.region_semantic_positions(region, has_semantic_content) for region in regions]
    kept = oracle.keep_maximal_by_semantic_sets(semantic_sets)
    return gap, penalty, _oracle_regions(regions), kept, candidates


def _assert_float_exact(actual: float, expected: float) -> None:
    """比较 Python 浮点的精确十六进制表示，禁止容差掩盖 0.15 选择边界差异。"""
    assert type(actual) is float
    assert actual.hex() == expected.hex()


def _assert_trace_exact(actual: tuple, expected: tuple) -> None:
    """逐项检查最终选择及三档候选，避免只比较最终 bbox 漏掉中间语义变化。"""
    gap, penalty, regions, kept, candidates = actual
    expected_gap, expected_penalty, expected_regions, expected_kept, expected_candidates = expected
    assert gap == expected_gap
    _assert_float_exact(penalty, expected_penalty)
    assert _native_regions(regions) == expected_regions
    assert list(kept) == expected_kept
    assert len(candidates) == len(expected_candidates) == 3
    for candidate, expected_candidate in zip(candidates, expected_candidates, strict=True):
        current_gap, current_penalty, current_regions, summary = candidate
        reference_gap, reference_penalty, reference_regions, reference_summary = expected_candidate
        assert current_gap == reference_gap
        _assert_float_exact(current_penalty, reference_penalty)
        assert _native_regions(current_regions) == reference_regions
        assert summary.keys() == reference_summary.keys()
        for field, value in reference_summary.items():
            if isinstance(value, float):
                _assert_float_exact(summary[field], value)
            else:
                assert type(summary[field]) is int
                assert summary[field] == value


def _named_cases() -> list[_Case]:
    """构造针对既有契约的反例，刻意分离起点、连接、语义及跨度。"""
    dense = tuple((row, col) for row in range(4) for col in range(5))
    return [
        _Case("empty-domain", max_row=-1, max_col=-1),
        _Case("empty-starts-have-unvisited-content", points=((1, 1),), semantic=((1, 1),)),
        _Case("unconditional-noncontent-seed", starts=((2, 2),)),
        _Case(
            "duplicate-input-facts-do-not-double-count",
            starts=((0, 0), (0, 0)),
            points=((0, 0), (0, 0), (0, 1)),
            rectangles=((0, 0, 0, 1), (0, 0, 0, 1)),
            semantic=((0, 0), (0, 0)),
            max_row=0,
            max_col=1,
        ),
        _Case("empty-semantic-merged-singleton", starts=((2, 2),), points=((2, 2),), spans=((2, 2, 3, 2),)),
        _Case(
            "two-ghost-seeds-share-global-visited-content",
            starts=((0, 0), (0, 2)),
            points=((0, 1),),
            semantic=((0, 1),),
            max_row=0,
            max_col=2,
        ),
        _Case(
            "bbox-semantic-point-is-not-a-connected-cell",
            starts=((0, 0),),
            points=((0, 0), (0, 1), (1, 1)),
            semantic=((0, 0), (1, 0)),
            max_row=1,
            max_col=1,
        ),
        _Case(
            "merge-only-island-does-not-seed",
            starts=((0, 0),),
            points=((0, 0),),
            rectangles=((4, 5, 4, 5),),
            semantic=((0, 0),),
            spans=((4, 4, 2, 2),),
        ),
        _Case(
            "preserve-unsorted-and-duplicate-seeds",
            starts=((4, 4), (0, 0), (4, 4), (4, 5)),
            points=((0, 0), (4, 4), (4, 5)),
            semantic=((0, 0), (4, 4)),
        ),
        _Case(
            "overlapping-and-clipped-semantic-spans",
            starts=((0, 0),),
            rectangles=((0, 2, 0, 3), (1, 2, 1, 3)),
            semantic=((0, 0), (1, 1), (2, 3)),
            spans=((0, 0, 3, 3), (1, 1, 4, 4), (2, 3, 2, 3)),
            max_row=2,
            max_col=3,
        ),
        _Case(
            "side-by-side-row-overlap",
            starts=((0, 0), (0, 5)),
            rectangles=((0, 3, 0, 0), (0, 3, 5, 5)),
            semantic=((0, 0), (0, 5), (3, 0), (3, 5)),
            max_row=3,
            max_col=5,
        ),
        _Case(
            "two-interior-separators",
            starts=((0, 0), (3, 0)),
            points=((0, 0), (3, 0)),
            semantic=((0, 0), (3, 0)),
            max_row=3,
            max_col=0,
        ),
        _Case("sparse-ratio-exact-035", starts=((0, 0),), points=dense, semantic=dense[:13], max_row=3, max_col=4),
        _Case("sparse-ratio-over-035", starts=((0, 0),), points=dense, semantic=dense[:12], max_row=3, max_col=4),
        _Case(
            "span-does-not-expand-region-bounds", starts=((1, 1),), points=((1, 1),), semantic=((1, 1),), spans=((1, 1, 5, 5),)
        ),
        _Case(
            "negative-seed-can-reach-bounded-content",
            starts=((-1, 0),),
            points=((0, 0), (1, 0)),
            semantic=((-1, 0), (1, 0)),
            max_row=1,
            max_col=0,
        ),
        _Case("positive-outside-seed-stays-present", starts=((8, 9),), semantic=((8, 9),), max_row=2, max_col=2),
        _Case("nonempty-start-with-empty-bounds", starts=((0, 0),), semantic=((0, 0),), max_row=-1, max_col=-1),
        _Case(
            "large-coordinate-small-bbox",
            starts=((2**52, 2**52),),
            rectangles=((2**52, 2**52 + 2, 2**52, 2**52 + 3),),
            semantic=((2**52, 2**52), (2**52 + 2, 2**52 + 3)),
            spans=((2**52, 2**52, 2, 2),),
            max_row=2**52 + 2,
            max_col=2**52 + 3,
        ),
    ]


def _random_cases() -> list[_Case]:
    """用固定种子独立生成稀疏、密集、合并和无语义形态，规模限制在小网格内。"""
    rng = random.Random(20261025)
    cases = []
    for index in range(96):
        rows, cols = rng.randint(1, 9), rng.randint(1, 10)
        row_offset, col_offset = rng.choice((0, 0, 7)), rng.choice((0, 0, 11))
        coordinates = [(row_offset + row, col_offset + col) for row in range(rows) for col in range(cols)]
        density = (0.08, 0.25, 0.65, 0.95)[index % 4]
        points = [point for point in coordinates if rng.random() < density]
        rectangles = []
        spans = {}
        for _ in range(rng.randint(0, 3)):
            row, col = rng.choice(coordinates)
            row_end = min(row + rng.randint(0, 3), row_offset + rows - 1)
            col_end = min(col + rng.randint(0, 3), col_offset + cols - 1)
            rectangles.append((row, row_end, col, col_end))
            spans[(row, col)] = (row_end - row + 1, col_end - col + 1)
        semantic = [point for point in coordinates if rng.random() < density * 0.7]
        if index % 7 == 0:
            semantic = []
        for point in rng.sample(coordinates, min(len(coordinates), rng.randint(0, 3))):
            spans[point] = (rng.randint(1, 5), rng.randint(1, 5))
        starts = [point for point in points if rng.random() < 0.65]
        if index % 3 == 0:
            starts.append(rng.choice(coordinates))
        if index % 5 == 0:
            rng.shuffle(starts)
        if starts and index % 4 == 0:
            starts.insert(rng.randrange(len(starts) + 1), rng.choice(starts))
        cases.append(
            _Case(
                f"seed-20261025-{index:02d}",
                starts=tuple(starts),
                points=tuple(points),
                rectangles=tuple(rectangles),
                semantic=tuple(semantic),
                spans=tuple((row, col, row_span, col_span) for (row, col), (row_span, col_span) in spans.items()),
                max_row=row_offset + rows - 1,
                max_col=col_offset + cols - 1,
            )
        )
    return cases


_CASES = _named_cases() + _random_cases()


def test_frozen_python_oracle_body_keeps_its_original_source_digest() -> None:
    """防止随 Rust 实现同步改写 oracle，从而让相同错误在两侧互相掩盖。"""
    lines = Path(oracle.__file__).read_bytes().splitlines(keepends=True)
    assert (
        lines[2].decode().strip() == "# Frozen-source-sha256: 6ac2cc9e797226cfde1d8b4839f41c349cfc357b2b9c18a45dbaf9247172a421"
    )
    assert hashlib.sha256(b"".join(lines[4:])).hexdigest() == "6ac2cc9e797226cfde1d8b4839f41c349cfc357b2b9c18a45dbaf9247172a421"


@pytest.mark.parametrize("case", _CASES, ids=[case.name for case in _CASES])
def test_native_candidate_trace_matches_frozen_python_exactly(case: _Case) -> None:
    """逐阶段对照每档连接集合、bbox、全部评分指标、精确浮点及最终筛选。"""
    actual = get_native().spreadsheet_select_regions(
        case.starts, case.points, case.rectangles, case.semantic, case.spans, case.max_row, case.max_col, diagnostics=True
    )
    _assert_trace_exact(actual, _oracle_trace(case))


@pytest.mark.parametrize("case", _CASES, ids=[case.name for case in _CASES])
@pytest.mark.parametrize("gap", [-3, -1, 0, 1, 2, 5])
def test_native_fixed_gap_preserves_seed_order_cells_and_bounds(case: _Case, gap: int) -> None:
    """固定 gap 独立核对 BFS，包含负 gap、重复起点、越界起点和非内容起点。"""
    has_content, _, _ = _callbacks(case)
    expected = oracle.discover_connected_regions(has_content, case.starts, case.max_row, case.max_col, gap)
    actual = get_native().spreadsheet_discover_regions(
        case.starts, case.points, case.rectangles, case.max_row, case.max_col, gap
    )
    assert _native_regions(actual) == _oracle_regions(expected)


@pytest.mark.parametrize(
    "semantic_sets",
    [
        [],
        [[], []],
        [[(1, 1)], [(1, 1)]],
        [[], [(1, 1)], [(2, 2)], [(1, 1), (2, 2)]],
        [[(0, 0), (1, 1)], [(1, 1)], [(0, 0), (1, 1)]],
        [[(0, 0), (0, 0)], [(2, 2)], []],
    ],
)
def test_native_strict_semantic_subset_keeps_equal_sets_and_order(semantic_sets) -> None:
    """空集合与相等集合只按严格包含判定，重复坐标不改变语义集合及保留顺序。"""
    expected = oracle.keep_maximal_by_semantic_sets([set(points) for points in semantic_sets])
    assert get_native().spreadsheet_keep_maximal_semantic_sets(semantic_sets) == expected


def test_native_strict_semantic_subset_matches_fixed_seed_set_families() -> None:
    """独立生成有重复、包含与交叉的集合族，避免筛选测试只覆盖嵌套链。"""
    rng = random.Random(2026102501)
    coordinates = [(row, col) for row in range(4) for col in range(5)]
    for _ in range(60):
        families = [rng.sample(coordinates, rng.randint(0, 12)) for _ in range(rng.randint(0, 12))]
        if families:
            families.append(list(rng.choice(families)))
        expected = oracle.keep_maximal_by_semantic_sets([set(points) for points in families])
        assert get_native().spreadsheet_keep_maximal_semantic_sets(families) == expected


def test_native_large_sparse_bbox_scores_without_materializing_dense_mask() -> None:
    """用一万余连接格包围六千余万槽位，按解析公式检查评分而不运行密集 Python oracle。"""
    size = 8192
    points = [(row, row) for row in range(size)] + [(row, row + 1) for row in range(size - 1)]
    semantic = [(0, 0), (size - 1, size - 1)]
    result = get_native().spreadsheet_select_regions([(0, 0)], points, [], semantic, [], size - 1, size - 1, diagnostics=True)
    gap, penalty, regions, kept, candidates = result
    blank_ratio = 1.0 - 2 / (size * size)
    expected_summary = {
        "real_singleton_ratio": 0.0,
        "weighted_blank_ratio": (size * size * blank_ratio) / (size * size),
        "interior_blank_line_ratio": 1.0,
        "sparse_large_table_ratio": 1.0,
        "severe_separator_count": 1,
        "row_overlap_excess_ratio": 0.0,
    }
    expected_penalty = oracle.gap_candidate_penalty(expected_summary)
    assert gap == 1
    assert list(kept) == [0]
    _assert_float_exact(penalty, expected_penalty)
    assert _native_regions(regions) == [(sorted(points), (0, size - 1, 0, size - 1))]
    assert len(candidates) == 3
    for (candidate_gap, candidate_penalty, candidate_regions, summary), expected_gap in zip(candidates, (0, 1, 2), strict=True):
        assert candidate_gap == expected_gap
        assert summary == expected_summary
        for field, value in expected_summary.items():
            if isinstance(value, float):
                _assert_float_exact(summary[field], value)
        _assert_float_exact(candidate_penalty, expected_penalty)
        assert _native_regions(candidate_regions) == _native_regions(regions)


def _xlsx_source() -> bytes:
    """生成包含空白值、布尔值、离散表和合并单元格的真实 XLSX 小包。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "left"
    sheet["A2"] = "   "
    sheet["A3"] = 0
    sheet["B1"] = False
    sheet["D1"] = "right"
    sheet["D2"] = "merged"
    sheet.merge_cells("D2:E3")
    sheet["G5"] = "separate"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _ods_source() -> bytes:
    """生成具有重复、合并和内部空白格的最小 ODS 包，使用真实 ODF 入口。"""
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        '<office:body><office:spreadsheet><table:table table:name="Sheet">'
        "<table:table-row><table:table-cell><text:p>left</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="2"/>'
        '<table:table-cell table:number-columns-spanned="2" table:number-rows-spanned="2">'
        "<text:p>merged</text:p></table:table-cell><table:covered-table-cell/></table:table-row>"
        "<table:table-row><table:table-cell><text:p>second</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="2"/>'
        '<table:covered-table-cell table:number-columns-repeated="2"/></table:table-row>'
        '<table:table-row table:number-rows-repeated="2"><table:table-cell/></table:table-row>'
        "<table:table-row><table:table-cell><text:p>separate</text:p></table:table-cell></table:table-row>"
        "</table:table></office:spreadsheet></office:body></office:document-content>"
    )
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", content)
    return output.getvalue()


def _full_public_output(result) -> tuple:
    """保留所有模型字段、原始素材字节和诊断，端到端对照不忽略任何结果字段。"""
    return (
        result.model_json.model_dump(mode="json"),
        result.middle_json.model_dump(mode="json"),
        dict(result.assets),
        [asdict(item) for item in result.diagnostics],
    )


@pytest.mark.parametrize("suffix", ["xlsx", "ods"])
def test_public_nonpdf_entry_uses_native_and_matches_original_callbacks(monkeypatch: pytest.MonkeyPatch, suffix: str) -> None:
    """验证真实非 PDF parse 命中新扩展，再与冻结 Python 回调路径比较完整公开输出。"""
    from docvortex import parse
    from docvortex.analyzers.native.office.odf import table as odf_table
    from docvortex.analyzers.native.office.spreadsheet import projector

    payload = _xlsx_source() if suffix == "xlsx" else _ods_source()
    native = get_native()
    original = native.spreadsheet_select_regions
    observed = []

    def tracked_select(*args, **kwargs):
        """记录真实 Rust 选择调用并返回其原结果，不以模拟结果替代原生计算。"""
        observed.append((args, kwargs))
        return original(*args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(native, "spreadsheet_select_regions", tracked_select)
        actual = parse(payload, file_suffix=suffix, keep_model_json=True)
    assert observed

    def unexpected_native(*_args, **_kwargs):
        """强制冻结回调路径不得悄悄回到同一 Rust 实现而形成自比较。"""
        pytest.fail("frozen callback reference reached the native spreadsheet implementation")

    with monkeypatch.context() as patch:
        for module in (projector, odf_table):
            patch.setattr(module, "discover_connected_regions", oracle.discover_connected_regions)
            patch.setattr(module, "select_best_gap_candidate", oracle.select_best_gap_candidate)
            patch.setattr(module, "keep_maximal_by_semantic_sets", oracle.keep_maximal_by_semantic_sets)
        for name in ("spreadsheet_select_regions", "spreadsheet_discover_regions", "spreadsheet_keep_maximal_semantic_sets"):
            patch.setattr(native, name, unexpected_native)
        expected = parse(payload, file_suffix=suffix, keep_model_json=True)
    assert _full_public_output(actual) == _full_public_output(expected)
