# Copyright (c) Opendatalab. All rights reserved.
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pytest

from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter

NO_BREAK_HYPHEN = "‑"


SOFT_HYPHEN = "­"


def _save_docx_bytes(doc: Document) -> bytes:
    """Serializes an in-memory DOCX document into bytes."""
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _convert_docx_bytes(file_bytes: bytes) -> list[dict]:
    """Calls the current Flash DOCX Converter and flattens the paging block."""
    converter = DocxConverter()
    converter.convert(BytesIO(file_bytes))
    return [block for page in converter.pages for block in page]


def _table_blocks(blocks: list[dict]) -> list[dict]:
    """Filter table blocks from the list of flattened blocks."""
    return [block for block in blocks if block["type"] == "table"]


def _build_docx_with_no_break_hyphen_table() -> bytes:
    """Construct an entire table loss use case with both non-breaking hyphens and numbered lists."""
    doc = Document()
    doc.add_heading("1 项目概述", level=1)

    table = doc.add_table(rows=2, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run("NCI")
    etree.SubElement(run._r, qn("w:noBreakHyphen"))
    paragraph.add_run("CTCAE 标准")
    numbered = table.cell(0, 1).paragraphs[0]
    numbered.style = doc.styles["List Number"]
    numbered.add_run("第一条编号内容")

    doc.add_heading("2 进度安排", level=1)
    plain = doc.add_table(rows=1, cols=2)
    plain.cell(0, 0).paragraphs[0].add_run("里程碑")
    plain.cell(0, 1).paragraphs[0].add_run("2026-08-01")
    return _save_docx_bytes(doc)


def _build_docx_with_sym_table(*, char: str) -> bytes:
    """Construct a table containing both the Symbol character and the numbered list."""
    doc = Document()
    doc.add_heading("符号表格", level=1)
    table = doc.add_table(rows=1, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run("符号")
    etree.SubElement(
        run._r,
        qn("w:sym"),
        {qn("w:char"): char, qn("w:font"): "Symbol"},
    )
    paragraph.add_run("结束")
    numbered = table.cell(0, 1).paragraphs[0]
    numbered.style = doc.styles["List Number"]
    numbered.add_run("编号项")
    return _save_docx_bytes(doc)


def _build_plain_docx_table() -> bytes:
    """Constructs a normal table control group without special characters."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("APTT")
    table.cell(0, 1).paragraphs[0].add_run("部分凝血酶时间")
    return _save_docx_bytes(doc)


def _build_nested_merged_docx_table() -> bytes:
    """Construct a complete contextual table with nested rows, vertical merging, and numbering styles."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(1, 0))
    merged.paragraphs[0].add_run("纵向合并")
    nested = table.cell(0, 1).add_table(rows=2, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run("嵌套首行")
    numbered = nested.cell(1, 0).paragraphs[0]
    numbered.style = doc.styles["List Number"]
    numbered.add_run("嵌套编号行")
    table.cell(1, 1).paragraphs[0].add_run("外层尾行")
    return _save_docx_bytes(doc)


def _build_docx_with_residual_vmerge_text(*, explicit_continue: bool) -> bytes:
    """Construct a table with continuation cells containing residual text and adjacent cells containing numbered lists."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(1, 0))
    merged.paragraphs[0].add_run("职责")
    continuation = table._tbl.tr_lst[1].tc_lst[0]
    vmerge = continuation.tcPr.vMerge
    if explicit_continue:
        vmerge.set(qn("w:val"), "continue")
    else:
        vmerge.attrib.pop(qn("w:val"), None)
    residual_run = etree.SubElement(continuation.p_lst[0], qn("w:r"))
    etree.SubElement(residual_run, qn("w:t")).text = "残留"
    numbered = table.cell(0, 1).paragraphs[0]
    numbered.style = doc.styles["List Number"]
    numbered.add_run("编号内容")
    table.cell(1, 1).paragraphs[0].add_run("尾部内容")
    return _save_docx_bytes(doc)


def test_no_break_hyphen_table_not_lost() -> None:
    """Verify that the entire table is not lost when a non-breaking hyphen coexists with a numbered list."""
    tables = _table_blocks(_convert_docx_bytes(_build_docx_with_no_break_hyphen_table()))

    assert len(tables) == 2
    assert f"NCI{NO_BREAK_HYPHEN}CTCAE" in tables[0]["content"]
    assert "<li>" in tables[0]["content"]
    assert "第一条编号内容" in tables[0]["content"]
    assert "里程碑" in tables[1]["content"]


def test_mapped_sym_table_takes_mammoth_path() -> None:
    """Verify that the common Symbol map characters are consistent with the Mammoth table signature."""
    tables = _table_blocks(_convert_docx_bytes(_build_docx_with_sym_table(char="0022")))

    assert len(tables) == 1
    assert "∀" in tables[0]["content"]
    assert "<li>" in tables[0]["content"]


def test_f0_sym_table_takes_low_byte_fallback() -> None:
    """Verify that F0 prefix Symbol characters are mapped using low byte fallback according to Mammoth rules."""
    tables = _table_blocks(_convert_docx_bytes(_build_docx_with_sym_table(char="F0B7")))

    assert len(tables) == 1
    assert "•" in tables[0]["content"]
    assert "<li>" in tables[0]["content"]


def test_plain_table_still_emitted() -> None:
    """Verify special character signature fix does not affect normal forms."""
    tables = _table_blocks(_convert_docx_bytes(_build_plain_docx_table()))

    assert len(tables) == 1
    assert "APTT" in tables[0]["content"]


def test_nested_merged_table_uses_recursive_row_signature_and_full_context() -> None:
    """Verify that XML/HTML recursive line number alignment does not enter an orphan fallback with missing numbering context."""
    file_bytes = _build_nested_merged_docx_table()
    converter = DocxConverter()
    preparsed = converter._preparse_tables_with_mammoth(file_bytes)

    assert len(preparsed) == 1
    assert preparsed[0] is not None
    tables = _table_blocks(_convert_docx_bytes(file_bytes))
    assert len(tables) == 1
    assert "纵向合并" in tables[0]["content"]
    assert "嵌套首行" in tables[0]["content"]
    assert "嵌套编号行" in tables[0]["content"]
    assert "外层尾行" in tables[0]["content"]


@pytest.mark.parametrize("explicit_continue", (True, False))
def test_vmerge_continuation_text_is_ignored_by_table_signature(explicit_continue: bool) -> None:
    """Verify that both OOXML and continuation writing methods are consistent with the Mammoth form text signature."""
    file_bytes = _build_docx_with_residual_vmerge_text(explicit_continue=explicit_continue)
    converter = DocxConverter()
    preparsed = converter._preparse_tables_with_mammoth(file_bytes)

    assert len(preparsed) == 1
    assert preparsed[0] is not None
    assert "编号内容" in preparsed[0]
    assert "残留" not in preparsed[0]
    tables = _table_blocks(_convert_docx_bytes(file_bytes))
    assert len(tables) == 1
    assert "rowspan=\"2\"" in tables[0]["content"]
    assert "编号内容" in tables[0]["content"]


def test_xml_table_signature_renders_special_chars() -> None:
    """Verify that the XML signature preserves non-breaking hyphens and soft hyphens."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run("NCI")
    etree.SubElement(run._r, qn("w:noBreakHyphen"))
    paragraph.add_run("CTCAE")
    second_paragraph = table.cell(0, 1).paragraphs[0]
    second_run = second_paragraph.add_run("软连")
    etree.SubElement(second_run._r, qn("w:softHyphen"))
    second_paragraph.add_run("字符")

    signature = DocxConverter._xml_table_signature(table._tbl)

    assert f"NCI{NO_BREAK_HYPHEN}CTCAE" in signature["text"]
    assert f"软连{SOFT_HYPHEN}字符" in signature["text"]


def test_xml_table_signature_renders_sym_like_mammoth() -> None:
    """Verified Symbol output character is mapped, true unmapped Symbol output empty string."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run("映射")
    etree.SubElement(
        run._r,
        qn("w:sym"),
        {qn("w:char"): "0022", qn("w:font"): "Symbol"},
    )
    paragraph.add_run("结束")
    second_paragraph = table.cell(0, 1).paragraphs[0]
    second_run = second_paragraph.add_run("未映射")
    etree.SubElement(
        second_run._r,
        qn("w:sym"),
        {qn("w:char"): "F0B7", qn("w:font"): "UnknownFont"},
    )
    second_paragraph.add_run("结束")

    signature = DocxConverter._xml_table_signature(table._tbl)

    assert signature["text"] == "映射∀结束未映射结束"


def test_xml_table_signature_excludes_omml_equation_text() -> None:
    """Verify that Mammoth does not render the OMML formula text that does not participate in the table signature."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run("正文")
    equation = etree.SubElement(run._r, qn("m:oMath"))
    equation_run = etree.SubElement(equation, qn("m:r"))
    etree.SubElement(equation_run, qn("m:t")).text = "E=mc2"
    table.cell(0, 1).paragraphs[0].add_run("单元格")

    signature = DocxConverter._xml_table_signature(table._tbl)

    assert "正文" in signature["text"]
    assert "E=mc2" not in signature["text"]
