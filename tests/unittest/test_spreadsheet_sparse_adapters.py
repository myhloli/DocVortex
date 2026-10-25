"""表格快照适配的语义分层、素材访问顺序与私有钩子回归。"""

from types import MethodType
from unittest.mock import Mock

from openpyxl import Workbook
import pytest

from docvortex.analyzers.native.office.odf import table as odf_module
from docvortex.analyzers.native.office.odf.models import GridCell, TableGrid
from docvortex.analyzers.native.office.spreadsheet import projector as projector_module
from docvortex.analyzers.native.office.spreadsheet.projector import SpreadsheetProjector


def _projector() -> SpreadsheetProjector:
    """建立只含一个工作表的中立投影器，避免格式解析掩盖适配边界。"""
    projector = SpreadsheetProjector()
    projector.workbook = Workbook()
    return projector


def test_hyperlinked_whitespace_is_filtered_after_materialization(monkeypatch: pytest.MonkeyPatch) -> None:
    """空白链接在评分中无文本语义，但物化后有链接标记，不能套用原生预筛选下标。"""
    projector = _projector()
    sheet = projector.workbook.active
    sheet["A1"] = " "
    sheet["A1"].hyperlink = "https://example.test/whitespace"
    sheet["F10"] = "visible"
    original = projector_module.select_sparse_regions
    selections = []

    def observe(*args: object) -> tuple:
        """记录原生评分筛选结果，确认格式适配最终仍按物化语义决定保留集合。"""
        selected = original(*args)
        selections.append(selected)
        return selected

    monkeypatch.setattr(projector_module, "select_sparse_regions", observe)
    _, _, tables = projector._select_best_gap_candidate(sheet)

    assert len(selections) == 1
    assert selections[0][3] == [1]
    assert [table.anchor for table in tables] == [(0, 0), (5, 9)]
    assert tables[0].data[0].text == '<a href="https://example.test/whitespace"> </a>'


def test_hidden_merged_media_does_not_replace_materialized_semantics(monkeypatch: pytest.MonkeyPatch) -> None:
    """合并隐藏格媒体参与评分但不进入物化表，最终空语义表仍按旧规则过滤。"""
    projector = _projector()
    sheet = projector.workbook.active
    sheet["A1"] = " "
    sheet.merge_cells("A1:B1")
    sheet["F10"] = "visible"
    projector.table_image_map[(0, 1)] = ['<img src="hidden.png"/>']
    original = projector_module.select_sparse_regions
    snapshots = []
    selections = []

    def observe(*args: object) -> tuple:
        """记录连接矩形、评分语义和跨度，核对三类证据没有提前合并。"""
        snapshots.append(args)
        selected = original(*args)
        selections.append(selected)
        return selected

    monkeypatch.setattr(projector_module, "select_sparse_regions", observe)
    _, _, tables = projector._select_best_gap_candidate(sheet)

    assert snapshots[0][2] == [(0, 0, 0, 1)]
    assert set(snapshots[0][3]) == {(0, 1), (9, 5)}
    assert snapshots[0][4] == [(0, 0, 1, 2)]
    assert selections[0][3] == [0, 1]
    assert [table.anchor for table in tables] == [(5, 9)]


def test_dispimg_keeps_candidate_and_materialization_access_order(monkeypatch: pytest.MonkeyPatch) -> None:
    """有失败副作用的图片函数保持三档评分和最终物化的原始访问顺序。"""
    projector = _projector()
    sheet = projector.workbook.active
    first, second = 'DISPIMG("first")', 'DISPIMG("second")'
    sheet["A1"] = first
    sheet["F10"] = second
    calls = []

    def missing_image(_projector: SpreadsheetProjector, text: str) -> str:
        """模拟每次都可能产生诊断的图片缺失路径。"""
        calls.append(text)
        return ""

    projector._resolve_cell_image = MethodType(missing_image, projector)
    select = Mock(side_effect=AssertionError("DISPIMG must retain the callback selection path"))
    monkeypatch.setattr(projector_module, "select_sparse_regions", select)
    _, _, tables = projector._select_best_gap_candidate(sheet)

    select.assert_not_called()
    assert calls == [first, second] * 4
    assert len(tables) == 2


@pytest.mark.parametrize(
    "hook",
    ["_discover_sheet_regions", "_sheet_semantic_predicates", "_get_non_empty_cell_positions", "_get_merged_cell_lookup"],
)
def test_selection_preserves_overridden_private_hooks(monkeypatch: pytest.MonkeyPatch, hook: str) -> None:
    """覆写会被批处理绕过的私有钩子时，仍按原候选链调用这些钩子。"""
    projector = _projector()
    sheet = projector.workbook.active
    sheet["A1"] = "one"
    sheet["C1"] = "two"
    original = getattr(projector, hook)
    calls = []

    def observe(_projector: SpreadsheetProjector, *args: object, **kwargs: object) -> object:
        """保留原方法行为并观察覆写后的调用次数和参数。"""
        calls.append(args)
        return original(*args, **kwargs)

    monkeypatch.setattr(projector, hook, MethodType(observe, projector))
    select = Mock(side_effect=AssertionError("Overridden hooks must not be bypassed"))
    monkeypatch.setattr(projector_module, "select_sparse_regions", select)
    _, _, tables = projector._select_best_gap_candidate(sheet)

    select.assert_not_called()
    assert calls
    if hook == "_discover_sheet_regions":
        assert [args[2] for args in calls] == [0, 1, 2]
    if hook == "_get_non_empty_cell_positions":
        assert len(calls) == 3
    assert [cell.text for table in tables for cell in table.data if cell.text] == ["one", "two"]


def test_fixed_gap_respects_merged_lookup_override(monkeypatch: pytest.MonkeyPatch) -> None:
    """固定 gap 原生发现不能绕过定制的合并覆盖查询。"""
    projector = _projector()
    sheet = projector.workbook.active
    sheet["A1"] = "anchor"
    sheet.merge_cells("A1:C1")
    lookup = projector._get_merged_cell_lookup(sheet)
    lookup.contains_merged_cell = Mock(return_value=False)
    discover = Mock(side_effect=AssertionError("Overridden merge queries must remain callbacks"))
    monkeypatch.setattr(projector_module, "discover_sparse_regions", discover)

    regions = projector._discover_sheet_regions(sheet, projector._find_true_data_bounds(sheet), 0)

    discover.assert_not_called()
    assert [region.cells for region in regions] == [frozenset({(0, 0)})]
    assert lookup.contains_merged_cell.called


def test_unrelated_subclass_keeps_native_selection(monkeypatch: pytest.MonkeyPatch) -> None:
    """未覆写发现语义的子类仍走原生路径，不按整个类身份粗略排除。"""

    class ChildProjector(SpreadsheetProjector):
        """继承全部发现行为，仅用于验证细粒度钩子识别。"""

    projector = ChildProjector()
    projector.workbook = Workbook()
    projector.workbook.active["A1"] = "value"
    select = Mock(wraps=projector_module.select_sparse_regions)
    monkeypatch.setattr(projector_module, "select_sparse_regions", select)

    projector._select_best_gap_candidate(projector.workbook.active)

    select.assert_called_once()


def test_odf_snapshot_separates_covered_points_and_semantic_anchors(monkeypatch: pytest.MonkeyPatch) -> None:
    """ODS 覆盖占位参与连接但不变成语义起点，越界覆盖点也不会进入快照。"""
    grid = TableGrid(
        rows=[[GridCell(html="head", row_span=2, col_span=2), None, None], [None, None, None]],
        covered={(0, 1), (1, 0), (1, 1), (99, 99)},
    )
    select = Mock(wraps=odf_module.select_sparse_regions)
    monkeypatch.setattr(odf_module, "select_sparse_regions", select)

    regions = odf_module.split_table_regions(grid)

    starts, points, rectangles, semantic, spans, max_row, max_col = select.call_args.args
    assert starts == semantic == [(0, 0)]
    assert set(points) == {(0, 0), (0, 1), (1, 0), (1, 1)}
    assert rectangles == []
    assert spans == [(0, 0, 2, 2)]
    assert (max_row, max_col) == (1, 2)
    assert len(regions) == 1
    assert regions[0].width == 2
    assert regions[0].covered == {(0, 1), (1, 0), (1, 1)}


def test_odf_span_does_not_invent_missing_connected_placeholders(monkeypatch: pytest.MonkeyPatch) -> None:
    """直接构造的网格里跨度不隐式扩展连接点，只有显式 covered 才提供占位连接。"""
    grid = TableGrid(rows=[[GridCell(html="", col_span=3), None, None, GridCell(html="text")]])
    select = Mock(wraps=odf_module.select_sparse_regions)
    monkeypatch.setattr(odf_module, "select_sparse_regions", select)

    regions = odf_module.split_table_regions(grid)

    assert select.call_args.args[0] == [(0, 3)]
    assert set(select.call_args.args[1]) == {(0, 0), (0, 3)}
    assert select.call_args.args[2] == []
    assert select.call_args.args[4] == [(0, 0, 1, 3)]
    assert len(regions) == 1
    assert regions[0].rows[0][0].html == "text"


def test_odf_preserves_monkeypatched_callback_selector(monkeypatch: pytest.MonkeyPatch) -> None:
    """ODS 的旧模块选择钩子被替换后仍会生效，避免原生路径绕过已有差分入口。"""
    original = odf_module.select_best_gap_candidate
    select = Mock(wraps=original)
    native = Mock(side_effect=AssertionError("Patched callback selector must remain active"))
    monkeypatch.setattr(odf_module, "select_best_gap_candidate", select)
    monkeypatch.setattr(odf_module, "select_sparse_regions", native)

    regions = odf_module.split_table_regions(TableGrid(rows=[[GridCell(html="text")]]))

    select.assert_called_once()
    native.assert_not_called()
    assert regions[0].rows[0][0].html == "text"


@pytest.mark.parametrize("field", ["row_span", "col_span"])
def test_odf_float_span_preserves_python_type_error(monkeypatch: pytest.MonkeyPatch, field: str) -> None:
    """手工构造的浮点跨度须回到原 Python range 错误，不能被原生默认单格吞掉。"""
    from docvortex._compute_backend import get_native

    cell = GridCell(html="visible", **{field: 1.0})
    grid = TableGrid(rows=[[cell]])
    native_select = Mock(side_effect=AssertionError("float spans must retain Python callback semantics"))
    monkeypatch.setattr(get_native(), "spreadsheet_select_regions", native_select)

    with pytest.raises(TypeError, match="float.*integer"):
        odf_module.split_table_regions(grid)

    native_select.assert_not_called()


def test_odf_unbounded_python_row_span_preserves_bbox_clipping(monkeypatch: pytest.MonkeyPatch) -> None:
    """超出 i64 的 Python 行跨度仍先按 bbox 裁剪评分，最终保留原跨度及回调结果。"""
    import _spreadsheet_regions_oracle as oracle
    from docvortex._compute_backend import get_native

    huge_span = 10**100
    grid = TableGrid(rows=[[GridCell(html="visible", row_span=huge_span, col_span=1)]])
    native_select = Mock(side_effect=AssertionError("unbounded Python spans must not enter the i64 binding"))
    monkeypatch.setattr(get_native(), "spreadsheet_select_regions", native_select)
    with monkeypatch.context() as patch:
        patch.setattr(odf_module, "discover_connected_regions", oracle.discover_connected_regions)
        patch.setattr(odf_module, "select_best_gap_candidate", oracle.select_best_gap_candidate)
        patch.setattr(odf_module, "keep_maximal_by_semantic_sets", oracle.keep_maximal_by_semantic_sets)
        expected = odf_module.split_table_regions(grid)

    actual = odf_module.split_table_regions(grid)

    native_select.assert_not_called()
    assert actual == expected
    assert len(actual) == 1
    assert len(actual[0].rows) == actual[0].width == 1
    assert actual[0].rows[0][0].row_span == huge_span


def _contract_override_source(suffix: str, *, single_row: bool) -> bytes:
    """构造能区分 gap 选择结果的真实 XLSX/ODS，避免仅观察被替换函数的调用次数。"""
    from io import BytesIO
    from zipfile import ZipFile

    values = [["left", None, "right"]] if single_row else [["a", "b", None, "c", "d"], ["e", "f", None, "g", "h"]]
    output = BytesIO()
    if suffix == "xlsx":
        workbook = Workbook()
        for row_index, values_in_row in enumerate(values, start=1):
            for col_index, value in enumerate(values_in_row, start=1):
                if value is not None:
                    workbook.active.cell(row_index, col_index, value)
        workbook.save(output)
        return output.getvalue()

    rows = []
    for values_in_row in values:
        cells = [
            f"<table:table-cell><text:p>{value}</text:p></table:table-cell>" if value is not None else "<table:table-cell/>"
            for value in values_in_row
        ]
        rows.append("<table:table-row>" + "".join(cells) + "</table:table-row>")
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        '<office:body><office:spreadsheet><table:table table:name="Sheet">'
        + "".join(rows)
        + "</table:table></office:spreadsheet></office:body></office:document-content>"
    )
    with ZipFile(output, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", content)
    return output.getvalue()


@pytest.mark.parametrize("suffix", ["xlsx", "ods"])
@pytest.mark.parametrize("override", ["candidates", "preference", "margin", "penalty", "summary"])
def test_default_format_entries_respect_changed_region_contract(
    monkeypatch: pytest.MonkeyPatch, suffix: str, override: str
) -> None:
    """五类运行时评分契约变化均绕开原生选择，并真实改变默认 XLSX/ODS 的分表结果。"""
    from io import BytesIO

    from bs4 import BeautifulSoup

    from docvortex._compute_backend import get_native
    from docvortex.analyzers.native import OdsModel, XlsxModel
    from docvortex.analyzers.native.office.spreadsheet import region_discovery
    from docvortex.schema import BlockType

    payload = _contract_override_source(suffix, single_row=override == "margin")
    model_type = XlsxModel if suffix == "xlsx" else OdsModel
    baseline = model_type().predict(BytesIO(payload))
    baseline_tables = [block for page in baseline for block in page if block["type"] == BlockType.TABLE]
    assert len(baseline_tables) == (0 if override == "margin" else 1)

    native_select = Mock(side_effect=AssertionError("patched region contract must retain the Python candidate path"))
    monkeypatch.setattr(get_native(), "spreadsheet_select_regions", native_select)
    observed = []
    if override == "candidates":
        monkeypatch.setattr(region_discovery, "AUTO_GAP_TOLERANCE_CANDIDATES", (0,))
    elif override == "preference":
        # 原地更新同一字典，验证 guard 比较冻结值而不只是比较对象身份。
        monkeypatch.setitem(region_discovery.AUTO_GAP_TOLERANCE_PREFERENCE, 0, 0)
        monkeypatch.setitem(region_discovery.AUTO_GAP_TOLERANCE_PREFERENCE, 1, 1)
    elif override == "margin":
        monkeypatch.setattr(region_discovery, "AUTO_GAP_TOLERANCE_PREFERENCE_MARGIN", 1.0)
    elif override == "penalty":

        def prefer_separate_tables(summary: dict[str, float | int]) -> float:
            """令同排双表得分更低，确认替换的惩罚函数而非 Rust 固定权重实际参与选择。"""
            observed.append(float(summary["row_overlap_excess_ratio"]))
            return -float(summary["row_overlap_excess_ratio"])

        monkeypatch.setattr(region_discovery, "gap_candidate_penalty", prefer_separate_tables)
    else:
        original_summary = region_discovery.summarize_connected_regions

        def mark_joined_table_as_severe(regions, has_semantic_content, span_at):
            """仅改变单表候选的严重分隔指标，保留原汇总的其它指标和浮点语义。"""
            observed.append(len(regions))
            summary = original_summary(regions, has_semantic_content, span_at)
            if len(regions) == 1:
                summary["severe_separator_count"] += 1
            return summary

        monkeypatch.setattr(region_discovery, "summarize_connected_regions", mark_joined_table_as_severe)

    actual = model_type().predict(BytesIO(payload))

    native_select.assert_not_called()
    tables = [block for page in actual for block in page if block["type"] == BlockType.TABLE]
    assert len(tables) == (1 if override == "margin" else 2)
    visible_cells = [
        cell.get_text()
        for table in tables
        for cell in BeautifulSoup(table["content"], "html.parser").find_all(["th", "td"])
        if cell.get_text().strip()
    ]
    assert visible_cells == (["left", "right"] if override == "margin" else ["a", "b", "e", "f", "c", "d", "g", "h"])
    if override == "penalty":
        assert observed == [1.0, 0.0, 0.0]
    elif override == "summary":
        assert observed == [2, 1, 1]
