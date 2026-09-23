"""DOCX 表格匹配与完整文档上下文回退的回归测试。"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import mammoth
import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from loguru import logger

from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter
from docvortex.api import parse

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
ISSUE_CASES = ("control", "footnote", "endnote", "checkbox", "alternate", "field")
PROJECT_ROOT = Path(__file__).resolve().parents[2]


def _save(document: Document) -> bytes:
    """把内存中的 DOCX 文档序列化为字节。"""
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _table_html(converter: DocxConverter) -> list[str]:
    """按正文顺序提取转换器输出的表格 HTML。"""
    return [block["content"] for page in converter.pages for block in page if block["type"] == "table"]


def _attach_note(document: Document, *, note_type: str) -> None:
    """为脚注或尾注引用添加真实的 OOXML 包成员。"""
    relation_type = RT.FOOTNOTES if note_type == "footnote" else RT.ENDNOTES
    part = Part(
        PackURI(f"/word/{note_type}s.xml"),
        f"application/vnd.openxmlformats-officedocument.wordprocessingml.{note_type}s+xml",
        (
            f'<w:{note_type}s {NS}><w:{note_type} w:id="1">'
            "<w:p><w:r><w:t>注释正文</w:t></w:r></w:p>"
            f"</w:{note_type}></w:{note_type}s>"
        ).encode(),
        document.part.package,
    )
    document.part.relate_to(part, relation_type)


def _issue_trigger(name: str, *, checked: bool) -> str | None:
    """生成 issue 中会造成 XML 与 Mammoth 文本签名不一致的元素。"""
    if name in {"footnote", "endnote"}:
        return f'<w:r {NS}><w:{name}Reference w:id="1"/></w:r>'
    if name == "checkbox":
        value = "1" if checked else "0"
        return (
            f'<w:sdt {NS}><w:sdtPr><w14:checkbox><w14:checked w14:val="{value}"/>'
            "</w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:t>☐</w:t></w:r>"
            "</w:sdtContent></w:sdt>"
        )
    if name == "alternate":
        return (
            f'<mc:AlternateContent {NS}><mc:Choice Requires="w14"><w:r><w:t>Choice</w:t></w:r>'
            "</mc:Choice><mc:Fallback><w:r><w:t>Fallback</w:t></w:r>"
            "</mc:Fallback></mc:AlternateContent>"
        )
    if name == "field":
        return f'<w:fldSimple {NS} w:instr=" FILENAME "><w:r><w:t>report.docx</w:t></w:r></w:fldSimple>'
    return None


def _build_issue_document(*, styled: bool, checked: bool) -> bytes:
    """构造普通表、脚注、尾注、复选框、兼容分支和简单域混合文档。"""
    document = Document()
    _attach_note(document, note_type="footnote")
    _attach_note(document, note_type="endnote")
    for name in ISSUE_CASES:
        table = document.add_table(rows=2, cols=2)
        first = table.cell(0, 0).paragraphs[0]
        if styled:
            first.style = document.styles["List Paragraph"]
        trigger = _issue_trigger(name, checked=checked)
        if trigger:
            first._p.append(parse_xml(trigger))
        first.add_run(f"{name} A1")
        for row, col in ((0, 1), (1, 0), (1, 1)):
            table.cell(row, col).text = f"{name} {'AB'[col]}{row + 1}"
    return _save(document)


@pytest.mark.parametrize("styled", (True, False))
@pytest.mark.parametrize("checked", (True, False))
def test_issue_11_tables_match_and_survive(styled: bool, checked: bool) -> None:
    """验证所有触发条件均能匹配原表，并保持六张表的内容与顺序。"""
    file_bytes = _build_issue_document(styled=styled, checked=checked)
    converter = DocxConverter()
    preparsed = converter._preparse_tables_with_mammoth(file_bytes)

    assert len(preparsed) == len(ISSUE_CASES)
    assert all(html is not None for html in preparsed)
    converter.convert(BytesIO(file_bytes))
    html_tables = _table_html(converter)
    assert len(html_tables) == len(ISSUE_CASES)
    for name, html in zip(ISSUE_CASES, html_tables):
        cells = BeautifulSoup(html, "html.parser").find_all("td")
        assert len(cells) == 4
        assert f"{name} A1" in cells[0].get_text()
        assert cells[-1].get_text() == f"{name} B2"
    checkbox = BeautifulSoup(html_tables[3], "html.parser").find("input")
    assert checkbox is not None
    assert checkbox.has_attr("checked") is checked


def test_issue_11_public_parse_keeps_all_tables() -> None:
    """从公开 parse 入口验证 issue 的各张表仍在结构化结果中。"""
    result = parse(_build_issue_document(styled=True, checked=True), file_suffix="docx")
    serialized = json.dumps(result.to_dict(), ensure_ascii=False)
    assert all(f"{name} B2" in serialized for name in ISSUE_CASES)


def test_note_marker_filter_keeps_literal_bracketed_number() -> None:
    """只从签名排除 Mammoth 生成的注释引用，保留单元格正文中的 [1]。"""
    html = BeautifulSoup(
        '<table><tr><td>[1]<sup><a id="footnote-ref-1" href="#footnote-1">[1]</a></sup>'
        '<sup><a id="endnote-ref-2" href="#endnote-2">[2]</a></sup></td></tr></table>',
        "html.parser",
    ).table
    assert html is not None
    assert DocxConverter._html_table_signature(html)["text"] == "[1]"


def _build_fallback_document() -> bytes:
    """构造带显式编号、样式、链接、图片、公式和纵向合并的回退表格。"""
    document = Document()
    for label in ("first", "second"):
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).merge(table.cell(1, 0))
        cell = table.cell(0, 0)
        paragraph = cell.paragraphs[0]
        paragraph.style = document.styles["List Number"]
        paragraph.add_run(label)
        if label == "first":
            ppr = paragraph._p.get_or_add_pPr()
            ppr.append(parse_xml(f'<w:numPr {NS}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'))
            rid = document.part.relate_to("https://example.com/docx-table", RT.HYPERLINK, is_external=True)
            paragraph._p.append(parse_xml(f'<w:hyperlink {NS} r:id="{rid}"><w:r><w:t>linked</w:t></w:r></w:hyperlink>'))
            png = BytesIO()
            Image.new("RGB", (3, 3), "red").save(png, format="PNG")
            png.seek(0)
            paragraph.add_run().add_picture(png)
            styled = cell.add_paragraph("styled")
            styled.style = document.styles["List Paragraph"]
            formula_paragraph = table.cell(0, 2).paragraphs[0]
            math = etree.SubElement(formula_paragraph._p, qn("m:oMath"))
            math_run = etree.SubElement(math, qn("m:r"))
            etree.SubElement(math_run, qn("m:t")).text = "x+y"
        table.cell(0, 1).text = f"{label} B1"
        table.cell(1, 1).text = f"{label} B2"
        table.cell(1, 2).text = f"{label} C2"
    output = BytesIO(_save(document))
    mammoth.embed_style_map(output, "p[style-name='List Paragraph'] => h6:fresh")
    return output.getvalue()


@pytest.mark.parametrize("preparsed", ([], [None, None]))
def test_forced_fallback_keeps_full_document_context(monkeypatch: pytest.MonkeyPatch, preparsed: list) -> None:
    """强制跳过预匹配，验证回退在两张表中保留全部受支持的上下文内容。"""
    file_bytes = _build_fallback_document()
    converter = DocxConverter()

    def skip_preparse(_data: bytes) -> list:
        """按参数强制表格走完整上下文回退。"""
        return preparsed

    monkeypatch.setattr(converter, "_preparse_tables_with_mammoth", skip_preparse)
    converter.convert(BytesIO(file_bytes))
    html_tables = _table_html(converter)
    assert len(html_tables) == 2
    first, second = [BeautifulSoup(html, "html.parser") for html in html_tables]
    assert first.find("li") is not None
    assert first.find("a", href="https://example.com/docx-table") is not None
    image = first.find("img")
    assert image is not None
    assert image.get("src", "").startswith("data:image/")
    assert first.find("eq") is not None
    assert first.find("h6") is not None
    assert first.find("td", rowspan="2") is not None
    assert "first B2" in first.get_text()
    assert "second B2" in second.get_text()
    assert converter._mammoth_fallback_context is None
    assert converter._fallback_docx_bytes is None


def test_alignment_skips_extra_non_body_table() -> None:
    """额外的文本框候选表不得抢占正文表格的匹配位置。"""
    file_bytes = _build_issue_document(styled=True, checked=False)
    document = Document(BytesIO(file_bytes))
    body_tables = [child for child in document.element.body if child.tag == qn("w:tbl")]
    converter = DocxConverter()
    raw = mammoth.convert_to_html(BytesIO(file_bytes), transform_document=converter._mammoth_top_level_table_document).value
    html_tables = BeautifulSoup(raw, "html.parser").find_all("table")
    extra = BeautifulSoup("<table><tr><td>textbox only</td></tr></table>", "html.parser").table
    aligned = converter._align_mammoth_tables_to_xml_tables([extra, *html_tables], body_tables, document.part)
    assert len(aligned) == len(ISSUE_CASES)
    assert all(html is not None for html in aligned)
    assert "textbox only" not in "".join(aligned)


def test_fallback_failure_does_not_skip_later_table(monkeypatch: pytest.MonkeyPatch) -> None:
    """一张回退表解析失败时，警告定位表号且后续表仍输出。"""
    converter = DocxConverter()

    def skip_preparse(_data: bytes) -> list:
        """强制后续两张表进入完整上下文回退。"""
        return []

    monkeypatch.setattr(converter, "_preparse_tables_with_mammoth", skip_preparse)
    original = converter._mammoth_fallback_html
    attempts = 0

    def fail_first(element: object) -> str:
        """仅模拟首张表发生上下文回退异常。"""
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise ValueError("injected failure")
        return original(element)

    warnings = []
    sink_id = logger.add(warnings.append, level="WARNING")
    try:
        monkeypatch.setattr(converter, "_mammoth_fallback_html", fail_first)
        converter.convert(BytesIO(_build_fallback_document()))
    finally:
        logger.remove(sink_id)
    assert len(_table_html(converter)) == 1
    assert "second B2" in _table_html(converter)[0]
    assert any("table #1, full-context fallback" in str(warning) for warning in warnings)
    assert converter._mammoth_fallback_context is None


def test_fallback_converter_reuse_and_error_release(monkeypatch: pytest.MonkeyPatch) -> None:
    """同一转换器的连续调用和失败重试都释放旧包及回退状态。"""
    converter = DocxConverter()

    def skip_preparse(_data: bytes) -> list:
        """强制各次转换均使用完整上下文回退。"""
        return []

    monkeypatch.setattr(converter, "_preparse_tables_with_mammoth", skip_preparse)
    first = _build_fallback_document()
    converter.convert(BytesIO(first))
    first_tables = _table_html(converter)
    assert converter._mammoth_fallback_context is None

    original_add_header_footer = converter._add_header_footer

    def fail_after_tables(_document: Document) -> None:
        """模拟完成表格回退后发生的文档级异常。"""
        raise RuntimeError("injected lifecycle failure")

    monkeypatch.setattr(converter, "_add_header_footer", fail_after_tables)
    with pytest.raises(RuntimeError, match="injected lifecycle failure"):
        converter.convert(BytesIO(first))
    assert converter._mammoth_fallback_context is None
    assert converter._fallback_docx_bytes is None

    monkeypatch.setattr(converter, "_add_header_footer", original_add_header_footer)
    with pytest.raises(Exception):
        converter.convert(BytesIO(b"invalid docx"))
    assert converter._mammoth_fallback_context is None
    assert converter._fallback_docx_bytes is None

    converter.convert(BytesIO(_build_issue_document(styled=True, checked=True)))
    assert len(_table_html(converter)) == len(ISSUE_CASES)
    assert "first B2" not in "".join(_table_html(converter))
    converter.convert(BytesIO(first))
    assert _table_html(converter) == first_tables
    assert converter._mammoth_fallback_context is None


@pytest.mark.parametrize(
    ("path", "expected_count", "expected_rows"),
    [
        ("tests/fixtures/docx/issue9-merged-cells.docx", 1, 37),
        ("demo/ms_office_docs/docx_01.docx", 8, None),
    ],
)
def test_existing_docx_samples_keep_tables(path: str, expected_count: int, expected_rows: int | None) -> None:
    """回放真实 DOCX 样本，检查表格数量及 issue #9 的完整行数。"""
    file_bytes = (PROJECT_ROOT / path).read_bytes()
    converter = DocxConverter()
    preparsed = converter._preparse_tables_with_mammoth(file_bytes)
    assert len(preparsed) == expected_count
    assert all(html is not None for html in preparsed)
    converter.convert(BytesIO(file_bytes))
    html_tables = _table_html(converter)
    assert len(html_tables) == expected_count
    assert html_tables == [converter._normalize_table_colspans(html) for html in preparsed]
    if expected_rows is not None:
        assert len(BeautifulSoup(html_tables[0], "html.parser").find_all("tr")) == expected_rows
