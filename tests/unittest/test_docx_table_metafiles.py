"""DOCX 表格图片的矢量转换与位图保真回归测试。"""

from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pytest
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from metafile_render import MetafileError
from PIL import Image

from _metafile_test_utils import basic_wmf, build_emf, emf_rectangle
from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter
from docvortex.api import parse, render
from docvortex.export.bundle import load_bundle
from docvortex.foundation._image_payload import extract_generated_svg_fallback
from docvortex.render.contracts import RenderFormat


def _image_bytes(image_format: str) -> bytes:
    """生成同时可验证原始字节与透明度的微型位图。"""
    image = Image.new("RGBA" if image_format == "PNG" else "RGB", (4, 3), (60, 120, 180, 80))
    output = BytesIO()
    image.save(output, format=image_format)
    return output.getvalue()


def _docx_with_image(
    payload: bytes,
    *,
    content_type: str,
    extension: str,
    nested: bool = False,
    next_table: bool = False,
) -> bytes:
    """从合法位图图片关系构造带指定 MIME 和图片载荷的 DOCX。"""
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 1).text = "right cell"
    if nested:
        table.cell(0, 0).paragraphs[0].add_run("outer cell")
        inner = table.cell(0, 0).add_table(rows=1, cols=1)
        paragraph = inner.cell(0, 0).paragraphs[0]
    else:
        paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("before ")
    picture = paragraph.add_run().add_picture(BytesIO(_image_bytes("PNG")))
    picture._inline.docPr.set("descr", "Table logo")
    paragraph.add_run(" after")
    if next_table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "later table"

    stream = BytesIO()
    document.save(stream)
    destination = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(destination, "w") as target:
        media_name = next(name for name in source.namelist() if name.startswith("word/media/"))
        replacement_name = media_name.rsplit(".", 1)[0] + f".{extension}"
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                for child in root:
                    if child.get("Extension") == "png":
                        child.set("Extension", extension)
                        child.set("ContentType", content_type)
                data = etree.tostring(root)
            elif member.filename.endswith(".rels"):
                data = data.replace(media_name.rsplit("/", 1)[1].encode(), replacement_name.rsplit("/", 1)[1].encode())
            elif member.filename == media_name:
                data = payload
            name = replacement_name if member.filename == media_name else member.filename
            target.writestr(name, data, compress_type=member.compress_type)
    return destination.getvalue()


def _parse_document(data: bytes, *, fallback: bool, monkeypatch: pytest.MonkeyPatch, keep_model_json: bool = False):
    """按真实公共入口解析，必要时强制表格进入完整上下文回退。"""
    if fallback:
        monkeypatch.setattr(DocxConverter, "_preparse_tables_with_mammoth", lambda _self, _bytes: [])
    return parse(data, file_suffix="docx", keep_model_json=keep_model_json)


def _table_content(result) -> list[str]:
    """提取物化后的表格 HTML，用于验证图片路径与文字。"""
    return [
        child["content"]
        for page in result.to_dict()["pages"]
        for block in page["blocks"]
        if block["type"] == "table"
        for child in block["content"]
        if child["type"] == "table_body"
    ]


@pytest.mark.parametrize("fallback", [False, True], ids=["preparse", "fallback"])
@pytest.mark.parametrize(
    ("content_type", "extension", "payload"),
    [
        ("image/x-wmf", "wmf", basic_wmf()),
        ("image/wmf", "wmf", basic_wmf()),
        ("image/x-emf", "emf", build_emf([emf_rectangle(0, 0, 80, 60)])),
        ("image/emf", "emf", build_emf([emf_rectangle(0, 0, 80, 60)])),
    ],
)
def test_table_metafile_survives_parse_and_materialization(
    monkeypatch: pytest.MonkeyPatch, content_type: str, extension: str, payload: bytes, fallback: bool
) -> None:
    """四种 MIME 均经两条表格路径生成可物化 SVG，并保留单元格文字。"""
    document = _docx_with_image(payload, content_type=content_type, extension=extension)
    result = _parse_document(document, fallback=fallback, monkeypatch=monkeypatch)
    table = BeautifulSoup(_table_content(result)[0], "html.parser")
    image = table.find("img")
    assert image is not None
    assert image["alt"] == "Table logo"
    assert "before" in table.get_text() and "after" in table.get_text()
    assert "right cell" in table.get_text()
    path = image["src"]
    assert path.endswith(".svg")
    assert list(result.assets) == [path]
    fallback_png, width, height = extract_generated_svg_fallback(result.assets[path])
    assert fallback_png.startswith(b"\x89PNG\r\n\x1a\n")
    assert width > 0 and height > 0


@pytest.mark.parametrize("fallback", [False, True], ids=["preparse", "fallback"])
@pytest.mark.parametrize(
    ("content_type", "extension"),
    [("application/x-msmetafile", "wmf"), ("image/png", "png")],
)
def test_table_metafile_recognized_by_mime_or_payload(
    monkeypatch: pytest.MonkeyPatch, content_type: str, extension: str, fallback: bool
) -> None:
    """别名及错误 MIME 下的真实 WMF 载荷仍生成 SVG。"""
    result = _parse_document(
        _docx_with_image(basic_wmf(), content_type=content_type, extension=extension),
        fallback=fallback,
        monkeypatch=monkeypatch,
    )
    assert len(result.assets) == 1
    assert next(iter(result.assets)).endswith(".svg")


@pytest.mark.parametrize("fallback", [False, True], ids=["preparse", "fallback"])
@pytest.mark.parametrize("image_format", ["PNG", "JPEG"])
def test_table_raster_bytes_and_alt_are_preserved(monkeypatch: pytest.MonkeyPatch, image_format: str, fallback: bool) -> None:
    """普通位图不重编码，透明 PNG 与替代文字原样保留。"""
    payload = _image_bytes(image_format)
    extension = "png" if image_format == "PNG" else "jpg"
    mime = "image/png" if image_format == "PNG" else "image/jpeg"
    result = _parse_document(
        _docx_with_image(payload, content_type=mime, extension=extension),
        fallback=fallback,
        monkeypatch=monkeypatch,
    )
    image = BeautifulSoup(_table_content(result)[0], "html.parser").find("img")
    assert image is not None and image["alt"] == "Table logo"
    assert result.assets[image["src"]] == payload
    if image_format == "PNG":
        with Image.open(BytesIO(result.assets[image["src"]])) as decoded:
            assert decoded.mode == "RGBA" and decoded.getpixel((0, 0))[3] == 80


@pytest.mark.parametrize("fallback", [False, True], ids=["preparse", "fallback"])
@pytest.mark.parametrize("failure", ["invalid", "render"])
def test_unrenderable_table_vector_uses_placeholder_and_keeps_next_table(
    monkeypatch: pytest.MonkeyPatch, failure: str, fallback: bool
) -> None:
    """损坏矢量图与渲染异常只影响该图片，后续表格仍输出。"""
    if failure == "render":
        from docvortex.analyzers.native.office import image as office_image

        def fail_render(*_args, **_kwargs):
            """模拟 metafile-render 无法处理输入载荷。"""
            raise MetafileError("test render failure")

        monkeypatch.setattr(office_image, "render_metafile", fail_render)
    payload = b"invalid wmf" if failure == "invalid" else basic_wmf()
    result = _parse_document(
        _docx_with_image(payload, content_type="image/x-wmf", extension="wmf", next_table=True),
        fallback=fallback,
        monkeypatch=monkeypatch,
    )
    tables = _table_content(result)
    assert len(tables) == 2 and "later table" in tables[1]
    image = BeautifulSoup(tables[0], "html.parser").find("img")
    assert image is not None and image["src"].endswith(".jpg")
    assert result.assets[image["src"]].startswith(b"\xff\xd8\xff")


@pytest.mark.parametrize("fallback", [False, True], ids=["preparse", "fallback"])
def test_nested_table_metafile_and_export_roundtrip(monkeypatch: pytest.MonkeyPatch, tmp_path, fallback: bool) -> None:
    """嵌套表格的矢量素材可保存结果包，并在 HTML 与 DOCX 中使用。"""
    result = _parse_document(
        _docx_with_image(basic_wmf(), content_type="image/x-wmf", extension="wmf", nested=True),
        fallback=fallback,
        monkeypatch=monkeypatch,
        keep_model_json=True,
    )
    table = BeautifulSoup(_table_content(result)[0], "html.parser")
    assert len(table.find_all("table")) == 2
    assert "outer cell" in table.get_text() and "right cell" in table.get_text()
    image = table.find("img")
    assert image is not None and image["src"] in result.assets

    result.save_bundle(tmp_path / "bundle")
    restored = load_bundle(tmp_path / "bundle")
    assert restored.model_json is not None
    assert restored.assets[image["src"]] == result.assets[image["src"]]
    html = render(restored.middle_json, RenderFormat.HTML, assets=restored.assets)
    assert image["src"].encode() in html.content
    docx = render(restored.middle_json, RenderFormat.DOCX, assets=restored.assets)
    exported_document = Document(BytesIO(docx.content))
    assert len(exported_document.tables) == 1
    assert len(exported_document.tables[0]._element.xpath(".//a:blip")) == 1
    with ZipFile(BytesIO(docx.content)) as package:
        png_names = [name for name in package.namelist() if name.startswith("word/media/") and name.endswith(".png")]
        assert len(png_names) == 1
        with Image.open(BytesIO(package.read(png_names[0]))) as fallback_image:
            fallback_image.verify()
