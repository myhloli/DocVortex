from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from _native_test_utils import analyze_native_test_document
from _ofd_test_utils import build_multi_document_ofd, build_ofd_package, page_xml, path_object, text_object
from _span_test_utils import inline_text
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from PIL import Image

from docvortex.analyzers.native import OfdModel
from docvortex.analyzers.native.ofd import OfdParseError, OfdResourceLimitError, detect_ofd
from docvortex.analyzers.native.ofd import images as ofd_images
from docvortex.analyzers.native.ofd import metadata as ofd_metadata
from docvortex.analyzers.native.ofd import scene as ofd_scene
from docvortex.analyzers.native.ofd import table as ofd_table
from docvortex.analyzers.native.ofd import text as ofd_text
from docvortex.analyzers.native.ofd.constants import (
    MAX_DELTA_TOKENS,
    MAX_DOCUMENT_COUNT,
    MAX_DRAW_PARAM_INHERITANCE,
    MAX_EXPANDED_GLYPHS,
    MAX_GLYPH_TOKENS,
    MAX_PAGE_COUNT,
    MAX_PATH_TOKENS,
)
from docvortex.analyzers.native.ofd.geometry import Affine, canonical_angle
from docvortex.analyzers.native.ofd.images import build_image_item
from docvortex.analyzers.native.ofd.models import AxisLine, MediaResource, OfdPageScene, ResourceRegistry, TextLine
from docvortex.analyzers.native.ofd.package import OfdPackage
from docvortex.analyzers.native.ofd.path import OfdPathBudget, _segments, build_axis_lines
from docvortex.analyzers.native.ofd.reading_order import OfdReadingOrderProjector
from docvortex.analyzers.native.ofd.resources import parse_resource_part, resolve_draw_param
from docvortex.analyzers.native.ofd.text import FontMetricResolver, OfdTextBudget, build_text_lines, parse_delta
from docvortex.render import render_docx, render_html, render_markdown
from docvortex.schema import BlockType

_PROJECT_ROOT = Path(__file__).resolve().parents[2]


_LOCAL_SAMPLE_DIR = _PROJECT_ROOT / "tmp" / "ofd_samples" / "ofdrw_issues_20260828" / "ofd"


def _minimal_payload(*, namespace: str = "http://www.ofdspec.org/2016", version: str = "1.0") -> bytes:
    """Constructs a minimal OFD containing a single line of text."""
    content = text_object(3, "你好，OFD！", boundary="10 10 50 12", delta_x="g 6 5")
    return build_ofd_package(
        [("Pages/Page_0/Content.xml", page_xml(content, namespace=namespace))],
        namespace=namespace,
        version=version,
    )


def _replace_package_part(payload: bytes, part_name: str, replacement: bytes | None) -> bytes:
    """Replace or delete a single member in test OFD ZIP."""
    source_buffer = BytesIO(payload)
    output_buffer = BytesIO()
    with ZipFile(source_buffer) as source, ZipFile(output_buffer, "w", ZIP_DEFLATED) as output:
        for info in source.infolist():
            if info.filename == part_name:
                if replacement is not None:
                    output.writestr(info, replacement)
                continue
            output.writestr(info, source.read(info.filename))
    return output_buffer.getvalue()


def test_ofd_legacy_namespace_and_declared_page_order() -> None:
    """Verify that old namespaces, non-directory sorted page trees, and empty pages maintain declaration order."""
    namespace = "http://www.ofdspec.org"
    pages = [
        (
            "Pages/Page_0/Content.xml",
            page_xml(text_object(1, "first", boundary="10 10 30 10"), namespace=namespace),
        ),
        (
            "Pages/Page_2/Content.xml",
            page_xml(text_object(2, "second", boundary="10 10 30 10"), namespace=namespace),
        ),
        ("Pages/Page_1/Content.xml", page_xml("", namespace=namespace)),
    ]
    middle, _model = analyze_native_test_document(
        build_ofd_package(pages, namespace=namespace, version="1.2"), file_suffix="ofd"
    )

    assert [page.page_idx for page in middle.pages] == [0, 1, 2]
    assert [inline_text(page.blocks[0].content) if page.blocks else "" for page in middle.pages] == ["first", "second", ""]


def test_ofd_multiple_doc_bodies_flatten_in_declared_order() -> None:
    """Verify that multiple DocBodys expand into contiguous physical pages in declaration order."""
    middle, model = analyze_native_test_document(build_multi_document_ofd(), file_suffix="ofd")

    assert [page.page_idx for page in middle.pages] == [0, 1]
    assert [inline_text(page[0]["content"]) for page in model.pages] == ["doc-zero", "doc-one"]


def test_ofd_declared_page_count_is_bounded_for_parser_and_metadata() -> None:
    """Verification of repeated references to the same page part is also subject to the full-text page budget."""
    payload = build_ofd_package([("Pages/Page_0/Content.xml", page_xml(""))])
    source_buffer = BytesIO(payload)
    output_buffer = BytesIO()
    with ZipFile(source_buffer) as source, ZipFile(output_buffer, "w", ZIP_DEFLATED) as output:
        document_root = etree.fromstring(source.read("Doc_0/Document.xml"))
        pages = next(element for element in document_root.iter() if etree.QName(element).localname == "Pages")
        pages.clear()
        namespace = etree.QName(document_root).namespace
        for page_index in range(MAX_PAGE_COUNT + 1):
            etree.SubElement(
                pages,
                f"{{{namespace}}}Page",
                ID=str(page_index + 1),
                BaseLoc="Pages/Page_0/Content.xml",
            )
        replacement = etree.tostring(document_root, xml_declaration=True, encoding="UTF-8")
        for info in source.infolist():
            output.writestr(info, replacement if info.filename == "Doc_0/Document.xml" else source.read(info.filename))
    oversized = output_buffer.getvalue()

    with pytest.raises(OfdResourceLimitError, match="max_page_count"):
        OfdModel().predict(BytesIO(oversized))
    with pytest.raises(OfdResourceLimitError, match="max_page_count"):
        ofd_metadata.extract_ofd_metadata(BytesIO(oversized))


def test_ofd_page_count_budget_is_shared_across_doc_bodies(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify parsing and metadata page budgets shared by multiple DocBodys."""
    monkeypatch.setattr(ofd_scene, "MAX_PAGE_COUNT", 1)
    monkeypatch.setattr(ofd_metadata, "MAX_PAGE_COUNT", 1)
    payload = build_multi_document_ofd()

    with pytest.raises(OfdResourceLimitError, match="max_page_count"):
        OfdModel().predict(BytesIO(payload))
    with pytest.raises(OfdResourceLimitError, match="max_page_count"):
        ofd_metadata.extract_ofd_metadata(BytesIO(payload))


def test_ofd_declared_document_count_is_bounded_before_page_parsing() -> None:
    """Verify large number of empty DocBody subject to independent budget before materializing document references."""
    payload = build_ofd_package([("Pages/Page_0/Content.xml", page_xml(""))])
    source_buffer = BytesIO(payload)
    output_buffer = BytesIO()
    with ZipFile(source_buffer) as source, ZipFile(output_buffer, "w", ZIP_DEFLATED) as output:
        ofd_root = etree.fromstring(source.read("OFD.xml"))
        for child in list(ofd_root):
            ofd_root.remove(child)
        namespace = etree.QName(ofd_root).namespace
        for _document_index in range(MAX_DOCUMENT_COUNT + 1):
            doc_body = etree.SubElement(ofd_root, f"{{{namespace}}}DocBody")
            doc_root = etree.SubElement(doc_body, f"{{{namespace}}}DocRoot")
            doc_root.text = "Doc_0/Document.xml"
        document_root = etree.fromstring(source.read("Doc_0/Document.xml"))
        pages = next(element for element in document_root.iter() if etree.QName(element).localname == "Pages")
        pages.clear()
        replacements = {
            "OFD.xml": etree.tostring(ofd_root, xml_declaration=True, encoding="UTF-8"),
            "Doc_0/Document.xml": etree.tostring(document_root, xml_declaration=True, encoding="UTF-8"),
        }
        for info in source.infolist():
            content = replacements[info.filename] if info.filename in replacements else source.read(info.filename)
            output.writestr(info, content)
    oversized = output_buffer.getvalue()

    with pytest.raises(OfdResourceLimitError, match="max_document_count"):
        OfdModel().predict(BytesIO(oversized))
    with pytest.raises(OfdResourceLimitError, match="max_document_count"):
        ofd_metadata.extract_ofd_metadata(BytesIO(oversized))


def test_ofd_textcode_geometry_does_not_use_oversized_boundary() -> None:
    """Verify that Foxit style oversized Boundary will not be the final text bbox."""
    content = text_object(
        81,
        "661016910189",
        boundary="-8.8175 -170.35411 361.51749 399.96179",
        size=9,
        x=83.00001,
        y=168.00009,
        delta_x="g 11 4.158",
        ctm="0.3527 0 0 0.3527 0 139.66919",
    )
    middle, _model = analyze_native_test_document(
        build_ofd_package([("Pages/Page_0/Content.xml", page_xml(content, physical_box="0 0 210 140"))]), file_suffix="ofd"
    )

    bbox = middle.pages[0].blocks[0].bbox
    assert bbox is not None
    assert bbox[2] - bbox[0] < 0.5
    assert bbox[3] - bbox[1] < 0.2


def test_ofd_cardinal_text_directions_keep_geometry_and_angle() -> None:
    """Verify that ReadDirection/CharDirection participates in the glyph quad and the sort direction."""
    content = (
        '<ofd:TextObject ID="8" Boundary="20 20 20 50" Font="1" Size="5" '
        'ReadDirection="90" CharDirection="90">'
        '<ofd:TextCode X="2" Y="-5" DeltaX="g 3 6">竖排文</ofd:TextCode>'
        "</ofd:TextObject>"
    )
    _middle, model = analyze_native_test_document(
        build_ofd_package([("Pages/Page_0/Content.xml", page_xml(content))]), file_suffix="ofd"
    )

    assert model.pages[0][0]["angle"] == 90
    assert model.pages[0][0]["bbox"][2] > model.pages[0][0]["bbox"][0]
    assert model.pages[0][0]["bbox"][3] > model.pages[0][0]["bbox"][1]


@pytest.mark.parametrize(("second_x", "expected"), [(16.0, "Hello"), (17.0, "Hel lo")])
def test_ofd_same_baseline_ascii_fragments_use_measured_gap(second_x: float, expected: str) -> None:
    """Verify adjacent English run Only pad spaces when there is a visible word gap."""
    content = "".join(
        [
            text_object(1, "Hel", boundary="10 10 10 10", size=4, y=5, delta_x="2 2"),
            text_object(2, "lo", boundary=f"{second_x} 10 10 10", size=4, y=5, delta_x="2"),
        ]
    )

    _middle, model = analyze_native_test_document(
        build_ofd_package([("Pages/Page_0/Content.xml", page_xml(content))]), file_suffix="ofd"
    )

    assert len(model.pages[0]) == 1
    assert inline_text(model.pages[0][0]["content"]) == expected


def test_ofd_textcode_preserves_boundary_whitespace_and_glyph_positions() -> None:
    """Verify that TextCode leading and trailing spaces participate in Delta expansion and CGTransform global position mapping."""
    text_element = etree.fromstring(
        b'<TextObject ID="9" Boundary="10 10 50 10" Font="1" Size="5">'
        b'<TextCode X="1" Y="5" DeltaX="5 7"> A </TextCode>'
        b'<CGTransform CodePosition="1" CodeCount="1"><Glyphs>42</Glyphs></CGTransform>'
        b"</TextObject>"
    )
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")
    package = OfdPackage(package_buffer.getvalue())

    lines = build_text_lines(
        text_element,
        parent_transform=Affine(),
        parent_clip=(0.0, 0.0, 100.0, 100.0),
        resources=ResourceRegistry(),
        package=package,
        font_metrics=FontMetricResolver(package),
        budget=OfdTextBudget(),
        paint_order=0,
        layer_type="body",
        template_id=None,
    )

    assert len(lines) == 1
    assert lines[0].text == " A "
    assert [glyph.glyph_id for glyph in lines[0].glyphs] == [None, 42, None]
    assert [glyph.origin[0] for glyph in lines[0].glyphs] == pytest.approx([11.0, 16.0, 23.0])


def test_ofd_cgtransform_expands_only_actual_text_positions() -> None:
    """Verify that the oversized CodeCount only maps character positions where TextCode actually exists."""
    glyphs = "42 " + "999 " * 100_000
    text_element = etree.fromstring(
        (
            '<TextObject ID="10" Boundary="0 0 10 10" Font="1" Size="5">'
            '<TextCode X="1" Y="5">A</TextCode>'
            f'<CGTransform CodePosition="0" CodeCount="1000000000"><Glyphs>{glyphs}</Glyphs></CGTransform>'
            "</TextObject>"
        ).encode()
    )
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")
    package = OfdPackage(package_buffer.getvalue())
    budget = OfdTextBudget()

    lines = build_text_lines(
        text_element,
        parent_transform=Affine(),
        parent_clip=(0.0, 0.0, 100.0, 100.0),
        resources=ResourceRegistry(),
        package=package,
        font_metrics=FontMetricResolver(package),
        budget=budget,
        paint_order=0,
        layer_type="body",
        template_id=None,
    )

    assert budget.glyph_count == 1
    assert budget.glyph_mapping_count == 1
    assert budget.glyph_token_count == 1
    assert len(lines) == 1
    assert [glyph.glyph_id for glyph in lines[0].glyphs] == [42]

    exhausted_budget = OfdTextBudget(glyph_mapping_count=MAX_EXPANDED_GLYPHS)
    with pytest.raises(OfdResourceLimitError, match="max_expanded_glyphs"):
        build_text_lines(
            text_element,
            parent_transform=Affine(),
            parent_clip=(0.0, 0.0, 100.0, 100.0),
            resources=ResourceRegistry(),
            package=package,
            font_metrics=FontMetricResolver(package),
            budget=exhausted_budget,
            paint_order=0,
            layer_type="body",
            template_id=None,
        )

    token_exhausted_budget = OfdTextBudget(glyph_token_count=MAX_GLYPH_TOKENS)
    with pytest.raises(OfdResourceLimitError, match="max_glyph_tokens"):
        build_text_lines(
            text_element,
            parent_transform=Affine(),
            parent_clip=(0.0, 0.0, 100.0, 100.0),
            resources=ResourceRegistry(),
            package=package,
            font_metrics=FontMetricResolver(package),
            budget=token_exhausted_budget,
            paint_order=0,
            layer_type="body",
            template_id=None,
        )


def test_ofd_delta_tokens_are_streamed_and_bounded() -> None:
    """Verification Delta Scans only the required token, with controlled failure if the full text accumulation exceeds the limit."""
    budget = OfdTextBudget()

    assert parse_delta("1 " + "2 " * 100_000, 1, budget) == [1.0]
    assert budget.delta_token_count == 1
    assert parse_delta("g 3 2", 4, budget) == [2.0, 2.0, 2.0, 2.0]

    exhausted = OfdTextBudget(delta_token_count=MAX_DELTA_TOKENS)
    with pytest.raises(OfdResourceLimitError, match="max_delta_tokens"):
        parse_delta("invalid 1", 1, exhausted)


def test_ofd_textcode_is_decoded_and_charged_in_bounded_chunks(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that hex escaping is recoverable across shards and does not decode the full TextCode first when exceeding the limit."""
    monkeypatch.setattr(ofd_text, "_TEXT_CODE_DECODE_CHUNK_SIZE", 3)
    escaped = etree.fromstring(b"<TextCode>before\\00<Part>41</Part>after</TextCode>")
    budget = OfdTextBudget()

    assert ofd_text._decode_text_code_element(escaped, budget) == "beforeAafter"
    assert budget.glyph_count == len("beforeAafter")

    monkeypatch.setattr(ofd_text, "_TEXT_CODE_DECODE_CHUNK_SIZE", 16)
    oversized = etree.fromstring(f"<TextCode>{'A' * 100_000}</TextCode>".encode())
    original_decode = ofd_text.decode_text_code
    decoded_chunk_lengths: list[int] = []

    def record_decode(value: str) -> str:
        """Record the size of a single decoding and entrust the actual escape implementation."""
        decoded_chunk_lengths.append(len(value))
        return original_decode(value)

    monkeypatch.setattr(ofd_text, "decode_text_code", record_decode)
    exhausted = OfdTextBudget(glyph_count=MAX_EXPANDED_GLYPHS)
    with pytest.raises(OfdResourceLimitError, match="max_expanded_glyphs"):
        ofd_text._decode_text_code_element(oversized, exhausted)

    assert decoded_chunk_lengths == [16]


def test_ofd_textcode_discards_glyphs_outside_object_boundary() -> None:
    """Verification TextObject Boundary clips characters and blank lines that are completely outside the boundary."""
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")
    package = OfdPackage(package_buffer.getvalue())
    font_metrics = FontMetricResolver(package)
    common = {
        "parent_transform": Affine(),
        "parent_clip": (0.0, 0.0, 100.0, 100.0),
        "resources": ResourceRegistry(),
        "package": package,
        "font_metrics": font_metrics,
        "budget": OfdTextBudget(),
        "paint_order": 0,
        "layer_type": "body",
        "template_id": None,
    }

    partial = etree.fromstring(
        b'<TextObject ID="10" Boundary="0 0 10 10" Size="5"><TextCode X="2" Y="5" DeltaX="20">AB</TextCode></TextObject>'
    )
    outside = etree.fromstring(
        b'<TextObject ID="11" Boundary="0 0 10 10" Size="5"><TextCode X="20" Y="5">X</TextCode></TextObject>'
    )

    partial_lines = build_text_lines(partial, **common)
    outside_lines = build_text_lines(outside, **common)

    assert len(partial_lines) == 1
    assert partial_lines[0].text == "A"
    assert [glyph.text for glyph in partial_lines[0].glyphs] == ["A"]
    assert outside_lines == []


def test_ofd_draw_param_inheritance_is_ordered_and_bounded() -> None:
    """Verification DrawParam maintains parent-to-child override order and provides controlled failure when the inheritance chain exceeds limits."""
    ordered = ResourceRegistry(
        draw_params={
            1: {"ID": "1", "Relative": "2", "LineWidth": "2"},
            2: {"ID": "2", "LineWidth": "1", "Color": "red"},
        }
    )
    assert resolve_draw_param(ordered, 1) == {"LineWidth": "2", "Color": "red"}

    chain_length = MAX_DRAW_PARAM_INHERITANCE + 1
    oversized = ResourceRegistry(
        draw_params={
            resource_id: (
                {"ID": str(resource_id), "Relative": str(resource_id + 1)}
                if resource_id + 1 < chain_length
                else {"ID": str(resource_id), "LineWidth": "1"}
            )
            for resource_id in range(chain_length)
        }
    )
    with pytest.raises(OfdResourceLimitError, match="max_draw_param_inheritance"):
        resolve_draw_param(oversized, 0)


@pytest.mark.parametrize(
    ("raw_angle", "expected_angle"),
    [
        (89.9, 90),
        (90.1, 90),
        (179.9, 180),
        (180.1, 180),
        (359.9, 0),
        (0.1, 0),
    ],
)
def test_ofd_image_near_cardinal_rotation_accepts_both_directions(raw_angle: float, expected_angle: int) -> None:
    """Verify that loading and reading order are preserved when image rotation is approached from both sides of a right angle block."""
    image_buffer = BytesIO()
    Image.new("RGB", (2, 2), "white").save(image_buffer, format="PNG")
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")
        archive.writestr("Res/image.png", image_buffer.getvalue())
    package = OfdPackage(package_buffer.getvalue())
    transform = Affine.rotation(raw_angle)
    ctm = " ".join(str(value) for value in (transform.a, transform.b, transform.c, transform.d, transform.e, transform.f))
    image_element = etree.fromstring(f'<ImageObject ID="1" Boundary="0 0 10 10" ResourceID="1" CTM="{ctm}"/>'.encode())
    resources = ResourceRegistry(
        media={
            1: MediaResource(
                resource_id=1,
                media_type="Image",
                media_format="PNG",
                media_part="Res/image.png",
            )
        }
    )

    item = build_image_item(
        image_element,
        parent_transform=Affine(),
        parent_clip=(0.0, 0.0, 100.0, 100.0),
        resources=resources,
        package=package,
        paint_order=0,
        layer_type="body",
        template_id=None,
    )

    assert item is not None
    assert item.diagnostic is None
    assert item.image_base64
    scene = OfdPageScene(
        page_idx=0,
        physical_box=(0.0, 0.0, 100.0, 100.0),
        content_box=None,
        images=[item],
    )
    blocks = OfdReadingOrderProjector([scene]).project_page(scene)
    assert len(blocks) == 1
    assert blocks[0]["type"] == BlockType.IMAGE
    assert blocks[0]["image_base64"] == item.image_base64
    assert canonical_angle(raw_angle) == expected_angle


def test_ofd_oversized_image_is_rejected_before_pixel_decode(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verification OFD out of limit raster downgraded to no load diagnostic before Pillow load."""
    image = MagicMock()
    image.__enter__.return_value = image
    image.size = (8_193, 1)
    image.load.side_effect = AssertionError("oversized image must not be decoded")
    monkeypatch.setattr(ofd_images.Image, "open", lambda _source: image)
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")
        archive.writestr("Res/image.png", b"oversized")
    package = OfdPackage(package_buffer.getvalue())
    image_element = etree.fromstring(b'<ImageObject ID="1" Boundary="0 0 10 10" ResourceID="1"/>')

    item = build_image_item(
        image_element,
        parent_transform=Affine(),
        parent_clip=(0.0, 0.0, 100.0, 100.0),
        resources=ResourceRegistry(
            media={1: MediaResource(1, "Image", "PNG", "Res/image.png")},
        ),
        package=package,
        paint_order=0,
        layer_type="body",
        template_id=None,
    )

    assert item is not None
    assert item.image_base64 is None
    assert item.diagnostic == "unsupported_image_payload"
    image.load.assert_not_called()


def test_ofd_referenced_media_resource_and_part_are_required() -> None:
    """Validation body ImageObject failed when referencing resource ID or member that does not exist."""
    image_element = etree.fromstring(b'<ImageObject ID="7" Boundary="0 0 10 10" ResourceID="1"/>')
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")

    with OfdPackage(package_buffer.getvalue()) as package:
        with pytest.raises(OfdParseError, match="missing media resource"):
            build_image_item(
                image_element,
                parent_transform=Affine(),
                parent_clip=(0.0, 0.0, 100.0, 100.0),
                resources=ResourceRegistry(),
                package=package,
                paint_order=0,
                layer_type="body",
                template_id=None,
            )

    with OfdPackage(package_buffer.getvalue()) as package:
        with pytest.raises(OfdParseError, match="missing required part"):
            build_image_item(
                image_element,
                parent_transform=Affine(),
                parent_clip=(0.0, 0.0, 100.0, 100.0),
                resources=ResourceRegistry(media={1: MediaResource(1, "Image", "PNG", "Res/missing.png")}),
                package=package,
                paint_order=0,
                layer_type="body",
                template_id=None,
            )


def test_ofd_path_rejects_unexpected_numeric_tokens_without_hanging() -> None:
    """Verify that no active command or number after C token directly degrades the current illegal path."""
    assert _segments("1 2", Affine(), OfdPathBudget()) == []
    assert _segments("M 0 0 C 1 2 L 3 4", Affine(), OfdPathBudget()) == []
    assert _segments("M 0 0 L 10 0", Affine(), OfdPathBudget()) == [((0.0, 0.0), (10.0, 0.0))]


def test_ofd_path_tokens_are_streamed_and_bounded() -> None:
    """Verification stops immediately when there are no active commands and fails when path token accumulation exceeds limit."""
    stray_budget = OfdPathBudget()

    assert _segments("1 " * 100_000, Affine(), stray_budget) == []
    assert stray_budget.token_count == 1
    assert stray_budget.command_count == 0

    exhausted = OfdPathBudget(token_count=MAX_PATH_TOKENS)
    with pytest.raises(OfdResourceLimitError, match="max_path_tokens"):
        _segments("M 0 0", Affine(), exhausted)


def test_ofd_path_segments_are_clipped_to_object_boundary() -> None:
    """Verify that complete out-of-bounds paths are discarded, and crossing paths are clipped to the intersection of the object and parent clipping."""
    outside = etree.fromstring(
        b'<PathObject ID="1" Boundary="10 10 10 1" LineWidth="0.2">'
        b"<AbbreviatedData>M 20 0.5 L 30 0.5</AbbreviatedData></PathObject>"
    )
    outside_lines = build_axis_lines(
        outside,
        parent_transform=Affine(),
        parent_clip=(0.0, 0.0, 100.0, 100.0),
        paint_order=0,
        template_id=None,
        budget=OfdPathBudget(),
    )
    assert outside_lines == []

    cases = [
        ("M -5 5 L 15 5", "horizontal", (12.0, 14.9, 18.0, 15.1)),
        ("M 5 -5 L 5 15", "vertical", (14.9, 12.0, 15.1, 18.0)),
    ]
    for data, orientation, expected_bbox in cases:
        crossing = etree.fromstring(
            (
                '<PathObject ID="2" Boundary="10 10 10 10" LineWidth="0.2">'
                f"<AbbreviatedData>{data}</AbbreviatedData></PathObject>"
            ).encode()
        )
        lines = build_axis_lines(
            crossing,
            parent_transform=Affine(),
            parent_clip=(12.0, 12.0, 18.0, 18.0),
            paint_order=0,
            template_id=None,
            budget=OfdPathBudget(),
        )

        assert len(lines) == 1
        assert lines[0].orientation == orientation
        assert lines[0].bbox == pytest.approx(expected_bbox)


def test_ofd_template_grid_recovers_table() -> None:
    """Validate template paths and page text together to restore high-confidence full-line tables."""
    paths = "".join(
        [
            path_object(10, boundary="10 20 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(11, boundary="10 50 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(12, boundary="10 80 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(13, boundary="10 20 0.2 60", data="M 0.1 0 L 0.1 60"),
            path_object(14, boundary="50 20 0.2 60", data="M 0.1 0 L 0.1 60"),
            path_object(15, boundary="90 20 0.2 60", data="M 0.1 0 L 0.1 60"),
        ]
    )
    template = page_xml(paths)
    body = "".join(
        [
            text_object(21, "A", boundary="20 30 10 8", size=4, y=4),
            text_object(22, "B", boundary="60 30 10 8", size=4, y=4),
            text_object(23, "C", boundary="20 60 10 8", size=4, y=4),
            text_object(24, "D", boundary="60 60 10 8", size=4, y=4),
        ]
    )
    payload = build_ofd_package(
        [("Pages/Page_0/Content.xml", page_xml(body, template_id=5))],
        templates={5: ("Tpls/Tpl_0/Content.xml", template)},
    )
    middle, model = analyze_native_test_document(payload, file_suffix="ofd")

    assert [block["type"] for block in model.pages[0]] == [BlockType.TABLE]
    assert all(value in model.pages[0][0]["content"] for value in ("A", "B", "C", "D"))
    assert middle.pages[0].blocks[0].type == BlockType.TABLE


def test_ofd_table_styles_use_standard_html_across_renderers() -> None:
    """Verify that the OFD table directly outputs standard labels, and the styles are retained by Markdown, HTML, and DOCX."""
    paths = "".join(
        [
            path_object(10, boundary="10 20 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(11, boundary="10 50 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(12, boundary="10 80 80 0.2", data="M 0 0.1 L 80 0.1"),
            path_object(13, boundary="10 20 0.2 60", data="M 0.1 0 L 0.1 60"),
            path_object(14, boundary="50 20 0.2 60", data="M 0.1 0 L 0.1 60"),
            path_object(15, boundary="90 20 0.2 60", data="M 0.1 0 L 0.1 60"),
        ]
    )
    body = "".join(
        [
            text_object(21, "A", boundary="20 30 10 8", size=4, y=4),
            text_object(22, "B", boundary="60 30 10 8", size=4, y=4),
            text_object(23, "C", boundary="20 60 10 8", size=4, y=4),
            text_object(24, "D", boundary="60 60 10 8", size=4, y=4),
        ]
    )
    page_resource = (
        '<ofd:Res xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Fonts>'
        '<ofd:Font ID="1" FontName="Styled" Bold="true" Italic="true"/></ofd:Fonts></ofd:Res>'
    )
    payload = build_ofd_package(
        [("Pages/Page_0/Content.xml", page_xml(paths + body, page_res="PageRes.xml"))],
        extra_parts={"Doc_0/Pages/Page_0/PageRes.xml": page_resource},
    )

    middle, model = analyze_native_test_document(payload, file_suffix="ofd")
    table_html = model.pages[0][0]["content"]

    assert "<text" not in table_html
    assert "<em><strong>A</strong></em>" in table_html
    assert "***A***" in render_markdown(middle)
    html_soup = BeautifulSoup(render_html(middle), "html.parser")
    assert html_soup.select_one("td em strong").get_text() == "A"
    document = Document(BytesIO(render_docx(middle)))
    styled_run = next(run for paragraph in document.tables[0].cell(0, 0).paragraphs for run in paragraph.runs if run.text)
    assert styled_run.text == "A" and styled_run.bold and styled_run.italic


def test_ofd_table_intersection_budget_is_shared_across_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify that the table segment comparison budget is accumulated by the same projector across pages."""

    def scene(page_idx: int) -> OfdPageScene:
        """Construct a single table page that requires six pairwise comparisons."""
        axis_lines = [
            AxisLine((0.0, 0.0, 10.0, 0.2), "horizontal", 0.2, 0, None),
            AxisLine((0.0, 10.0, 10.0, 10.2), "horizontal", 0.2, 1, None),
            AxisLine((0.0, 0.0, 0.2, 10.0), "vertical", 0.2, 2, None),
            AxisLine((10.0, 0.0, 10.2, 10.0), "vertical", 0.2, 3, None),
        ]
        text_lines = [
            TextLine("A", (2.0, 2.0, 4.0, 4.0), [], 0, 2.0, 4, 1, "Body", None),
            TextLine("B", (2.0, 6.0, 4.0, 8.0), [], 0, 2.0, 5, 2, "Body", None),
        ]
        return OfdPageScene(page_idx, (0.0, 0.0, 20.0, 20.0), None, text_lines, axis_lines)

    monkeypatch.setattr(ofd_table, "MAX_TABLE_INTERSECTION_CHECKS", 8)

    with pytest.raises(OfdResourceLimitError, match="max_table_intersection_checks"):
        OfdReadingOrderProjector([scene(0), scene(1)]).project()


def test_ofd_page_resource_overrides_document_resource() -> None:
    """Verification exception duplicate resource ID is resolved by the rules of PageRes which is higher than PublicRes."""
    page_resource = (
        '<ofd:Res xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Fonts>'
        '<ofd:Font ID="1" FontName="Page Bold" Bold="true"/></ofd:Fonts></ofd:Res>'
    )
    payload = build_ofd_package(
        [
            (
                "Pages/Page_0/Content.xml",
                page_xml(text_object(1, "page-style", boundary="10 10 40 10"), page_res="PageRes.xml"),
            )
        ],
        extra_parts={"Doc_0/Pages/Page_0/PageRes.xml": page_resource},
    )
    _middle, model = analyze_native_test_document(payload, file_suffix="ofd")

    assert model.pages[0][0]["content"] == [{"type": "text", "content": "page-style", "styles": ["bold"]}]


@pytest.mark.parametrize(
    ("replacement", "error"),
    [
        (None, "invalid required location"),
        (b"<broken", "invalid XML part"),
        (b'<Res xmlns="urn:unsupported"/>', "unsupported namespace"),
    ],
)
def test_ofd_declared_resource_part_is_required(replacement: bytes | None, error: str) -> None:
    """Validation failed entirely when a declared resource member is missing, corrupted, or with a namespace error."""
    payload = _replace_package_part(_minimal_payload(), "Doc_0/PublicRes.xml", replacement)

    with pytest.raises(OfdParseError, match=error):
        OfdModel().predict(BytesIO(payload))


def test_ofd_absent_resource_declaration_is_optional_but_empty_declaration_is_invalid() -> None:
    """Verify that the undeclared resource is legal, and explicit null resource references are treated as broken packages."""
    payload = _minimal_payload()
    with ZipFile(BytesIO(payload)) as package:
        document_root = etree.fromstring(package.read("Doc_0/Document.xml"))
    public_res = next(element for element in document_root.iter() if etree.QName(element).localname == "PublicRes")

    parent = public_res.getparent()
    assert parent is not None
    parent.remove(public_res)
    without_declaration = _replace_package_part(
        payload,
        "Doc_0/Document.xml",
        etree.tostring(document_root, xml_declaration=True, encoding="UTF-8"),
    )
    assert OfdModel().predict(BytesIO(without_declaration))

    public_res = etree.SubElement(parent, "{http://www.ofdspec.org/2016}PublicRes")
    public_res.text = ""
    empty_declaration = _replace_package_part(
        payload,
        "Doc_0/Document.xml",
        etree.tostring(document_root, xml_declaration=True, encoding="UTF-8"),
    )
    with pytest.raises(OfdParseError, match="invalid required location"):
        OfdModel().predict(BytesIO(empty_declaration))


def test_parse_resource_part_none_returns_empty_registry() -> None:
    """Verify that optional members are not read when the caller explicitly passes in an undeclared resource."""
    package_buffer = BytesIO()
    with ZipFile(package_buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("OFD.xml", "<OFD/>")

    with OfdPackage(package_buffer.getvalue()) as package:
        registry = parse_resource_part(package, None)

    assert not registry.fonts and not registry.media and not registry.composites and not registry.draw_params


def test_ofd_does_not_open_malformed_unreferenced_custom_tag() -> None:
    """Verifying extension XML that is corrupted but does not participate in the body does not break parsing."""
    payload = build_ofd_package(
        [("Pages/Page_0/Content.xml", page_xml(text_object(1, "visible", boundary="10 10 30 10")))],
        extra_parts={"Doc_0/Tags/CustomTag.xml": "<broken"},
    )

    middle, _model = analyze_native_test_document(payload, file_suffix="ofd")
    assert inline_text(middle.pages[0].blocks[0].content) == "visible"


def test_ofd_rejects_foreign_namespace_and_non_v1_version() -> None:
    """Verify that unknown namespaces and non-1.x versions will not be accepted loosely."""
    assert not detect_ofd(_minimal_payload(namespace="https://example.com/ofd"))
    assert not detect_ofd(_minimal_payload(version="2.0"))
    with pytest.raises(OfdParseError):
        OfdModel().predict(BytesIO(_minimal_payload(namespace="https://example.com/ofd")))


def test_ofd_package_rejects_unsafe_member_and_dtd() -> None:
    """Verify ZIP up-hop member with any DTD rejected before body parsing."""
    unsafe = BytesIO()
    with ZipFile(unsafe, "w", ZIP_DEFLATED) as package:
        package.writestr("OFD.xml", '<ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016" Version="1.0"/>')
        package.writestr("../escape.xml", "unsafe")
    with pytest.raises(OfdParseError, match="unsafe member"):
        OfdPackage(unsafe.getvalue())

    dtd = BytesIO()
    with ZipFile(dtd, "w", ZIP_DEFLATED) as package:
        package.writestr(
            "OFD.xml",
            '<!DOCTYPE OFD [<!ENTITY x "boom">]><ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016" '
            'Version="1.0" DocType="OFD"><ofd:DocBody><ofd:DocRoot>&x;</ofd:DocRoot></ofd:DocBody></ofd:OFD>',
        )
    with OfdPackage(dtd.getvalue()) as package:
        with pytest.raises(OfdParseError, match="DTD"):
            package.root()


_LOCAL_SAMPLE_CASES = [
    ("issue_10_issue_10_a7b6f9eceb_ofdrw-converter_src_test_resources_helloworld_a7b6f9eceb.ofd", 1),
    ("issue_200_issue_200_a2b9080ae3_ofdrw-converter_src_test_resources_999_a2b9080ae3.ofd", 5),
    ("issue_208_issue_208_dfed483fa0_ofdrw-layout_src_test_resources_AddWatermarkAnnot_dfed483fa0.ofd", 6),
    ("issue_208_issue_208_e7cac8d149_ofdrw-layout_src_test_resources_no_page_container_e7cac8d149.ofd", 1),
    ("issue_385_issue_385_bff244b113_ofdrw-layout_src_test_resources_keyword2_bff244b113.ofd", 1),
    ("issue_183_issue_183_08a57b9736_48.3_2_08a57b9736.ofd", 1),
    ("issue_190_issue_190_6eadbaeb7b_文档2_6eadbaeb7b.ofd", 2),
    ("issue_385_issue_385_2be07ec79d_ofdrw-layout_src_test_resources_1-1_2be07ec79d.ofd", 3),
    ("issue_271_源文件_bfb11cbcab.ofd", 9),
    ("issue_293_issue_293_87024f3f31_test.zip_test_fcb3a84887.ofd", 2),
    ("issue_312_issue_312_b8973c8f48_测试_b8973c8f48.ofd", 1),
]


@pytest.mark.parametrize(("filename", "page_count"), _LOCAL_SAMPLE_CASES)
def test_ofd_local_real_samples_preserve_pages_and_normalized_bboxes(filename: str, page_count: int) -> None:
    """Verify page count, stable parsing, and bbox contract when user local real corpus exists."""
    source = _LOCAL_SAMPLE_DIR / filename
    if not source.exists():
        pytest.skip("local OFD sample corpus is unavailable")
    payload = source.read_bytes()
    assert hashlib.sha256(payload).hexdigest()[:10] in filename
    middle, model = analyze_native_test_document(payload, file_suffix="ofd")
    assert len(model.pages) == len(middle.pages) == page_count
    for page in middle.pages:
        for block in page.blocks:
            assert block.bbox is not None
            assert all(0 <= value <= 1 for value in block.bbox)
