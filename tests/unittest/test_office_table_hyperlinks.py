"""表格链接按源 XML 顺序保留，公式重建与回退不丢失链接。"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches

from docvortex.api import parse, render
from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter


URLS = ("https://example.org/one?a=1&b=2", "https://example.org/two")


def _docx_link(paragraph: Any, url: str) -> None:
    """在真实 relationship 上构造两个 run 组成的同名链接。"""
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    for text in ("same", " label"):
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
        link.append(run)
    paragraph._p.append(link)


def make_table_document(suffix: str) -> bytes:
    """生成含重复标签、换行、合并单元格和不安全链接的 Office 表格。"""
    if suffix == "docx":
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        para = table.cell(0, 0).paragraphs[0]
        _docx_link(para, URLS[0])
        from docvortex.render._internal.docx.math import latex_to_omml

        para._p.append(latex_to_omml("x+1", display=False))
        para.add_run().add_break()
        _docx_link(para, URLS[1])
        _docx_link(table.cell(1, 0).paragraphs[0], "javascript:alert(1)")
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(7), Inches(2))
        table = shape.table
        table.cell(0, 0).merge(table.cell(0, 1))
        para = table.cell(0, 0).text_frame.paragraphs[0]
        for index, url in enumerate(URLS):
            if index:
                para.add_line_break()
            for label in ("same", " label"):
                run = para.add_run()
                run.text = label
                run.hyperlink.address = url
                run.font.bold = True
        run = table.cell(1, 0).text_frame.paragraphs[0].add_run()
        run.text = "same label"
        run.hyperlink.address = "javascript:alert(1)"
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize("suffix,fallback", [("pptx", False), ("docx", False), ("docx", True)])
def test_table_links_survive_parse_render(suffix: str, fallback: bool, monkeypatch: pytest.MonkeyPatch) -> None:
    """公开解析和三种输出均保留 URL 对应关系，不安全 URL 仅保留可见文本。"""
    if fallback:
        monkeypatch.setattr(DocxConverter, "_preparse_tables_with_mammoth", lambda *args: [])
    result = parse(make_table_document(suffix), file_suffix=suffix)
    raw = json.dumps(result.to_dict(), ensure_ascii=False)
    assert all(url.replace("&", "&amp;") in raw or url in raw for url in URLS)
    assert "javascript:" not in raw
    for fmt in ("markdown", "html", "structured_content"):
        content = render(result.middle_json, fmt).content.decode()
        assert all(url.replace("&", "&amp;") in content or url in content for url in URLS)
        assert "javascript:" not in content
    table_html = next(
        block.content[0].content for page in result.middle_json.pages for block in page.blocks if block.type == "table"
    )
    soup = BeautifulSoup(table_html, "html.parser")
    grouped: list[tuple[str, str]] = []
    for link in soup.find_all("a"):
        label, url = link.get_text(), link["href"]
        if grouped and grouped[-1][1] == url:
            grouped[-1] = (grouped[-1][0] + label, url)
        else:
            grouped.append((label, url))
    assert grouped == [("same label", url) for url in URLS]
    assert soup.find("br") is not None
    assert soup.find(attrs={"colspan": "2"}) is not None
