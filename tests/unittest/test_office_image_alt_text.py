"""图片替代文本（pptx cNvPr/descr、docx docPr/descr、odf svg:title+desc）进入图片块 content 的测试。"""

from __future__ import annotations

import base64
from io import BytesIO
from zipfile import ZIP_STORED, ZipFile

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from docvortex.analyzers.native import DocxModel, OdtModel, PptxModel
from docvortex.schema import BlockType

# 1x1 透明 PNG
PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _image_blocks(pages: list[list[dict]]) -> list[dict]:
    """收集转换输出中的图片块（含 image_body 子类型）。"""
    return [
        block
        for page in pages
        for block in page
        if block.get("type") in {BlockType.IMAGE, BlockType.IMAGE_BODY}
    ]


def _build_pptx_with_descr() -> bytes:
    """构造带 cNvPr descr 替代文本的单图演示文稿。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = slide.shapes.add_picture(BytesIO(PIXEL_PNG), 0, 0, width=Inches(4), height=Inches(3))
    picture._element.nvPicPr.cNvPr.set("descr", "Quarterly numbers")
    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def _build_docx_with_descr() -> bytes:
    """构造带 wp:docPr descr 替代文本的单图文档。"""
    document = Document()
    document.add_picture(BytesIO(PIXEL_PNG))
    document.inline_shapes[-1]._inline.docPr.set("descr", "A cute mascot")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _build_odt_with_alt() -> bytes:
    """构造 svg:title/svg:desc 替代文本随 draw:frame 存储的最小 ODT 包。"""
    content_xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
    xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
    xmlns:xlink="http://www.w3.org/1999/xlink"
    xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
    xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
  <office:body><office:text>
    <text:p>before image</text:p>
    <draw:frame>
      <svg:title>Mascot</svg:title>
      <svg:desc>A cute mascot</svg:desc>
      <draw:image>
        <office:binary-data>{base64.b64encode(PIXEL_PNG).decode()}</office:binary-data>
      </draw:image>
    </draw:frame>
  </office:text></office:body>
</office:document-content>"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
  <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>
  <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""
    output = BytesIO()
    with ZipFile(output, "w") as package:
        # mimetype 必须是首个成员且不压缩，ODF 规范要求。
        package.writestr(_stored_info("mimetype"), b"application/vnd.oasis.opendocument.text")
        package.writestr("META-INF/manifest.xml", manifest)
        package.writestr("content.xml", content_xml)
    return output.getvalue()


def _stored_info(name: str):
    """构造 STORED 存储方式的 zip 成员信息。"""
    from zipfile import ZipInfo

    info = ZipInfo(name)
    info.compress_type = ZIP_STORED
    return info


def test_pptx_descr_becomes_image_content() -> None:
    """pptx 图片 cNvPr 的 descr 应写入图片块 content。"""
    pages = PptxModel().predict(BytesIO(_build_pptx_with_descr()))
    image_blocks = _image_blocks(pages)
    assert image_blocks
    assert any(block.get("content") == "Quarterly numbers" for block in image_blocks)


def test_docx_descr_becomes_image_content() -> None:
    """docx DrawingML wp:docPr 的 descr 应写入图片块 content。"""
    pages = DocxModel().predict(BytesIO(_build_docx_with_descr()))
    image_blocks = _image_blocks(pages)
    assert image_blocks
    assert any(block.get("content") == "A cute mascot" for block in image_blocks)


def test_odt_alt_becomes_image_content() -> None:
    """odt draw:frame 的 svg:title+svg:desc 应写入图片块 content。"""
    pages = OdtModel().predict(BytesIO(_build_odt_with_alt()))
    image_blocks = _image_blocks(pages)
    assert image_blocks
    assert any(block.get("content") == "Mascot A cute mascot" for block in image_blocks)
