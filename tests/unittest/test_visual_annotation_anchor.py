"""Issue #35：视觉说明书签的模型、解析、导出及 HTML 往返回归。"""

from __future__ import annotations

import base64
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pydantic import ValidationError
from pypdf import PdfReader
from jsonschema import Draft202012Validator
from _native_test_utils import analyze_native_test_document

from docvortex import analyze, parse
from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter
from docvortex.postprocess.pages import blocks_to_page_info
from docvortex.render import (
    PdfLayout,
    RenderMode,
    render_docx,
    render_epub,
    render_html,
    render_latex,
    render_markdown,
    render_pdf,
    render_structured_content,
)
from docvortex.schema import IndexBlock, MiddleJson, PageInfo, TextBlock, parse_block

PREFIXES = ("image", "table", "chart", "code")
ANNOTATION_TYPES = tuple(f"{prefix}_{role}" for prefix in PREFIXES for role in ("caption", "footnote"))
WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
XHTML_NS = {"x": "http://www.w3.org/1999/xhtml"}
# 由修复前 a75bbe5d 的独立源码进程生成，防止可选空字段改变既有 EPUB 标识。
NO_ANCHOR_EPUB_IDS = {
    "image": "urn:uuid:3ff77694-9c91-58d6-b715-afa7804cb881",
    "table": "urn:uuid:4e072362-3a49-5d8b-b191-f69578cd6cf8",
    "chart": "urn:uuid:b95ff33e-f93e-5160-ac37-5062e34d4461",
    "code": "urn:uuid:6b671ad4-1cb7-5953-bb81-1a12f2aeba7c",
}


def _inline(text: str) -> list[dict[str, str]]:
    """构造可序列化的最小行内内容。"""
    return [{"type": "text", "content": text}]


def _png() -> bytes:
    """生成各渲染器都能实际解码的图片，避免占位输出掩盖书签问题。"""
    output = BytesIO()
    Image.new("RGB", (40, 20), (30, 80, 130)).save(output, format="PNG")
    return output.getvalue()


def _visual_document(
    prefix: str, *, anchors: tuple[str | None, str | None] = ("caption target", "footnote target")
) -> MiddleJson:
    """构造目录及行内链接前向引用两个视觉说明的两页严格文档。"""
    body = {"type": f"{prefix}_body", "index": 0, "bbox": [0.1, 0.2, 0.9, 0.45], "content": ""}
    if prefix in {"image", "chart"}:
        body["image_base64"] = f"data:image/png;base64,{base64.b64encode(_png()).decode('ascii')}"
    elif prefix == "table":
        body["content"] = "<table><tr><td>Cell A</td><td>Cell B</td></tr></table>"
    else:
        body["content"] = "print('hello')"
    visual = {
        "type": prefix,
        "index": 0,
        "bbox": body["bbox"],
        "content": [
            {
                "type": f"{prefix}_caption",
                "index": 1,
                "bbox": [0.1, 0.12, 0.9, 0.18],
                "anchor": anchors[0],
                "content": _inline(f"{prefix} caption"),
            },
            body,
            {
                "type": f"{prefix}_footnote",
                "index": 2,
                "bbox": [0.1, 0.5, 0.9, 0.58],
                "anchor": anchors[1],
                "content": _inline(f"{prefix} footnote"),
            },
        ],
    }
    if prefix == "code":
        visual.update(sub_type="code", guess_lang="python")
    links = [
        {"type": "hyperlink", "url": f"#{anchor}", "content": _inline(f"Jump {role}")}
        for anchor, role in zip(anchors, ("caption", "footnote"), strict=True)
        if anchor
    ]
    return MiddleJson(
        pages=[
            PageInfo(
                page_idx=0,
                blocks=[
                    TextBlock(type="text", index=0, bbox=(0.1, 0.1, 0.9, 0.15), content=links),
                    IndexBlock(
                        type="index",
                        index=1,
                        bbox=(0.1, 0.2, 0.9, 0.3),
                        content=[
                            TextBlock(type="text", anchor=anchor, content=_inline(f"{prefix} {role}\t2"))
                            for anchor, role in zip(anchors, ("caption", "footnote"), strict=True)
                        ],
                    ),
                ],
            ),
            PageInfo(page_idx=1, blocks=[parse_block(visual)]),
        ],
        is_full_document=True,
        metadata={"file_suffix": "docx", "producer": {"name": "docvortex", "version": "test"}},
        extensions={},
    )


def _zip_xml(payload: bytes, name: str) -> etree._Element:
    """解析导出容器里的 XML，验证真实书签和超链接而非文本子串。"""
    with ZipFile(BytesIO(payload)) as archive:
        return etree.fromstring(archive.read(name))


@pytest.mark.parametrize("block_type", ANNOTATION_TYPES)
def test_annotation_anchor_is_optional_strict_and_roundtrips(block_type: str) -> None:
    """验证八种说明接受可选字符串书签，同时保持未知字段和类型严格校验。"""
    payload = {"type": block_type, "content": _inline("Note")}
    assert parse_block(payload).anchor is None
    assert parse_block({**payload, "anchor": None}).anchor is None
    anchored = parse_block({**payload, "anchor": "_Hlk46718182"})
    assert parse_block(anchored.to_dict()) == anchored
    assert type(anchored).model_validate_json(anchored.to_json()) == anchored
    for extra in ({"anchor": 12}, {"unexpected": True}):
        with pytest.raises(ValidationError):
            parse_block({**payload, **extra})


@pytest.mark.parametrize("prefix", PREFIXES)
def test_checked_middle_schema_and_structured_anchor_defaults(prefix: str) -> None:
    """验证提交的 JSON Schema 接受说明书签，缺省字段不改变结构化输出。"""
    middle = _visual_document(prefix)
    schema = json.loads((Path(__file__).resolve().parents[2] / "schemas/middle-2.0.json").read_text())
    Draft202012Validator(schema).validate(middle.to_dict())
    assert MiddleJson.from_json(middle.to_json()) == middle
    for child in middle.pages[1].blocks[0].content:
        if str(child.type).endswith(("caption", "footnote")):
            assert schema["$defs"][type(child).__name__] == MiddleJson.model_json_schema()["$defs"][type(child).__name__]
            child.anchor = None
    visual = render_structured_content(middle)["pages"][1]["blocks"][0]
    assert "anchor" not in visual["captions"][0]
    assert "anchor" not in visual["footnotes"][0]


@pytest.mark.parametrize("prefix", PREFIXES)
def test_default_annotation_anchor_preserves_existing_epub_identifier(prefix: str) -> None:
    """验证无书签文档的自动 EPUB 标识与修复前真实包成员保持一致。"""
    package = _zip_xml(render_epub(_visual_document(prefix, anchors=(None, None))), "EPUB/package.opf")
    assert package.xpath('string(//*[local-name()="identifier"])') == NO_ANCHOR_EPUB_IDS[prefix]


@pytest.mark.parametrize("prefix", PREFIXES)
def test_raw_visual_annotations_and_unmatched_fallback_preserve_anchor(prefix: str) -> None:
    """验证通用说明分组及未匹配回退为正文的路径不丢失书签。"""
    middle = _visual_document(prefix)
    visual = middle.pages[1].blocks[0]
    body = next(child for child in visual.content if str(child.type).endswith("body"))
    raw_body = body.to_dict()
    raw_body["type"] = prefix
    if prefix == "code":
        raw_body["guess_lang"] = "python"
    caption = {"type": "caption", "content": _inline("Caption"), "anchor": "caption-target"}
    note = {"type": "footnote", "content": _inline("Footnote"), "anchor": "footnote-target"}
    grouped = blocks_to_page_info([caption, raw_body, note], use_bbox=False)
    assert [child.anchor for child in grouped.blocks[0].content if str(child.type).endswith(("caption", "footnote"))] == [
        "caption-target",
        "footnote-target",
    ]
    unmatched = blocks_to_page_info([caption], use_bbox=False)
    assert isinstance(unmatched.blocks[0], TextBlock)
    assert unmatched.blocks[0].anchor == "caption-target"


@pytest.mark.parametrize("prefix", ["image", "chart"])
def test_annotation_index_discovery_accepts_a_raw_body_without_text(prefix: str) -> None:
    """验证目录目标扫描仍接受内容为空的原生图片和图表主体。"""
    from docvortex.postprocess.lists import fix_office_index_title_blocks

    blocks = [{"type": prefix, "content": None, "image_base64": "data:image/png;base64,AA=="}]
    fix_office_index_title_blocks([blocks])
    assert blocks[0]["content"] is None


@pytest.mark.parametrize("prefix", PREFIXES)
@pytest.mark.parametrize("mode", [RenderMode.DEFAULT, RenderMode.FULL])
def test_annotation_markdown_html_and_exact_html_roundtrip(prefix: str, mode: RenderMode) -> None:
    """验证说明目标、行内和目录链接以及 canonical HTML 往返均保留原始书签。"""
    middle = _visual_document(prefix)
    markdown = render_markdown(middle, mode=mode)
    for role in ("caption", "footnote"):
        assert f'<a id="{role} target"></a>\n{prefix} {role}' in markdown
        assert f"[Jump {role}](#{role}%20target)" in markdown
        assert f"[{prefix} {role}](#{role}%20target)" in markdown
    rendered = render_html(middle, mode=mode)
    soup = BeautifulSoup(rendered, "html.parser")
    for role in ("caption", "footnote"):
        target = soup.find("p", attrs={"data-block-type": f"{prefix}_{role}"})
        assert target["data-anchor"] == f"{role} target"
        assert target["id"] == f"{role}-target"
        assert len(soup.find_all("a", href=f"#{role}-target")) == 2
    decoded, _ = analyze_native_test_document(rendered.encode(), file_suffix="html")
    decoded_visual = next(block for page in decoded.pages for block in page.blocks if block.type == prefix)
    assert [
        (child.type, child.anchor) for child in decoded_visual.content if str(child.type).endswith(("caption", "footnote"))
    ] == [(f"{prefix}_caption", "caption target"), (f"{prefix}_footnote", "footnote target")]
    decoded_text = next(block for page in decoded.pages for block in page.blocks if isinstance(block, TextBlock))
    assert [span.url for span in decoded_text.content] == ["#caption target", "#footnote target"]
    decoded_index = next(block for page in decoded.pages for block in page.blocks if isinstance(block, IndexBlock))
    assert [child.anchor for child in decoded_index.content] == ["caption target", "footnote target"]


@pytest.mark.parametrize("prefix", PREFIXES)
def test_annotation_docx_epub_latex_and_structured_targets(prefix: str) -> None:
    """验证 Word 书签、EPUB id、LaTeX 目标和结构化说明字段落在正确说明上。"""
    middle = _visual_document(prefix)
    document = _zip_xml(render_docx(middle), "word/document.xml")
    for role in ("caption", "footnote"):
        bookmark = document.xpath(f"//w:bookmarkStart[@w:name='{role}_target']", namespaces=WORD_NS)
        assert len(bookmark) == 1
        assert "".join(bookmark[0].getparent().xpath(".//w:t/text()", namespaces=WORD_NS)) == f"{prefix} {role}"
        assert len(document.xpath(f"//w:hyperlink[@w:anchor='{role}_target']", namespaces=WORD_NS)) == 2
    epub = render_epub(middle, modified_at=datetime(2026, 1, 1, tzinfo=timezone.utc))
    content = _zip_xml(epub, "EPUB/text/content.xhtml")
    for role in ("caption", "footnote"):
        assert content.xpath(f"//x:p[@id='{role}-target']/text()", namespaces=XHTML_NS) == [f"{prefix} {role}"]
        assert len(content.xpath(f"//x:a[@href='#{role}-target']", namespaces=XHTML_NS)) == 2
    latex = render_latex(middle)
    assert latex.count(r"\hypertarget{") == 2
    assert latex.count(r"\hyperlink{") == 4
    structured = render_structured_content(middle)
    visual = structured["pages"][1]["blocks"][0]
    assert visual["captions"] == [{"bbox": [0.1, 0.12, 0.9, 0.18], "anchor": "caption target", "content": f"{prefix} caption"}]
    assert visual["footnotes"] == [
        {"bbox": [0.1, 0.5, 0.9, 0.58], "anchor": "footnote target", "content": f"{prefix} footnote"}
    ]


@pytest.mark.parametrize("prefix", PREFIXES)
@pytest.mark.parametrize("layout", [PdfLayout.AUTO, PdfLayout.ORIGINAL])
def test_annotation_pdf_destinations_point_to_the_description_page(prefix: str, layout: PdfLayout) -> None:
    """验证两种 PDF 布局中的前向跳转指向各自说明，而不假设重排保留源分页。"""
    middle = _visual_document(prefix)
    if layout is PdfLayout.ORIGINAL:
        middle.metadata.file_suffix = "pdf"
        middle.extensions = {
            "docvortex_layout": {
                "version": 1,
                "pages": [{"page_idx": i, "width_pt": 595.28, "height_pt": 841.89} for i in (0, 1)],
            }
        }
    reader = PdfReader(BytesIO(render_pdf(middle, layout=layout)))
    links = [ref.get_object() for page in reader.pages for ref in (page.get("/Annots") or []) if ref.get_object().get("/Dest")]
    assert len(links) == 4
    target_page = reader.pages[-1]
    assert all(link["/Dest"][0].idnum == target_page.indirect_reference.idnum for link in links)
    assert f"{prefix} footnote" in target_page.extract_text()
    # 目标纵坐标必须不同，避免两个链接仅指向同一页而没有落到各自说明。
    assert len({round(float(link["/Dest"][3]), 2) for link in links}) == 2


@pytest.mark.parametrize("prefix", ["table", "image"])
@pytest.mark.parametrize("seq_caption", [False, True])
@pytest.mark.parametrize("force_fallback", [False, True])
def test_real_docx_bookmarked_caption_preserves_content_and_forward_link(
    prefix: str,
    seq_caption: bool,
    force_fallback: bool,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """通过公共解析入口覆盖普通标题、SEQ 标题及表格完整上下文回退。"""
    if force_fallback:
        from unittest.mock import Mock

        monkeypatch.setattr(DocxConverter, "_preparse_tables_with_mammoth", Mock(return_value=[]))
    document = Document()
    paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "_Hlk46718182")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Forward caption link"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    if prefix == "table":
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell survives"
    else:
        document.add_picture(BytesIO(_png()))
    caption_text = "Table 1 Caption" if prefix == "table" else "Figure 1 Caption"
    caption = document.add_paragraph(caption_text, style="Caption" if seq_caption else "Normal")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), "_Hlk46718182")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    caption._p.insert(0, start)
    caption._p.append(end)
    if seq_caption:
        run = OxmlElement("w:r")
        field = OxmlElement("w:instrText")
        field.text = " SEQ Table \\* ARABIC "
        run.append(field)
        caption._p.append(run)
    document.add_paragraph("Following text survives")
    output = BytesIO()
    document.save(output)
    result = parse(output.getvalue(), file_suffix="docx")
    visual = next(block for page in result.middle_json.pages for block in page.blocks if block.type == prefix)
    annotation = next(child for child in visual.content if child.type == f"{prefix}_caption")
    assert annotation.anchor == "_Hlk46718182"
    assert caption_text in "".join(span.content for span in annotation.content)
    if prefix == "table":
        assert "Cell survives" in next(child for child in visual.content if child.type == "table_body").content
    markdown = render_markdown(result.middle_json)
    assert "Following text survives" in markdown
    assert "[Forward caption link](#_Hlk46718182)" in markdown
    raw = analyze(output.getvalue(), file_suffix="docx").model_json
    raw_caption = next(block for page in raw.pages for block in page if block.get("anchor") == "_Hlk46718182")
    assert raw_caption["type"] == ("caption" if seq_caption else "text")


def test_duplicate_empty_special_and_missing_annotation_anchors() -> None:
    """验证特殊字符转义、重复书签首目标、空说明及不存在的目录目标。"""
    anchor = 'target "& 空格'
    middle = _visual_document("table", anchors=(anchor, anchor))
    visual = middle.pages[1].blocks[0]
    empty = type(visual.content[0]).model_validate({**visual.content[0].to_dict(), "content": _inline(" "), "index": 3})
    visual.content.insert(0, empty)
    missing = TextBlock(type="text", anchor="missing", content=_inline("Missing target"))
    middle.pages[0].blocks[1].content.append(missing)
    soup = BeautifulSoup(render_html(middle), "html.parser")
    targets = soup.find_all("p", id=True)
    assert len(targets) == 1 and targets[0].get_text() == "table caption"
    assert soup.find("a", href="#missing") is None
    assert render_markdown(middle).count('<a id="') == 1
    import markdown

    parsed_markdown = BeautifulSoup(markdown.markdown(render_markdown(middle)), "html.parser")
    assert all(unquote(link["href"][1:]) == anchor for link in parsed_markdown.find_all("a", href=True))
    decoded, _ = analyze_native_test_document(render_html(middle).encode(), file_suffix="html")
    decoded_text = next(block for page in decoded.pages for block in page.blocks if isinstance(block, TextBlock))
    assert all(span.url == f"#{anchor}" for span in decoded_text.content)
    document = _zip_xml(render_docx(middle), "word/document.xml")
    bookmarks = document.xpath("//w:bookmarkStart", namespaces=WORD_NS)
    assert len(bookmarks) == 1
    assert "".join(bookmarks[0].getparent().xpath(".//w:t/text()", namespaces=WORD_NS)) == "table caption"
    epub = _zip_xml(render_epub(middle), "EPUB/text/content.xhtml")
    assert len(epub.xpath("//x:figure/x:p[@id]", namespaces=XHTML_NS)) == 1
    assert render_latex(middle).count(r"\hypertarget{") == 1
    assert len(PdfReader(BytesIO(render_pdf(middle))).pages) >= 1


@pytest.mark.parametrize("text_first", [True, False])
def test_annotation_and_body_same_anchor_share_one_target(text_first: bool) -> None:
    """验证说明与正文同名时按文档顺序选取首个目标，EPUB 不产生重复 id。"""
    middle = _visual_document("table", anchors=("same", None))
    body = TextBlock(type="text", index=2, anchor="same", content=_inline("Body target"))
    middle.pages[0 if text_first else 1].blocks.append(body)
    expected = "Body target" if text_first else "table caption"
    soup = BeautifulSoup(render_html(middle), "html.parser")
    assert soup.find(id="same").get_text() == expected
    assert len(soup.find_all(id="same")) == 1
    content = _zip_xml(render_epub(middle), "EPUB/text/content.xhtml")
    target = content.xpath("//*[@id='same']")
    assert len(target) == 1 and "".join(target[0].itertext()) == expected


def test_empty_body_does_not_claim_the_visual_annotation_target() -> None:
    """验证空正文与说明同名时，所有格式均把跳转目标挂到可见说明上。"""
    middle = _visual_document("table", anchors=("same", None))
    middle.pages[0].blocks.insert(0, TextBlock(type="text", index=2, anchor="same", content=_inline(" ")))
    soup = BeautifulSoup(render_html(middle), "html.parser")
    assert soup.find(id="same").get_text() == "table caption"
    document = _zip_xml(render_docx(middle), "word/document.xml")
    bookmark = document.xpath("//w:bookmarkStart[@w:name='same']", namespaces=WORD_NS)[0]
    assert "".join(bookmark.getparent().xpath(".//w:t/text()", namespaces=WORD_NS)) == "table caption"
    assert '<a id="same"></a>\ntable caption' in render_markdown(middle)
    assert r"{\small\itshape \hypertarget{" in render_latex(middle)
    epub = _zip_xml(render_epub(middle), "EPUB/text/content.xhtml")
    assert epub.xpath("//x:p[@id='same']/text()", namespaces=XHTML_NS) == ["table caption"]
    reader = PdfReader(BytesIO(render_pdf(middle)))
    destinations = [ref.get_object()["/Dest"] for ref in reader.pages[0].get("/Annots", []) if "/Dest" in ref.get_object()]
    assert len({float(destination[3]) for destination in destinations}) == 1
