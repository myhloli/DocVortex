"""验证 DOCX 表格预解析只复用本次文档，并减少无表和失配表格的重复读取。"""

from io import BytesIO
from zipfile import BadZipFile

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
import mammoth
import pytest

from docvortex.analyzers.native.office.docx import docx_converter as converter_module
from docvortex.analyzers.native.office.docx import tables as tables_module
from docvortex.analyzers.native.office.docx.docx_converter import DocxConverter
from docvortex.schema import BlockType


def _document_bytes(*, paragraph: str = "", tables: tuple[str, ...] = ()) -> bytes:
    """构造可区分正文和顶层表格内容的 DOCX。"""
    document = Document()
    if paragraph:
        document.add_paragraph(paragraph)
    for text in tables:
        document.add_table(rows=1, cols=1).cell(0, 0).text = text
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("table_values", [(), ("first table", "second table")])
def test_convert_opens_ooxml_document_once_and_skips_mammoth_without_tables(monkeypatch, table_values) -> None:
    """正文与预解析共用一次 OOXML 读取，无顶层表格时不进入 Mammoth。"""
    payload = _document_bytes(paragraph="paragraph", tables=table_values)
    opened = []
    converted = []
    original_convert = mammoth.convert_to_html

    def open_document(stream):
        """记录完整 OOXML 文档的实际读取，不替换其解析结果。"""
        opened.append(stream)
        return Document(stream)

    def convert_tables(*args, **kwargs):
        """记录 Mammoth 的实际转换调用并保留完整上下文。"""
        converted.append(True)
        return original_convert(*args, **kwargs)

    monkeypatch.setattr(converter_module, "Document", open_document)
    monkeypatch.setattr(tables_module, "Document", open_document)
    monkeypatch.setattr(mammoth, "convert_to_html", convert_tables)
    converter = DocxConverter()
    converter.convert(BytesIO(payload))

    assert len(opened) == 1
    assert len(converted) == bool(table_values)
    tables = [block["content"] for page in converter.pages for block in page if block["type"] == BlockType.TABLE]
    assert [BeautifulSoup(table, "html.parser").get_text() for table in tables] == list(table_values)
    assert converter._fallback_docx_bytes is None


def test_standalone_preparse_after_convert_reads_its_own_document() -> None:
    """独立预解析第二份字节时不能借用转换结束后仍可访问的第一份文档。"""
    converter = DocxConverter()
    converter.convert(BytesIO(_document_bytes(tables=("old table",))))
    result = converter._preparse_tables_with_mammoth(_document_bytes(tables=("new table",)))

    assert len(result) == 1
    assert BeautifulSoup(result[0], "html.parser").get_text() == "new table"
    assert converter._fallback_docx_bytes is None


def test_failed_convert_does_not_reuse_document_on_retry() -> None:
    """一次打开失败后复用实例，表格和原包状态必须来自后续成功输入。"""
    converter = DocxConverter()
    converter.convert(BytesIO(_document_bytes(tables=("old table",))))
    with pytest.raises(BadZipFile):
        converter.convert(BytesIO(b"broken docx"))
    assert converter.docx_obj is None
    assert converter._fallback_docx_bytes is None
    converter.convert(BytesIO(_document_bytes(tables=("retry table",))))

    tables = [block["content"] for page in converter.pages for block in page if block["type"] == BlockType.TABLE]
    assert len(tables) == 1
    assert BeautifulSoup(tables[0], "html.parser").get_text() == "retry table"


def test_table_alignment_reuses_unmodified_signatures_after_missing_match(monkeypatch) -> None:
    """缺失表格匹配导致再次扫描候选时，每棵 XML 和 HTML 表只统计一次。"""
    document = Document(BytesIO(_document_bytes(tables=("missing table", "matching table"))))
    xml_tables = list(document.element.body.iterchildren(tag=qn("w:tbl")))
    soup = BeautifulSoup(
        "<table><tr><td>extra table</td></tr></table><table><tr><td>matching table</td></tr></table>",
        "html.parser",
    )
    converter = DocxConverter()
    calls = {"xml": 0, "html": 0}
    original_xml = converter._xml_table_signature
    original_html = converter._html_table_signature

    def xml_signature(table):
        """统计源表签名提取次数并保留现有字符匹配语义。"""
        calls["xml"] += 1
        return original_xml(table)

    def html_signature(table):
        """统计候选表签名提取次数并保留现有 HTML 过滤语义。"""
        calls["html"] += 1
        return original_html(table)

    monkeypatch.setattr(converter, "_xml_table_signature", xml_signature)
    monkeypatch.setattr(converter, "_html_table_signature", html_signature)
    result = converter._align_mammoth_tables_to_xml_tables(soup.find_all("table"), xml_tables, document.part)

    assert calls == {"xml": 2, "html": 2}
    assert result[0] is None
    assert BeautifulSoup(result[1], "html.parser").get_text() == "matching table"
