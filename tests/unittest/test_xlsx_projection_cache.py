"""XLSX 只读投影期间的样式与 DISPIMG 复用边界。"""

from io import BytesIO
from unittest.mock import Mock
from zipfile import ZipFile

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image
import pytest

from _image_mtef_test_utils import build_gif_with_mtef
from _mtef_v5_test_utils import v5_formula_corpus
from docvortex.analyzers.native.office.spreadsheet.projector import SpreadsheetProjector
from docvortex.analyzers.native.office.xlsx import xlsx_converter as module
from docvortex.analyzers.native.office.xlsx.xlsx_converter import XlsxConverter


def _bitmap(color: str) -> bytes:
    """生成内容不同的小位图，验证同名包成员不会跨文档复用。"""
    stream = BytesIO()
    Image.new("RGB", (4, 3), color).save(stream, "PNG")
    return stream.getvalue()


def _dispimg_workbook(payload: bytes, *, part_name: str = "image.png") -> bytes:
    """构造六格引用同一 WPS 图片的完整工作簿以覆盖三档评分和物化。"""
    workbook = Workbook()
    for row in range(1, 4):
        for col in range(1, 3):
            workbook.active.cell(row, col, 'DISPIMG("image-id")')
    source = BytesIO()
    workbook.save(source)
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w") as archive:
        for info in original.infolist():
            archive.writestr(info, original.read(info.filename))
        archive.writestr(
            "xl/cellimages.xml",
            '<etc:cellImages xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData" '
            'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<etc:cellImage><xdr:cNvPr name="image-id"/><a:blip r:embed="rId1"/></etc:cellImage>'
            "</etc:cellImages>",
        )
        archive.writestr(
            "xl/_rels/cellimages.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Target="media/{part_name}"/></Relationships>',
        )
        archive.writestr(f"xl/media/{part_name}", payload)
    return target.getvalue()


def test_repeated_styles_are_projected_once_and_keep_style_values(monkeypatch: pytest.MonkeyPatch) -> None:
    """相同样式只投影一次，同时保留表格内部样式值和最终单元格顺序。"""
    workbook = Workbook()
    for row in range(1, 11):
        for col in range(1, 3):
            cell = workbook.active.cell(row, col, f"{row}:{col}")
            cell.font = Font(b=col == 1, i=col == 2)
            cell.fill = PatternFill(patternType="solid", fgColor="FF112233" if col == 1 else "FF445566")
    stream = BytesIO()
    workbook.save(stream)
    original = SpreadsheetProjector._extract_cell_style
    observed_calls = []

    def observed(projector: SpreadsheetProjector, cell: object) -> dict:
        """保留方法绑定并记录实际样式提取次数。"""
        result = original(projector, cell)
        observed_calls.append(result)
        return result

    monkeypatch.setattr(SpreadsheetProjector, "_extract_cell_style", observed)
    converter = XlsxConverter()
    converter.convert(BytesIO(stream.getvalue()))

    assert len(observed_calls) == 2
    cells = BeautifulSoup(converter.pages[0][0]["content"], "html.parser").find_all(["th", "td"])
    assert [cell.get_text() for cell in cells] == [f"{row}:{col}" for row in range(1, 11) for col in range(1, 3)]
    assert observed_calls == [
        {"font-weight": "bold", "background-color": "#112233"},
        {"font-style": "italic", "background-color": "#445566"},
    ]
    assert not converter._projection_cache_enabled
    assert converter._cell_style_cache == {}


def test_style_results_do_not_alias_or_cross_workbooks() -> None:
    """复用结果必须独立，赋新样式、其他工作簿及非转换期查询仍读取当前值。"""
    converter = XlsxConverter()
    converter.workbook = Workbook()
    converter._projection_cache_enabled = True
    cell = converter.workbook.active["A1"]
    cell.font = Font(b=True)
    first = converter._extract_cell_style(cell)
    first["font-weight"] = "changed"
    assert converter._extract_cell_style(cell)["font-weight"] == "bold"
    cell.font = Font(i=True)
    assert converter._extract_cell_style(cell) == {"font-style": "italic"}

    other = Workbook().active["A1"]
    other.font = Font(strike=True)
    assert converter._extract_cell_style(other) == {"text-decoration": "line-through"}
    converter._reset_conversion_caches()
    cell.font = Font(u="single")
    assert converter._extract_cell_style(cell) == {"text-decoration": "underline"}


def test_dispimg_scoring_and_materialization_share_successful_bitmap(monkeypatch: pytest.MonkeyPatch) -> None:
    """三档评分及六格物化只编码一次图片，转换结束释放缓存并允许实例复用。"""
    original = module.serialize_office_image
    serialize = Mock(wraps=original)
    monkeypatch.setattr(module, "serialize_office_image", serialize)
    converter = XlsxConverter()
    outputs = []
    for color in ("red", "blue"):
        payload = _bitmap(color)
        converter.convert(BytesIO(_dispimg_workbook(payload)))
        images = BeautifulSoup(converter.pages[0][0]["content"], "html.parser").find_all("img")
        assert len(images) == 6
        expected = original(payload, part_name="xl/media/image.png", content_type=None)
        assert all(image["src"] == expected for image in images)
        outputs.append(expected)
        assert converter._resolved_cell_image_cache == {}
        assert not converter._projection_cache_enabled
    assert serialize.call_count == 2
    assert outputs[0] != outputs[1]


def test_dispimg_formula_reuse_preserves_unique_candidate_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    """成功公式复用不能重复收费，也不能转而输出图片预览。"""
    _name, mtef, expected = v5_formula_corpus()[0]
    payload = build_gif_with_mtef(mtef)
    serialize = Mock(side_effect=AssertionError("formula must not serialize its preview"))
    monkeypatch.setattr(module, "serialize_office_image", serialize)
    converter = XlsxConverter()
    converter.convert(BytesIO(_dispimg_workbook(payload, part_name="image.gif")))
    table = converter.pages[0][0]["content"]
    assert table.count(f"<eq>{expected}</eq>") == 6
    assert converter._image_equation_decoder.candidate_total_bytes == len(mtef)
    serialize.assert_not_called()


@pytest.mark.parametrize("failure", ["missing", "invalid", "vector"])
def test_unsuccessful_or_vector_images_keep_repeated_diagnostics(monkeypatch: pytest.MonkeyPatch, failure: str) -> None:
    """缺失、损坏及矢量占位路径仍逐次执行并保留原有告警。"""
    payload = _bitmap("red") if failure == "vector" else b"broken"
    part_name = "image.emf" if failure == "vector" else "image.png"
    package = BytesIO()
    with ZipFile(package, "w") as archive:
        archive.writestr(f"xl/media/{part_name}", payload)
    converter = XlsxConverter()
    converter._projection_cache_enabled = True
    converter.cell_image_map = {"image-id": f"media/{part_name}" if failure != "missing" else "media/absent.png"}
    warning = Mock()
    monkeypatch.setattr(module.logger, "warning", warning)
    with ZipFile(BytesIO(package.getvalue())) as archive:
        converter.zf = archive
        first = converter._resolve_cell_image('DISPIMG("image-id")')
        first_count = warning.call_count
        second = converter._resolve_cell_image('DISPIMG("image-id")')
    assert first == second
    assert first_count > 0
    assert warning.call_count == first_count * 2
    assert converter._resolved_cell_image_cache == {}


def test_failed_conversion_releases_all_projection_caches(monkeypatch: pytest.MonkeyPatch) -> None:
    """已填充缓存的转换失败后仍清理状态，随后同实例可正常读取另一份文档。"""
    converter = XlsxConverter()
    original = converter._convert_sheet

    def fail_after_projection(sheet: object) -> None:
        """在真实投影已填充两种缓存后模拟后续失败。"""
        original(sheet)
        assert converter._cell_style_cache
        assert converter._resolved_cell_image_cache
        raise ValueError("controlled projection failure")

    monkeypatch.setattr(converter, "_convert_sheet", fail_after_projection)
    with pytest.raises(ValueError, match="controlled projection failure"):
        converter.convert(BytesIO(_dispimg_workbook(_bitmap("red"))))
    assert not converter._projection_cache_enabled
    assert converter._cell_style_cache == {}
    assert converter._resolved_cell_image_cache == {}
    monkeypatch.setattr(converter, "_convert_sheet", original)
    converter.convert(BytesIO(_dispimg_workbook(_bitmap("blue"))))
    assert len(BeautifulSoup(converter.pages[0][0]["content"], "html.parser").find_all("img")) == 6
