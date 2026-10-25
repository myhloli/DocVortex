"""一次生成并冻结非 PDF 真实/压力输入，配对版本只读取同一份字节清单。"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo


ROOT = Path(__file__).resolve().parents[1]
SUFFIXES = frozenset(
    {"doc", "docx", "ppt", "pptx", "xls", "xlsx", "rtf", "csv", "tsv", "html", "mhtml", "epub", "ofd", "odt", "ods", "odp"}
)


def _stable_zip(payload: bytes) -> bytes:
    """固定容器时间和成员顺序，文档语义与原成员字节保持不变。"""
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w") as target:
        names = sorted(source.namelist(), key=lambda name: (name != "mimetype", name))
        for name in names:
            member = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            member.compress_type = ZIP_STORED if name == "mimetype" else ZIP_DEFLATED
            target.writestr(member, source.read(name))
    return output.getvalue()


def generate_pressure(kind: str, parameters: dict) -> bytes:
    """使用显式规模生成真实格式字节，生成成本不进入解析计时。"""
    if any(type(value) is not int or value < 1 for value in parameters.values()):
        raise ValueError("Pressure parameters must be positive integers")
    if kind in {"csv", "tsv"}:
        rows, columns = parameters["rows"], parameters["columns"]
        separator = "," if kind == "csv" else "\t"
        lines = [separator.join(f"column_{column}" for column in range(columns))]
        lines.extend(separator.join(f"value_{row}_{column}" for column in range(columns)) for row in range(rows))
        return ("\n".join(lines) + "\n").encode("utf-8")
    if kind == "rtf":
        paragraphs = "".join(
            rf"{{\b Paragraph {index}}}\tab value {index} \u20013?\u25991?\par " for index in range(parameters["paragraphs"])
        )
        return (r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0 Arial;}}" + paragraphs + "}").encode("ascii")
    if kind == "ods":
        rows, columns = parameters["rows"], parameters["columns"]
        content = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
            '<office:body><office:spreadsheet><table:table table:name="Data">'
            + "".join(
                "<table:table-row>"
                + "".join(
                    f'<table:table-cell office:value-type="string"><text:p>value_{row}_{column}</text:p></table:table-cell>'
                    for column in range(columns)
                )
                + "</table:table-row>"
                for row in range(rows)
            )
            + "</table:table></office:spreadsheet></office:body></office:document-content>"
        )
        output = BytesIO()
        with ZipFile(output, "w") as package:
            package.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            package.writestr("content.xml", content)
        return _stable_zip(output.getvalue())
    if kind == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        book.properties.created = book.properties.modified = datetime(2000, 1, 1)
        sheet = book.active
        sheet.title = "Data"
        for row in range(parameters["rows"]):
            sheet.append([f"value_{row}_{column}" for column in range(parameters["columns"])])
        output = BytesIO()
        book.save(output)
        # openpyxl 写出时更新 modified；冻结该核心属性以保证跨时间生成一致。
        rebuilt = BytesIO()
        with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(rebuilt, "w") as target:
            for name in source.namelist():
                content = source.read(name)
                if name == "docProps/core.xml":
                    content = re.sub(
                        rb"(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)", rb"\g<1>2000-01-01T00:00:00Z\2", content
                    )
                target.writestr(name, content)
        return _stable_zip(rebuilt.getvalue())
    if kind in {"docx_paragraphs", "docx_tables"}:
        from docx import Document

        document = Document()
        document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        if kind == "docx_paragraphs":
            for index in range(parameters["paragraphs"]):
                document.add_paragraph(f"Paragraph {index}: fixed native document benchmark content.")
        else:
            for index in range(parameters["tables"]):
                table = document.add_table(rows=parameters["rows"], cols=parameters["columns"])
                for row_index, row in enumerate(table.rows):
                    for column_index, cell in enumerate(row.cells):
                        cell.text = f"table_{index}_row_{row_index}_column_{column_index}"
        output = BytesIO()
        document.save(output)
        return _stable_zip(output.getvalue())
    raise ValueError(f"Unknown pressure generator: {kind}")


def load_frozen_manifest(path: Path) -> list[dict]:
    """每次读取都核对完整字节身份与解析参数，拒绝被修改或含隐藏依赖的清单。"""
    value = json.loads(path.read_text(encoding="utf-8"))
    if value.get("schema") != "docvortex.benchmark.inputs" or value.get("schema_version") != 1:
        raise ValueError("Unsupported frozen input manifest")
    inputs = value.get("inputs")
    if not isinstance(inputs, list) or not inputs:
        raise ValueError("Frozen input manifest must not be empty")
    names = set()
    for entry in inputs:
        if not isinstance(entry, dict) or set(entry) != {
            "name",
            "path",
            "file_suffix",
            "category",
            "source_context",
            "bytes",
            "sha256",
        }:
            raise ValueError("Frozen input manifest entry has missing or unknown fields")
        name = entry["name"]
        if not isinstance(name, str) or not name or name in names:
            raise ValueError("Frozen input names must be nonempty and unique")
        names.add(name)
        if entry["file_suffix"] not in SUFFIXES or entry["category"] not in {"real", "fixture", "pressure"}:
            raise ValueError("Invalid frozen input format or category")
        context = entry["source_context"]
        if context is not None and (
            not isinstance(context, dict)
            or set(context) - {"source_uri", "transport_encoding"}
            or any(item is not None and not isinstance(item, str) for item in context.values())
        ):
            raise ValueError("Frozen source_context must not depend on external local files")
        source = Path(entry["path"])
        if not source.is_absolute() or not source.is_file():
            raise ValueError(f"Frozen input must be an absolute file path: {source}")
        payload = source.read_bytes()
        if (
            type(entry["bytes"]) is not int
            or entry["bytes"] != len(payload)
            or entry["sha256"] != hashlib.sha256(payload).hexdigest()
        ):
            raise ValueError(f"Frozen input bytes changed: {name}")
    return inputs


def freeze_corpus(spec_path: Path, output: Path, *, root: Path = ROOT, selected: list[str] | None = None) -> Path:
    """把真实样本和压力输入复制到新目录，供两个源码版本共同解析。"""
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    entries = spec["cases"]
    if selected is not None:
        missing = set(selected) - {entry["name"] for entry in entries}
        if missing:
            raise ValueError(f"Unknown corpus cases: {sorted(missing)}")
        entries = [entry for entry in entries if entry["name"] in selected]
    if not entries or len({entry["name"] for entry in entries}) != len(entries):
        raise ValueError("Corpus must contain unique nonempty cases")
    output.mkdir(parents=True, exist_ok=False)
    inputs = []
    for entry in entries:
        name, suffix = entry["name"], entry["file_suffix"]
        if not re.fullmatch(r"[A-Za-z0-9_-]+", name) or suffix not in SUFFIXES:
            raise ValueError("Invalid corpus name or file_suffix")
        payload = (
            (root / entry["path"]).read_bytes()
            if "path" in entry
            else generate_pressure(entry["generator"], entry["parameters"])
        )
        path = output / f"{name}.{suffix}"
        path.write_bytes(payload)
        inputs.append(
            {
                "name": name,
                "path": str(path.resolve()),
                "file_suffix": suffix,
                "category": entry["category"],
                "source_context": entry.get("source_context"),
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
            }
        )
    manifest = output / "manifest.json"
    value = {
        "schema": "docvortex.benchmark.inputs",
        "schema_version": 1,
        "input_mode": "bytes_with_explicit_file_suffix",
        "spec_sha256": hashlib.sha256(spec_path.read_bytes()).hexdigest(),
        "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "inputs": inputs,
    }
    manifest.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    load_frozen_manifest(manifest)
    return manifest


def main() -> None:
    """按固定规格冻结输入，可选择少量案例用于批次验收或工具冒烟。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", type=Path, default=ROOT / "tests/fixtures/nonpdf_performance_manifest.json")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--case", action="append")
    args = parser.parse_args()
    print(freeze_corpus(args.spec, args.output, selected=args.case))


if __name__ == "__main__":
    main()
